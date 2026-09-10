"""Reproducible, lossless source normalization for CUMCM 2026 problem C.

Run: python src/preprocess.py
Only reads raw/*.xlsx and templates/*.xlsx. Writes processed/ and diagnostics/.
Time convention: unzoned local clock, samples as preceding interval averages
for approximate energies. This is an explicit provisional modeling assumption.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import platform
from datetime import datetime, time
from pathlib import Path
import re

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
STEP = pd.Timedelta(minutes=10)
YEAR_START = pd.Timestamp('2025-01-01')
YEAR_END = pd.Timestamp('2026-01-01')


def require(condition, message):
    if not condition:
        raise ValueError(message)


def write_json(path, obj):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')


def write_csv(frame, path):
    frame.to_csv(path, index=False, encoding='utf-8-sig', date_format='%Y-%m-%d %H:%M:%S')


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_book(path):
    # openpyxl is used strictly for source extraction, never workbook authoring.
    w = openpyxl.load_workbook(path, read_only=True, data_only=False)
    result = {}
    for s in w:
        rows = []
        for row in s.iter_rows():
            vals = []
            for c in row:
                require(c.data_type not in ('f', 'e'), f'Unexpected formula/error: {path.name}/{s.title}/{getattr(c, "coordinate", "blank cell")}')
                vals.append(c.value)
            rows.append(tuple(vals))
        result[s.title] = rows
    w.close()
    return result


def clock_minutes(value):
    """Normalize Excel time, text clock and explicit next-day suffix without UTC conversion."""
    if isinstance(value, (time, datetime)):
        require(value.second == 0, f'Nonzero seconds in source time {value}')
        return value.hour * 60 + value.minute
    if isinstance(value, (float, int)):
        require(0 <= value <= 1, f'Unexpected numeric clock {value}')
        return int(round(value * 1440))
    text = str(value).strip()
    match = re.fullmatch(r'(\d{1,2}):(\d{2})(?::00)?(?:\+(\d+))?', text)
    require(match is not None, f'Unrecognized clock {value!r}')
    h, m, extra = int(match[1]), int(match[2]), int(match[3] or 0)
    require(0 <= h <= 24 and 0 <= m < 60 and (h < 24 or m == 0), f'Invalid clock {value!r}')
    return h * 60 + m + extra * 1440


def numeric_matrix(rows, description):
    a = np.asarray(rows, dtype=float)
    require(np.isfinite(a).all(), f'Missing/nonfinite numeric value in {description}; no automatic imputation allowed')
    require((a >= 0).all(), f'Negative source value in {description}; manual review required')
    return a


def wide_daily(rows, description):
    require(len(rows) == 366 and all(len(r) == 145 for r in rows), f'Unexpected dimensions: {description}')
    minutes = [clock_minutes(v) for v in rows[0][1:]]
    require(minutes == list(range(10, 1441, 10)), f'Unexpected clock sequence: {description}')
    dates = pd.DatetimeIndex([pd.Timestamp(r[0]) for r in rows[1:]])
    require(dates.equals(pd.date_range('2025-01-01', '2025-12-31', freq='D')), f'Missing/duplicate/unordered dates: {description}')
    return numeric_matrix([r[1:] for r in rows[1:]], description), dates, minutes


def build_actuals(books):
    l, dates, minutes = wide_daily(books[2]['小区负载'], 'load')
    p, dates2, minutes2 = wide_daily(books[2]['光伏发电实际功率'], 'PV')
    r, dates3, minutes3 = wide_daily(books[4]['Sheet1'], 'price')
    require(dates.equals(dates2) and dates.equals(dates3) and minutes == minutes2 == minutes3, 'Actual source axes disagree')
    frame = pd.DataFrame({
        'source_date': np.repeat(dates.to_numpy(), 144),
        'slot': np.tile(np.arange(1, 145), 365),
        'sample_offset_minutes': np.tile(minutes, 365),
        'load_kw': l.ravel(), 'pv_kw': p.ravel(), 'price_actual_yuan_per_kwh': r.ravel(),
        'source_excel_row': np.repeat(np.arange(2, 367), 144),
        'source_excel_column': np.tile(np.arange(2, 146), 365),
    })
    frame['sample_time'] = frame['source_date'] + pd.to_timedelta(frame['sample_offset_minutes'], unit='min')
    frame['interval_end'] = frame['sample_time']
    frame['interval_start'] = frame['interval_end'] - STEP
    frame['observed_available_at_assumed'] = frame['sample_time']
    frame['net_load_kw'] = frame['load_kw'] - frame['pv_kw']
    for col in ('load', 'pv', 'net_load'):
        frame[f'{col}_kwh_approx'] = frame[f'{col}_kw'] / 6
    frame['pv_exceeds_load'] = frame['pv_kw'] > frame['load_kw']
    frame['period_role'] = np.where(frame['source_date'] < pd.Timestamp('2025-02-01'), 'initial_history', 'evaluation')
    return frame


def build_q1(books):
    rows = books[1]['Sheet1']
    require(len(rows) == 145 and all(len(r) == 4 for r in rows), 'Unexpected attachment1 dimensions')
    minutes = [clock_minutes(r[0]) for r in rows[1:]]
    require(minutes == list(range(10, 1441, 10)), 'Q1 clock mismatch')
    a = numeric_matrix([r[1:] for r in rows[1:]], 'Q1')
    frame = pd.DataFrame({
        'slot': np.arange(1, 145), 'source_time_label': [str(r[0]) for r in rows[1:]],
        'sample_offset_minutes': minutes, 'interval_start_minute': np.arange(0, 1440, 10),
        'interval_end_minute': minutes, 'price_yuan_per_kwh': a[:, 0],
        'load_kw': a[:, 1], 'pv_forecast_kw': a[:, 2], 'source_excel_row': np.arange(2, 146),
    })
    for c in ('load', 'pv_forecast'):
        frame[c + '_kwh_approx'] = frame[c + '_kw'] / 6
    frame['net_load_kw'] = frame['load_kw'] - frame['pv_forecast_kw']
    return frame


def build_hourly(books):
    rows = books[3]['Sheet1']
    require(len(rows) == 1461 and all(len(r) == 26 for r in rows), 'Unexpected forecast dimensions')
    require(list(rows[0][2:]) == [f'预报{i}小时' for i in range(1, 25)], 'Unexpected forecast lead columns')
    records = []
    current_date = None
    filled = 0
    for excel_row, row in enumerate(rows[1:], 2):
        if row[0] not in (None, ''):
            current_date = pd.Timestamp(row[0])
        else:
            require(current_date is not None, 'Forecast date missing before any group date')
            filled += 1
        clock = clock_minutes(row[1])
        require(clock in (0, 360, 720, 1080), 'Unexpected forecast release hour')
        issue = current_date + pd.Timedelta(minutes=clock)
        values = numeric_matrix([row[2:]], 'forecast')[0]
        for lead, value in enumerate(values, 1):
            target = issue + pd.Timedelta(hours=lead)
            records.append((issue, target, lead, value, excel_row, lead + 2,
                            YEAR_START < target <= YEAR_END))
    frame = pd.DataFrame(records, columns=['issue_time', 'target_time', 'lead_hours', 'pv_forecast_kw',
                                          'source_excel_row', 'source_excel_column', 'target_in_actual_coverage'])
    require(not frame.duplicated(['issue_time', 'target_time']).any(), 'Duplicate forecast vintage/target')
    issues = pd.DatetimeIndex(frame.issue_time.drop_duplicates())
    require(issues.equals(pd.date_range('2025-01-01', '2025-12-31 18:00', freq='6h')), 'Missing/misordered forecast issues')
    return frame, filled


def build_forecast10(hourly, actuals):
    """Within-vintage linear interpolation. The initial anchor may use observed PV at issue time.

    Measurements are assumed instantly available at their timestamp. If unavailable,
    use the first issued hourly forecast as a flat first-hour anchor; never backfill
    from a later actual observation or combine forecasts issued at different times.
    """
    observed = actuals.set_index('sample_time').pv_kw
    records = []
    for issue, g in hourly.groupby('issue_time', sort=True):
        g = g.sort_values('lead_hours')
        if issue in observed.index:
            anchor, kind, available = float(observed.loc[issue]), 'observed_at_issue_assumed_available', issue
        else:
            anchor, kind, available = float(g.pv_forecast_kw.iloc[0]), 'first_hour_forecast_flat_fallback', issue
        y = np.r_[anchor, g.pv_forecast_kw.to_numpy()]
        lead = np.arange(10, 1441, 10)
        values = np.interp(lead, np.arange(0, 1441, 60), y)
        for offset, value in zip(lead, values):
            end = issue + pd.Timedelta(minutes=int(offset))
            records.append((issue, end - STEP, end, int(offset), value, value / 6,
                            kind, anchor, available, YEAR_START < end <= YEAR_END))
    return pd.DataFrame(records, columns=['issue_time', 'interval_start', 'interval_end', 'lead_minutes',
                                         'pv_forecast_kw_interpolated', 'pv_forecast_kwh_approx',
                                         'initial_anchor_kind', 'initial_anchor_kw', 'anchor_available_at',
                                         'target_in_actual_coverage'])


def build_features(actuals, q1):
    """Features for a plan made at 00:00; no target-day actuals, no full-year fitting."""
    base = actuals[['source_date', 'slot', 'interval_start', 'interval_end']].copy()
    base = base.rename(columns={'source_date': 'origin_time'})
    base['day_of_week'] = base.origin_time.dt.dayofweek
    base['month'] = base.origin_time.dt.month
    base['day_of_year'] = base.origin_time.dt.dayofyear
    base['is_weekend'] = base.day_of_week >= 5
    angle = 2 * np.pi * (base.slot - 1) / 144
    base['time_sin'] = np.sin(angle)
    base['time_cos'] = np.cos(angle)
    base['price_q2_q3_known_yuan_per_kwh'] = np.tile(q1.price_yuan_per_kwh.to_numpy(), 365)
    variables = [('load_kw', 'load'), ('pv_kw', 'pv'), ('price_actual_yuan_per_kwh', 'price')]
    for column, prefix in variables:
        wide = actuals.pivot(index='source_date', columns='slot', values=column).sort_index()
        for lag in (1, 7):
            base[f'{prefix}_lag{lag}d'] = wide.shift(lag).to_numpy().ravel()
        for days in (7, 28):
            base[f'{prefix}_past{days}d_mean'] = wide.shift(1).rolling(days, min_periods=days).mean().to_numpy().ravel()
    base['latest_history_sample_time'] = base.interval_end - pd.Timedelta(days=1)
    base['earliest_history_sample_time'] = base.interval_end - pd.Timedelta(days=28)
    base = base[base.origin_time >= pd.Timestamp('2025-02-01')].reset_index(drop=True)
    require(not base.isna().any().any(), 'Unexpected missing feature value in evaluation period')
    require((base.latest_history_sample_time <= base.origin_time).all(), 'History newer than plan origin')
    return base


def source_stats(frame):
    rows = []
    for col in ('load_kw', 'pv_kw', 'price_actual_yuan_per_kwh', 'net_load_kw'):
        a = frame[col]
        rows.append(dict(variable=col, count=len(a), missing=int(a.isna().sum()),
                         zeros=int((a == 0).sum()), negative=int((a < 0).sum()),
                         minimum=float(a.min()), p01=float(a.quantile(.01)),
                         median=float(a.median()), p99=float(a.quantile(.99)), maximum=float(a.max())))
    return pd.DataFrame(rows)


def review_flags(actuals):
    """Descriptive full-year IQR flags, NOT preprocessing masks or model inputs."""
    chunks = []
    for col in ('load_kw', 'pv_kw', 'price_actual_yuan_per_kwh'):
        x = actuals[col]
        q1, q3 = float(x.quantile(.25)), float(x.quantile(.75))
        low, high = q1 - 3 * (q3 - q1), q3 + 3 * (q3 - q1)
        mask = (x < low) | (x > high)
        item = actuals.loc[mask, ['sample_time', 'source_excel_row', 'source_excel_column']].copy()
        item['variable'], item['value'] = col, x[mask]
        item['reason'], item['action'] = 'outside_global_3IQR_descriptive_only', 'retained_unchanged'
        chunks.append(item)
    return pd.concat(chunks, ignore_index=True)


def template_map(template_dir):
    records = []
    for path in sorted(template_dir.glob('*.xlsx')):
        book = read_book(path)
        for name, rows in book.items():
            if name not in ('计划购电量', '调整购电量'):
                continue
            if path.name == 'result1.xlsx':
                labels = [r[0] for r in rows[1:145]]
                coords = [f'B{i}' for i in range(2, 146)]
            else:
                labels = list(rows[0][1:145])
                coords = [f'{openpyxl.utils.get_column_letter(i)}1' for i in range(2, 146)]
            require(len(labels) == 144, 'Unexpected template interval count')
            for slot, (label, coord) in enumerate(zip(labels, coords), 1):
                start, end = (slot - 1) * 10, slot * 10
                def fmt(m):
                    return f'{m // 60:02d}:{m % 60:02d}'
                records.append((path.name, name, slot, coord, label, f'{fmt(start)}-{fmt(end)}',
                                'UNRESOLVED_DO_NOT_EXPORT_BY_POSITION'))
    return pd.DataFrame(records, columns=['template_file', 'sheet', 'slot_by_position', 'template_cell',
                                         'original_interval_label', 'provisional_physical_interval', 'status'])


def validate(actuals, q1, hourly, forecast10, features):
    checks = []
    def check(name, condition, detail):
        ok = bool(condition)
        checks.append({'check': name, 'passed': ok, 'detail': detail})
        require(ok, name + ': ' + detail)
    check('actual_count', len(actuals) == 52560, '365 days x 144 records')
    check('unique_actual_time', actuals.sample_time.is_unique, 'No duplicate sample timestamps')
    check('continuous_actual_time', actuals.sample_time.diff().dropna().eq(STEP).all(), '10-minute spacing including day boundaries')
    check('full_calendar_intervals', actuals.interval_start.iloc[0] == YEAR_START and actuals.interval_end.iloc[-1] == YEAR_END, 'Under provisional end-sample mapping')
    check('forecast_hourly_count', len(hourly) == 35040, '1460 vintages x 24 leads')
    check('forecast10_count', len(forecast10) == 210240, '1460 vintages x 144 leads')
    check('unique_forecast10_key', not forecast10.duplicated(['issue_time', 'interval_end']).any(), 'One value per vintage and target')
    check('causal_forecast_anchors', (forecast10.anchor_available_at <= forecast10.issue_time).all(), 'No actual anchor after publication')
    check('forecast_horizon', (forecast10.interval_start >= forecast10.issue_time).all(), 'Only future intervals per vintage')
    join = forecast10[forecast10.lead_minutes % 60 == 0].merge(hourly, left_on=['issue_time', 'interval_end'], right_on=['issue_time', 'target_time'], validate='one_to_one')
    check('original_hourly_knots_preserved', len(join) == len(hourly) and np.array_equal(join.pv_forecast_kw_interpolated, join.pv_forecast_kw), 'Interpolated values at whole hours equal source exactly')
    check('nonnegative_interpolated_pv', (forecast10.pv_forecast_kw_interpolated >= 0).all(), 'No overshoot from linear interpolation')
    check('features_count', len(features) == 48096, '334 evaluation days x 144 slots')
    check('features_no_missing', not features.isna().any().any(), '28-day history exists from February')
    check('features_causal', (features.latest_history_sample_time <= features.origin_time).all(), 'All historical features available by 00:00 under instant-observation assumption')
    check('no_actual_target_columns_in_features', not set(['load_kw', 'pv_kw', 'net_load_kw', 'price_actual_yuan_per_kwh']) & set(features.columns), 'Target-day actuals are separate from features')
    for col in ('load', 'pv', 'net_load'):
        check('energy_conversion_' + col, np.allclose(actuals[col + '_kwh_approx'] * 6, actuals[col + '_kw'], atol=1e-9, rtol=0), 'Power x 1/6 hour, no rounded intermediate energies')
    return checks


def data_dictionary(tables):
    definitions = {
        'source_date': '原工作表行日期；次日00:00采样仍归该来源日',
        'slot': '原数据日内顺序1..144，不是模板已确认的时间映射',
        'sample_time': '原始采样时间，0:00+1转为次日00:00',
        'interval_start': '暂按采样值代表前10分钟平均功率构造的区间起点',
        'interval_end': '暂定区间终点，等于采样时刻/预测目标时刻',
        'origin_time': '日前计划作出时刻，当天00:00',
        'issue_time': '该版本预报的发布时间',
        'target_time': '该预报对应的未来整点',
        'latest_history_sample_time': '该行历史特征用到的最新实际采样时刻',
        'earliest_history_sample_time': '28天同期均值所用最早实际采样时刻',
        'source_excel_row': '来源Excel行号（1开始）',
        'source_excel_column': '来源Excel列号（1开始）；实际负载/光伏/电价三表坐标一致',
        'observed_available_at_assumed': '暂假设测量无延迟，在采样时刻已可用于决策',
        'anchor_available_at': '插值起始锚点的可用时刻',
        'initial_anchor_kind': '插值时刻0锚点为当前已观测光伏或首小时预报回退值',
        'target_in_actual_coverage': '目标是否位于实际观测覆盖范围；false不是预报无效',
        'price_q2_q3_known_yuan_per_kwh': '仅供问题2/3的已知固定日电价，来自附件1；不可当作问题4已知实际电价',
        'period_role': '1月initial_history，2—12月evaluation；不是可随机切分的训练测试标签',
        'sample_offset_minutes': '采样相对来源日00:00的分钟数，10..1440',
        'load_kw': '负载功率：q1表来自附件1；实际及标签表来自附件2',
        'pv_kw': '附件2中的实际光伏功率，保留全部零值与峰值',
        'price_actual_yuan_per_kwh': '附件4实际电价；未来是否提前可知未明确',
        'net_load_kw': '负载功率减光伏功率，负值表示光伏剩余',
        'pv_exceeds_load': '光伏功率严格大于负载功率',
        'source_time_label': '读取后保留的原始时钟标签，混合Excel时间类型和字符串',
        'interval_start_minute': '暂定区间起点相对典型日00:00的分钟数',
        'interval_end_minute': '暂定区间终点相对典型日00:00的分钟数',
        'price_yuan_per_kwh': '附件1典型日给定电价',
        'pv_forecast_kw': '原始光伏预报功率：q1来自附件1；带发布时间的表来自附件3',
        'lead_hours': '目标整点距离预报发布时间的小时数，1..24',
        'lead_minutes': '目标区间终点距离预报发布时间的分钟数，10..1440',
        'pv_forecast_kw_interpolated': '同一发布版本内分段线性插值的光伏功率',
        'initial_anchor_kw': '该发布版本插值在0小时时刻的锚点功率',
        'day_of_week': '星期编号，周一0至周日6；不等同于法定工作日',
        'month': '日历月份1..12', 'day_of_year': '日历年内序号1..365',
        'is_weekend': '是否为周六或周日；未导入法定节假日或调休信息',
        'time_sin': 'sin(2*pi*(slot-1)/144)，区间起点的日内周期编码',
        'time_cos': 'cos(2*pi*(slot-1)/144)，区间起点的日内周期编码',
        'date': '来源日期，用于逐日回测',
        'role': 'initial_history为1月；evaluation为2—12月',
        'paper_report_date': '是否为题面指定的3月20日、6月21日、9月23日、12月21日',
    }
    records = []
    for table, frame in tables.items():
        for col in frame:
            unit = 'kWh (approx)' if 'kwh_approx' in col else ('yuan/kWh' if 'price' in col else ('kW' if col.endswith('_kw') or col.startswith(('load_lag', 'load_past', 'pv_lag', 'pv_past')) or col in ('initial_anchor_kw', 'pv_forecast_kw_interpolated') else ''))
            description = definitions.get(col, '')
            if '_lag' in col:
                description = '与预测目标同一日内时点、前1或7天的实际值；仅历史'
            elif '_past' in col:
                description = '截至前一天的7或28天同一时点均值；不包含当天'
            elif 'kwh_approx' in col:
                description = '对应功率乘1/6小时；依赖区间平均功率近似，不是题面额外给定的电量'
            records.append((table, col, str(frame[col].dtype), unit, description))
    return pd.DataFrame(records, columns=['table', 'column', 'dtype', 'unit', 'description'])


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=ROOT)
    args = parser.parse_args()
    root = args.root.resolve()
    out, diag = root / 'processed', root / 'diagnostics'
    out.mkdir(exist_ok=True); diag.mkdir(exist_ok=True)
    rawpaths = [root / 'raw' / f'附件{i}.xlsx' for i in range(1, 5)]
    paths = rawpaths + sorted((root / 'templates').glob('*.xlsx'))
    before = {str(p.relative_to(root)): sha256(p) for p in paths}
    books = {i: read_book(p) for i, p in enumerate(rawpaths, 1)}
    actuals, q1 = build_actuals(books), build_q1(books)
    hourly, filled = build_hourly(books)
    forecast10 = build_forecast10(hourly, actuals)
    features = build_features(actuals, q1)
    labels = actuals.loc[actuals.period_role == 'evaluation', ['source_date', 'slot', 'interval_end', 'load_kw', 'pv_kw', 'net_load_kw', 'price_actual_yuan_per_kwh']].copy()
    calendar = pd.DataFrame({'date': pd.date_range('2025-01-01', '2025-12-31')})
    calendar['role'] = np.where(calendar.date < pd.Timestamp('2025-02-01'), 'initial_history', 'evaluation')
    calendar['paper_report_date'] = calendar.date.isin(pd.to_datetime(['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']))
    tables = {'actuals_10min': actuals, 'q1_typical_day': q1, 'pv_forecast_hourly': hourly,
              'pv_forecast_10min': forecast10, 'day_ahead_features': features,
              'evaluation_labels_DO_NOT_USE_AS_FEATURES': labels, 'evaluation_calendar': calendar}
    for name, frame in tables.items():
        write_csv(frame, out / (name + '.csv'))
    write_csv(data_dictionary(tables), root / 'data_dictionary.csv')
    checks = validate(actuals, q1, hourly, forecast10, features)
    stats = source_stats(actuals)
    write_csv(stats, diag / 'numeric_summary.csv')
    flags = review_flags(actuals)
    write_csv(flags, diag / 'review_flags_NOT_FOR_TRAINING.csv')
    tmap = template_map(root / 'templates')
    write_csv(tmap, diag / 'template_time_mapping_UNRESOLVED.csv')
    daily = actuals.groupby('source_date').agg(load_kwh_approx=('load_kwh_approx', 'sum'), pv_kwh_approx=('pv_kwh_approx', 'sum'),
                    price_min=('price_actual_yuan_per_kwh', 'min'), price_max=('price_actual_yuan_per_kwh', 'max'),
                    surplus_slots=('pv_exceeds_load', 'sum')).reset_index()
    write_csv(daily, diag / 'daily_summary_DESCRIPTIVE.csv')
    annual_means = actuals.groupby('slot')[['load_kw', 'pv_kw', 'price_actual_yuan_per_kwh']].mean()
    differences = {'load_kw': float(np.max(abs(q1.load_kw.to_numpy() - annual_means.load_kw.to_numpy()))),
                   'pv_kw': float(np.max(abs(q1.pv_forecast_kw.to_numpy() - annual_means.pv_kw.to_numpy()))),
                   'price_yuan_per_kwh': float(np.max(abs(q1.price_yuan_per_kwh.to_numpy() - annual_means.price_actual_yuan_per_kwh.to_numpy())))}
    after = {str(p.relative_to(root)): sha256(p) for p in paths}
    require(before == after, 'Raw input/template mutated')
    checks.append({'check': 'source_files_unchanged', 'passed': True, 'detail': 'SHA256 unchanged for all input workbooks'})
    manifest = {'schema_version': '1.0', 'source_sha256': before,
                'source_column_mapping': {
                    'actuals_10min.load_kw': 'raw/附件2.xlsx / 小区负载 / source_excel_row,source_excel_column',
                    'actuals_10min.pv_kw': 'raw/附件2.xlsx / 光伏发电实际功率 / source_excel_row,source_excel_column',
                    'actuals_10min.price_actual_yuan_per_kwh': 'raw/附件4.xlsx / Sheet1 / source_excel_row,source_excel_column',
                    'q1_typical_day': 'raw/附件1.xlsx / Sheet1 / A时间 B电价 C负载 D光伏预报',
                    'pv_forecast_hourly.pv_forecast_kw': 'raw/附件3.xlsx / Sheet1 / source_excel_row,source_excel_column'},
                'output_sha256': {str(p.relative_to(root)): sha256(p) for p in sorted(out.glob('*.csv'))},
                'row_counts': {n: len(f) for n, f in tables.items()},
                'assumptions': {
                    'time_axis': 'Unzoned local clock of source files; no UTC or user-location conversion',
                    'interval_alignment': 'PROVISIONAL: sample at t mapped to [t-10min,t)',
                    'energy_approximation': 'Power treated as preceding interval average; kWh = kW/6',
                    'instant_observation_availability': True,
                    'forecast_interpolation': 'Linear within each issued vintage; anchor actual at issue if present, else flat first hourly forecast',
                    'preserve_outliers_zeros_negative_net_load': True,
                    'scaling_imputation_model_fitting': 'None',
                    'forecast_tail': 'All forecasts preserved, including targets after 2025 actual coverage',
                    'price_information': 'Q2/3 fixed daily prices known from attachment1. Q4 actual prices are outcomes/history; future visibility unspecified.',
                    'template_time_mapping': 'UNRESOLVED; original templates unchanged; positional export prohibited'},
                'filled_group_dates': filled,
                'out_of_coverage_hourly_forecasts': int((~hourly.target_in_actual_coverage).sum()),
                'out_of_coverage_10min_forecasts': int((~forecast10.target_in_actual_coverage).sum()),
                'descriptive_review_flags_count': len(flags),
                'attachment1_max_deviation_from_full_year_slot_mean': differences,
                'runtime': {'python': platform.python_version(), 'pandas': pd.__version__, 'numpy': np.__version__, 'openpyxl': openpyxl.__version__}}
    write_json(root / 'manifest.json', manifest)
    write_json(diag / 'validation_results.json', {'status': 'passed', 'checks': checks})
    report = f'''# C题数据预处理检查报告

## 处理结果
- 实际数据：{len(actuals):,}行，144点/天，365天。
- 原始整点预报：{len(hourly):,}行，1460次独立发布。
- 派生10分钟预报：{len(forecast10):,}行，保留所有版本。
- 日前预测特征：{len(features):,}行，仅2—12月，未加入当天实际标签。
- 预报分组日期向下补全：{filled}个日期单元格；只在标准数据中补全。
- 数值缺失、非有限值、原始负值：均为0。重复实际时间和重复预报键：均为0。
- 全年3倍IQR描述性标记：{len(flags)}条，仅保留供查看，没有删除或替换。
- 跨出2025实际覆盖范围的预报：整点{manifest['out_of_coverage_hourly_forecasts']}行，10分钟{manifest['out_of_coverage_10min_forecasts']}行，全部保留并标记。
- 自动校验：{len(checks)}项通过；来源文件SHA256未变。

## 必须保留的限制
1. 时间标签歧义未解决。标准区间暂按采样时刻之前10分钟处理，电量列均为近似派生值。
2. 原模板未改。不得按数组位置直接把标准区间结果塞入模板。映射问题见template_time_mapping_UNRESOLVED.csv。
3. 即时观测可用性是建模假设。预报插值第一个小时使用发布时间的实际光伏作锚点；没有该时刻观测则使用首小时预报回退。未用后续实际数据。
4. 所有小时之间的预报值是线性插值，不是新增官方预报，也不代表预测准确性提高。
5. 附件1曲线接近全年同时间点平均，最大偏差：负载{differences['load_kw']:.8f}kW、光伏{differences['pv_kw']:.8f}kW、电价{differences['price_yuan_per_kwh']:.8f}元/kWh。此项是描述性核对，不能自行把用全年数据拟合的曲线作为历史可知预测。
6. 缺少天气、位置、电价提前公布规则，未杜撰补充数据、时区或未来电价信息。

## 未进行的处理
未删峰值；未把负净负载归零；未把负载/光伏截到5000kW；未对全年拟合标准化器；未随机打乱时间切分；未补入未来实际值；未生成优化答案。
'''
    (diag / '预处理检查报告.md').write_text(report, encoding='utf-8')
    print(json.dumps({'status': 'passed', 'row_counts': manifest['row_counts'], 'checks': len(checks), 'date_groups_filled': filled,
                      'hourly_tail': manifest['out_of_coverage_hourly_forecasts'], 'forecast10_tail': manifest['out_of_coverage_10min_forecasts'],
                      'review_flags': len(flags)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
