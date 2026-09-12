"""Causal forecast, continuation-value teacher, learning, dispatch and execution.

Energy at public interfaces is kWh; optimization internally uses MWh.
No realized future observations are used in continuation-value labels.
"""
from __future__ import annotations

import sys
sys.dont_write_bytecode = True
import hashlib
import json
import pickle
import time
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.optimize import linprog, minimize_scalar, nnls
from scipy.sparse import coo_matrix, csc_matrix, eye, vstack
from sklearn.decomposition import PCA
from sklearn.linear_model import Ridge
from sklearn.preprocessing import StandardScaler

ROOT = Path(__file__).resolve().parents[2]
TMP = ROOT / '.tmp/q4_terminal_value'
EMIN, EMAX, E0, CAP, ETA, POWER = 1200., 10800., 6000., 9600., .9, 5000.
XGRID = np.linspace(0., CAP, 7)
Q = .85
TOL = .003  # kWh; below one hundred-thousandth of a dispatch interval's capacity
VERSION = 'terminal-value-v1'
sys.path.insert(0, str(TMP / 'deps'))
try:
    import osqp
except ImportError:
    osqp = None
try:
    import clarabel
except ImportError:
    clarabel = None


def signature(paths=()):
    h = hashlib.sha256(VERSION.encode())
    for p in [Path(__file__), *paths]:
        h.update(Path(p).read_bytes())
    return h.hexdigest()


def save_pickle(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix('.partial')
    tmp.write_bytes(pickle.dumps(obj, protocol=5))
    tmp.replace(path)


def load_data():
    """Reuse Q2 forecast algorithm, never fit a price/value model on future data."""
    TMP.mkdir(parents=True, exist_ok=True)
    sys.path.insert(0, str(ROOT / 'submit'))
    import q2
    from openpyxl import load_workbook
    paths = [ROOT / 'submit/q2.py', *[
        ROOT / f'submit/Data/附件{i}.xlsx' for i in [1, 2, 4]]]
    key = signature(paths)
    path = TMP / 'data.pkl'
    if path.exists():
        cached = pickle.loads(path.read_bytes())
        if cached['key'] == key:
            return cached
    dates, load, pv, _ = q2.load_inputs(paths[2], paths[1])
    wb = load_workbook(paths[3], read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True)); wb.close()
    assert pd.DatetimeIndex(pd.to_datetime([r[0] for r in rows[1:]])).equals(dates)
    prices = np.asarray([r[1:145] for r in rows[1:]], float)
    assert prices.shape == load.shape == pv.shape == (365, 144)
    assert np.isfinite(prices).all() and (prices > 0).all()
    pred_l, pred_p, scenarios = q2.forecast_scenarios(dates, load, pv)
    nets = {d: np.quantile(a, Q, axis=0) / 6 for d, a in scenarios.items()}
    result = dict(key=key, dates=dates, load=load, pv=pv, prices=prices,
                  pred_load=pred_l, pred_pv=pred_p, nets=nets)
    save_pickle(path, result)
    return result


def historical_future(data, d):
    """At day d midnight, predict d+1..d+6 with realized d-6..d-1."""
    if d < 7:
        raise ValueError('At least seven completed days required')
    ref = np.arange(d-6, d)
    assert np.all(ref < d)
    return {k: np.asarray(data[k][ref], float).copy()
            for k in ('prices', 'load', 'pv')} | {'reference_indices': ref}


def features(future):
    p = future['prices'].reshape(6, 24, 6).mean(axis=2)
    net = (future['load']-future['pv']).reshape(6, 24, 6).mean(axis=2)
    econ, names = [], []
    for k in range(6):
        econ.extend([p[k].mean(), p[k].min(), p[k].max(),
                     np.maximum(net[k], 0).sum(), np.maximum(-net[k], 0).sum()])
        names.extend([f'day{k+1}_{s}' for s in ('mean_price', 'min_price',
                     'max_price', 'deficit_kwh', 'surplus_kwh')])
    econ.extend([p[0, :6].mean(), np.maximum(net[0, :6], 0).sum(), np.argmin(p[0])])
    names.extend(['tomorrow_00_06_price', 'tomorrow_00_06_deficit', 'next24_min_hour'])
    return dict(econ=np.array(econ), price=p.ravel(), net=net.ravel(), names=names)


@dataclass
class Reward:
    kind: str
    a: float = 0.
    b: float = 0.
    values: np.ndarray | None = None

    def value(self, x):
        if self.kind == 'fixed':
            return -.45 * abs(x + EMIN - E0)
        if self.kind == 'pwl':
            return float(np.interp(x, XGRID, self.values))
        return self.a*x - .5*self.b*x*x


def fit_rewards(values):
    z = XGRID / CAP
    # Nonnegative weights parameterize high-SOC marginal value and marginal gap.
    w, _ = nnls(np.column_stack([z, z-.5*z*z]), np.asarray(values)/CAP)
    quad = Reward('quadratic', a=float(w.sum()), b=float(w[1]/CAP))
    linear = Reward('linear', a=max(0., float(XGRID @ values / (XGRID @ XGRID))))
    err = np.array([quad.value(x) for x in XGRID]) - values
    return quad, linear, float(np.sqrt(np.mean(err**2)))


@lru_cache(maxsize=12)
def balance_matrix(n, dt):
    # variables: grid, charge, discharge, unused, E[0..n], epigraph
    m = 5*n+2
    t = np.arange(n)
    rr = np.r_[t,t,t,t,n+t,n+t,n+t,n+t]
    cc = np.r_[t,n+t,2*n+t,3*n+t,n+t,2*n+t,4*n+t,4*n+t+1]
    vv = np.r_[np.ones(n),-np.ones(n),np.ones(n),-np.ones(n),
               -ETA*np.ones(n),np.ones(n)/ETA,-np.ones(n),np.ones(n)]
    return coo_matrix((vv,(rr,cc)),shape=(2*n,m)).tocsc()


def problem_arrays(net, prices, initial, dt, reward, terminal):
    n = len(net); m = 5*n+2; ie = 5*n; it = m-1
    A = balance_matrix(n, dt)
    rhs = np.r_[net/1000, np.zeros(n)]
    lo, hi = np.zeros(m), np.full(m, np.inf)
    hi[n:3*n] = POWER*dt/1000
    lo[4*n:5*n+1], hi[4*n:5*n+1] = EMIN/1000, EMAX/1000
    lo[4*n] = hi[4*n] = initial/1000
    if terminal is not None:
        lo[ie] = hi[ie] = terminal/1000
    obj = np.zeros(m); obj[:n] = prices
    diag = np.zeros(m)
    if reward.kind in ('linear', 'quadratic'):
        diag[ie] = reward.b*1000
        obj[ie] = -reward.a-reward.b*EMIN
    ar, ac, av, ub = [], [], [], []
    if reward.kind == 'fixed':
        lo[it] = -np.inf; obj[it] = 1.
        for r, (s, c) in enumerate([(.45,-.45*E0/1000),(-.45,.45*E0/1000)]):
            ar.extend([r,r]); ac.extend([ie,it]); av.extend([s,-1]); ub.append(-c)
    elif reward.kind == 'pwl':
        lo[it] = -np.inf; obj[it] = 1.
        slopes = np.diff(reward.values)/np.diff(XGRID)
        if np.max(np.diff(slopes)) > 1e-6:
            raise ValueError('Nonconcave continuation values')
        for r, slope in enumerate(slopes):
            intercept = reward.values[r]-slope*XGRID[r]
            ar.extend([r,r]); ac.extend([ie,it]); av.extend([-slope,-1])
            ub.append((intercept-slope*EMIN)/1000)
    else:
        hi[it] = 0.
    U = coo_matrix((av,(ar,ac)),shape=(len(ub),m)).tocsc()
    return A, rhs, U, np.array(ub), lo, hi, obj, diag


def clean_plan(v, net, prices, dt, reward, solver, seconds):
    n = len(net)
    g,c,d,u = [v[j*n:(j+1)*n].copy()*1000 for j in range(4)]
    e = v[4*n:5*n+1].copy()*1000
    for a in (g,c,d,u):
        if a.min() < -TOL:
            raise ValueError('Negative primal decision')
        a[a < 0] = 0
    # Cancel a zero-state-change charge/discharge cycle, divert saved energy to unused.
    overlap = np.minimum(c, d/ETA**2)
    c -= overlap; d -= ETA**2*overlap; u += (1-ETA**2)*overlap
    balance = g+d-c-u-net
    recursion = np.diff(e)-ETA*c+d/ETA
    if max(np.abs(balance).max(), np.abs(recursion).max()) > TOL:
        raise ValueError('Primal feasibility tolerance exceeded')
    if e.min() < EMIN-TOL or e.max() > EMAX+TOL:
        raise ValueError('Storage bound violated')
    if max(c.max(),d.max()) > POWER*dt+TOL:
        raise ValueError('Power bound violated')
    return dict(grid=g, charge=c, discharge=d, unused=u, storage=e, solver=solver,
                seconds=seconds, objective=float(prices@g-reward.value(e[-1]-EMIN)),
                balance_error=float(np.abs(balance).max()),
                recursion_error=float(np.abs(recursion).max()))


def dispatch(net, prices, initial=E0, dt=1/6, reward=None, terminal=None,
             force_fallback=False):
    start = time.perf_counter()
    reward = reward or Reward('none')
    net, prices = np.asarray(net,float), np.asarray(prices,float)
    assert net.ndim == 1 and net.shape == prices.shape and np.isfinite(net).all()
    assert np.isfinite(prices).all() and (prices > 0).all()
    if not np.isfinite(initial) or initial < EMIN-TOL or initial > EMAX+TOL:
        raise ValueError('Initial storage is outside physical limits')
    initial = float(np.clip(initial, EMIN, EMAX))
    A,rhs,U,ub,lo,hi,obj,diag = problem_arrays(net,prices,initial,dt,reward,terminal)
    quadratic = diag.max() > 1e-12
    reason = None
    if quadratic and clarabel is not None and not force_fallback:
        # Interior-point QP handles flat LP faces much better than ADMM here.
        ident=eye(len(obj),format='csc')
        upper=np.flatnonzero(np.isfinite(hi)); lower=np.flatnonzero(np.isfinite(lo))
        AA=vstack([A,U,ident[upper,:],-ident[lower,:]],format='csc')
        bb=np.r_[rhs,ub,hi[upper],-lo[lower]]
        P=csc_matrix((diag,(np.arange(len(obj)),np.arange(len(obj)))),shape=(len(obj),len(obj)))
        settings=clarabel.DefaultSettings()
        settings.verbose=False
        settings.tol_gap_abs=1e-9;settings.tol_gap_rel=1e-9;settings.tol_feas=1e-9
        settings.max_iter=200
        cones=[clarabel.ZeroConeT(len(rhs)),clarabel.NonnegativeConeT(len(bb)-len(rhs))]
        solved=clarabel.DefaultSolver(P,obj,AA,bb,cones,settings).solve()
        if str(solved.status) in ('Solved','AlmostSolved'):
            try:
                return clean_plan(np.asarray(solved.x),net,prices,dt,reward,'Clarabel',time.perf_counter()-start)
            except ValueError as exc:
                reason=str(exc)
        else:
            reason=str(solved.status)
    if quadratic and osqp is not None and not force_fallback:
        qp = osqp.OSQP()
        AA = vstack([A,U,eye(len(obj),format='csc')],format='csc')
        qp.setup(P=csc_matrix((diag,(np.arange(len(obj)),np.arange(len(obj)))),shape=(len(obj),len(obj))),
                 q=obj,A=AA,l=np.r_[rhs,np.full(len(ub),-np.inf),lo],u=np.r_[rhs,ub,hi],
                 verbose=False,eps_abs=1e-8,eps_rel=1e-8,max_iter=5000,
                 polishing=True,check_termination=25)
        res = qp.solve(raise_error=False)
        if res.info.status_val in (1,2):
            try:
                return clean_plan(res.x,net,prices,dt,reward,'OSQP',time.perf_counter()-start)
            except ValueError as exc:
                reason = str(exc)
        else:
            reason = res.info.status
    if quadratic:
        # Same convex objective, LP evaluation of daily cost conditional on end state.
        cache = {}
        def value(end):
            end = float(end)
            if end not in cache:
                try:
                    plan = dispatch(net,prices,initial,dt,Reward('none'),end)
                    plan['objective'] = float(prices@plan['grid']-reward.value(end-EMIN))
                    cache[end] = plan
                except RuntimeError:
                    return np.inf
            return cache[end]['objective']
        lower = max(EMIN, initial-POWER*dt*len(net)/ETA)
        upper = min(EMAX, initial+ETA*POWER*dt*len(net))
        if terminal is not None:
            value(terminal)
        else:
            opt = minimize_scalar(value,bounds=(lower,upper),method='bounded',
                                  options={'xatol':1e-4,'maxiter':80})
            value(lower); value(upper); value(opt.x)
        if not cache:
            raise RuntimeError('No feasible fallback terminal state')
        best = min(cache.values(),key=lambda p:p['objective']).copy()
        best.update(solver='LP_scalar_fallback',seconds=time.perf_counter()-start,
                    fallback_reason=reason or ('forced' if force_fallback else 'OSQP unavailable'))
        return best
    res = linprog(obj,A_eq=A,b_eq=rhs,A_ub=U if len(ub) else None,
                  b_ub=ub if len(ub) else None,bounds=np.column_stack([lo,hi]),method='highs')
    if not res.success:
        raise RuntimeError(res.message)
    return clean_plan(res.x,net,prices,dt,reward,'LP',time.perf_counter()-start)


def teacher(future, terminal=E0, hourly=True):
    start = time.perf_counter()
    net = future['load']-future['pv']
    if hourly:
        net = net.reshape(6,24,6).mean(axis=2).ravel()
        p = future['prices'].reshape(6,24,6).mean(axis=2).ravel(); dt=1.
    else:
        net = net.ravel()/6; p=future['prices'].ravel(); dt=1/6
    costs = np.array([dispatch(net,p,EMIN+x,dt,terminal=terminal)['objective'] for x in XGRID])
    values = costs[0]-costs
    slopes = np.diff(values)/np.diff(XGRID)
    assert values.min() >= -1e-4 and slopes.min() >= -1e-7
    assert np.max(np.diff(slopes)) <= 1e-6
    quad,lin,rmse = fit_rewards(values)
    return dict(costs=costs,values=values,a=quad.a,b=quad.b,linear=lin.a,
                marginal_low=quad.a,marginal_high=quad.a-quad.b*CAP,
                fit_rmse_yuan=rmse,seconds=time.perf_counter()-start)


def build_labels(data, progress=print):
    path = TMP/'labels.pkl'
    saved = pickle.loads(path.read_bytes()) if path.exists() else {}
    rows = saved.get('rows',{}) if saved.get('key') == data['key'] else {}
    for d in range(7,len(data['dates'])):
        if d not in rows:
            future = historical_future(data,d)
            rows[d] = teacher(future) | features(future) | dict(
                index=d,date=str(data['dates'][d].date()),
                information_cutoff=str(data['dates'][d-1].date()),
                references=[str(data['dates'][i].date()) for i in future['reference_indices']],
                forecast_config='weekly_lag7_load_pv_price',
                teacher_config='six_days_hourly_terminal6000_grid7')
        if d % 30 == 0:
            save_pickle(path,dict(key=data['key'],rows=rows))
            progress(f'价值标签：{len(rows)}/358 天')
    save_pickle(path,dict(key=data['key'],rows=rows))
    serial = [{k:v for k,v in r.items() if np.isscalar(v)} for r in rows.values()]
    pd.DataFrame(serial).to_csv(TMP/'labels.csv',index=False)
    return rows


class ValueModel:
    def __init__(self, kind, feature, alpha=1., components=3):
        self.kind,self.feature,self.alpha,self.components=kind,feature,alpha,components

    def transform(self, rows, fit=False):
        parts=[]
        if self.feature in ('economic','hybrid'):
            parts.append(np.stack([r['econ'] for r in rows]))
        if self.feature in ('pca','hybrid'):
            if fit:
                self.pcas={k:PCA(n_components=self.components,svd_solver='full') for k in ('price','net')}
            for k in ('price','net'):
                raw=np.stack([r[k] for r in rows])
                parts.append(self.pcas[k].fit_transform(raw) if fit else self.pcas[k].transform(raw))
        X=np.column_stack(parts)
        if fit:
            self.scale=StandardScaler().fit(X)
        return self.scale.transform(X)

    def fit(self, rows):
        self.train_indices=tuple(r['index'] for r in rows)
        y=np.array([[r['marginal_low'],r['marginal_high']] if self.kind=='quadratic'
                    else [r['linear']] for r in rows])
        self.reg=Ridge(alpha=self.alpha).fit(self.transform(rows,True),y)
        return self

    def predict(self, rows):
        assert min(r['index'] for r in rows) > max(self.train_indices)
        y=np.asarray(self.reg.predict(self.transform(rows))).reshape(len(rows),-1)
        rewards=[]
        for vals in y:
            if self.kind=='linear':
                rewards.append(Reward('linear',max(0.,float(vals[0]))))
            else:
                low,high=map(float,vals)
                if low < high:
                    low=high=(low+high)/2
                low,high=max(0.,low),max(0.,high)
                rewards.append(Reward('quadratic',low,(low-high)/CAP))
        return rewards


def execute(grid, actual_net, initial, dt=1/6):
    n=len(grid); e=np.empty(n+1); e[0]=initial
    c=np.zeros(n); d=np.zeros(n); h=np.zeros(n); u=np.zeros(n)
    for t in range(n):
        surplus=grid[t]-actual_net[t]
        if surplus >= 0:
            c[t]=min(surplus,POWER*dt,max(0.,(EMAX-e[t])/ETA))
            u[t]=surplus-c[t]
        else:
            d[t]=min(-surplus,POWER*dt,max(0.,(e[t]-EMIN)*ETA))
            h[t]=-surplus-d[t]
        e[t+1]=e[t]+ETA*c[t]-d[t]/ETA
    return dict(charge=c,discharge=d,emergency=h,unused=u,storage=e)


def backtest(data, indices, rewards, keep_slots=False):
    daily=[]; slots=[]; state=E0
    for d,reward in zip(indices,rewards,strict=True):
        p=data['prices'][d-7]
        plan=dispatch(data['nets'][d],p,state,reward=reward)
        actual_net=(data['load'][d]-data['pv'][d])/6
        ex=execute(plan['grid'],actual_net,state)
        price=data['prices'][d]
        planned=float(price@plan['grid']); emergency=float(5*price@ex['emergency'])
        e=ex['storage']
        balance=plan['grid']+ex['discharge']+ex['emergency']-actual_net-ex['charge']-ex['unused']
        recursion=np.diff(e)-ETA*ex['charge']+ex['discharge']/ETA
        assert np.abs(balance).max()<TOL and np.abs(recursion).max()<TOL
        assert e.min()>=EMIN-TOL and e.max()<=EMAX+TOL
        assert not ((ex['charge']>TOL)&(ex['discharge']>TOL)).any()
        assert max(ex['charge'].max(),ex['discharge'].max())<=POWER/6+TOL
        row=dict(date=str(data['dates'][d].date()),index=d,initial_storage=state,
                 end_storage=float(e[-1]),planned_end=float(plan['storage'][-1]),
                 planned_cost=planned,emergency_cost=emergency,total_cost=planned+emergency,
                 planned_kwh=float(plan['grid'].sum()),emergency_kwh=float(ex['emergency'].sum()),
                 unused_kwh=float(ex['unused'].sum()),reward_kind=reward.kind,
                 a=reward.a,b=reward.b,terminal_gap=float(e[-1]-plan['storage'][-1]),
                 boundary_slots=int(((e[1:]<=EMIN+TOL)|(e[1:]>=EMAX-TOL)).sum()),
                 solver=plan['solver'],solve_seconds=plan['seconds'],
                 fallback_reason=plan.get('fallback_reason',''),
                 balance_error=float(np.abs(balance).max()),
                 recursion_error=float(np.abs(recursion).max()),
                 plan_balance_error=plan['balance_error'],plan_recursion_error=plan['recursion_error'])
        daily.append(row)
        if keep_slots:
            slots.append(pd.DataFrame(dict(date=row['date'],slot=np.arange(1,145),
                forecast_price=p,actual_price=price,actual_net_kwh=actual_net,
                planned_grid_kwh=plan['grid'],planned_charge=plan['charge'],planned_discharge=plan['discharge'],
                planned_storage_start=plan['storage'][:-1],planned_storage_end=plan['storage'][1:],
                charge=ex['charge'],discharge=ex['discharge'],emergency_kwh=ex['emergency'],unused=ex['unused'],
                storage_start=e[:-1],storage_end=e[1:],planned_cost=price*plan['grid'],
                emergency_cost=5*price*ex['emergency'])))
        state=float(e[-1])
    return pd.DataFrame(daily),pd.concat(slots,ignore_index=True) if keep_slots else None


def summary(daily):
    cost=float(daily.total_cost.sum()); change=float(daily.end_storage.iloc[-1]-daily.initial_storage.iloc[0])
    return dict(days=len(daily),cash_cost=cost,planned_cost=float(daily.planned_cost.sum()),
                emergency_cost=float(daily.emergency_cost.sum()),emergency_kwh=float(daily.emergency_kwh.sum()),
                initial_storage=float(daily.initial_storage.iloc[0]),end_storage=float(daily.end_storage.iloc[-1]),
                inventory_adjusted_045=cost-.45*change,inventory_adjusted_090=cost-.9*change,
                mean_abs_terminal_gap=float(daily.terminal_gap.abs().mean()),
                boundary_fraction=float(daily.boundary_slots.sum()/(144*len(daily))),
                solve_seconds=float(daily.solve_seconds.sum()),
                fallbacks=int((daily.solver=='LP_scalar_fallback').sum()),
                max_balance_error=float(daily.balance_error.max()),
                max_recursion_error=float(daily.recursion_error.max()),
                max_plan_balance_error=float(daily.plan_balance_error.max()),
                max_plan_recursion_error=float(daily.plan_recursion_error.max()))
