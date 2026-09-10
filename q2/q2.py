"""
CUMCM 2026 C题 Q2：单文件最终版

默认运行：
    python q2.py

作用：
1. 读取 q2/input/q2_actual_load_pv.csv 和 q2/input/q2_fixed_prices.csv；
2. 每天 00:00 仅使用此前已经结束的历史日数据构造情景；
3. 使用 28 日历史情景 SAA + 储能作为最终策略；
4. 按因果实时规则执行储能与紧急购电；
5. 完成物理约束、费用、历史信息泄漏等检查；
6. 直接填写附件5中的官方 result2.xlsx；
7. 保存后重新读取 Excel 进行核验。

默认只生成：
    q2/results/result2.xlsx

不生成图片，不生成中间 CSV/JSON/Markdown，不依赖 parameters.json、
q2_core.py、run_q2.py、check_q2.py 或 JavaScript 导出脚本。

可选：
    python q2.py --compare
仅额外在终端比较 7日均值、28日SAA+储能、28日SAA无储能，
仍不生成中间文件。
"""

from __future__ import annotations

from copy import copy
from pathlib import Path
import argparse
import math
import shutil
import time

import numpy as np
import pandas as pd
from scipy.optimize import Bounds, LinearConstraint, linprog, milp
from scipy.sparse import coo_matrix, csr_matrix, hstack

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise ImportError(
        "需要 openpyxl 才能填写 result2.xlsx。请运行: pip install openpyxl"
    ) from exc


# ============================================================
# 路径
# ============================================================

Q2_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = Q2_DIR.parent

INPUT_DIR = Q2_DIR / "input"
DEFAULT_ACTUAL = INPUT_DIR / "q2_actual_load_pv.csv"
DEFAULT_PRICE = INPUT_DIR / "q2_fixed_prices.csv"
DEFAULT_OUTPUT = Q2_DIR / "results" / "result2.xlsx"


def default_template_path() -> Path:
    """优先使用整理后的 templates；没有时直接读取附件5中的官方原模板。"""
    candidates = [
        PROJECT_ROOT / "Data_preprocessed" / "templates" / "result2.xlsx",
        PROJECT_ROOT / "Data_preprocessed" / "raw" / "附件5" / "result2.xlsx",
    ]
    for path in candidates:
        if path.exists():
            return path
    return candidates[0]


# ============================================================
# 固定模型参数
# ============================================================

CFG = {
    "interval_minutes": 10,
    "charge_efficiency": 0.9,
    "discharge_efficiency": 0.9,
    "storage_min_kwh": 1200.0,
    "storage_max_kwh": 10800.0,
    "initial_storage_kwh": 6000.0,
    "planned_terminal_kwh": 6000.0,
    "max_charge_power_kw": 5000.0,
    "max_discharge_power_kw": 5000.0,
    "emergency_multiplier": 5.0,
    "scenario_days": 28,
    "mean_days": 7,
    "tolerance": 1e-6,
}

PRIMARY = "saa28_battery"
METHODS = ("mean7_battery", "saa28_battery", "saa28_no_battery")
EVALUATION_START = pd.Timestamp("2025-02-01")
EVALUATION_END = pd.Timestamp("2025-12-31")


# ============================================================
# 通用检查
# ============================================================

def require(condition, message):
    if not bool(condition):
        raise ValueError(message)


def close(actual, expected, label, tol=1e-6):
    actual = float(actual)
    expected = float(expected)
    if not (math.isfinite(actual) and math.isfinite(expected)):
        raise ValueError(f"{label}: non-finite value")
    if abs(actual - expected) > tol:
        raise ValueError(
            f"{label}: {actual} != {expected}, tolerance={tol}"
        )


def clock(minute: int) -> str:
    return f"{minute // 60:02d}:{minute % 60:02d}"


def interval_label(slot: int) -> str:
    start = (slot - 1) * 10
    end = slot * 10
    return f"{clock(start)}-{clock(end)}"


# ============================================================
# 输入
# ============================================================

def load_inputs(actual_path: Path, price_path: Path):
    frame = pd.read_csv(actual_path, float_precision="round_trip")
    price = pd.read_csv(price_path, float_precision="round_trip")

    require(
        list(frame.columns) == ["date", "slot", "load_kw", "pv_kw"],
        "q2_actual_load_pv.csv 列名必须为 date, slot, load_kw, pv_kw",
    )
    require(
        list(price.columns) == ["slot", "price_yuan_per_kwh"],
        "q2_fixed_prices.csv 列名必须为 slot, price_yuan_per_kwh",
    )
    require(len(frame) == 52560, "实际负载/光伏数据应为 365×144=52560 行")
    require(len(price) == 144, "固定电价应为 144 个十分钟时段")
    require(
        frame.groupby("date").size().eq(144).all(),
        "每天必须恰有 144 个十分钟时段",
    )
    require(
        not frame.duplicated(["date", "slot"]).any(),
        "存在重复 date-slot",
    )

    frame = frame.sort_values(["date", "slot"]).reset_index(drop=True)
    dates = pd.DatetimeIndex(sorted(pd.to_datetime(frame.date.unique())))

    require(
        dates.equals(pd.date_range("2025-01-01", "2025-12-31")),
        "日期范围必须完整覆盖 2025-01-01 至 2025-12-31",
    )
    require(
        np.array_equal(
            frame.slot.to_numpy(),
            np.tile(np.arange(1, 145), 365),
        ),
        "slot 必须在每天按 1..144 排列",
    )
    require(
        np.array_equal(price.slot.to_numpy(), np.arange(1, 145)),
        "电价 slot 必须为 1..144",
    )

    numeric = frame[["load_kw", "pv_kw"]].to_numpy(float)
    p = price.price_yuan_per_kwh.to_numpy(float)

    require(np.isfinite(numeric).all(), "负载/光伏存在非有限数值")
    require(np.isfinite(p).all() and (p > 0).all(), "电价必须为有限正数")
    require((numeric >= 0).all(), "负载和光伏功率不能为负")

    load = frame.load_kw.to_numpy(float).reshape(365, 144)
    pv = frame.pv_kw.to_numpy(float).reshape(365, 144)

    return load, pv, p, dates


# ============================================================
# 仅使用历史信息构造情景
# ============================================================

def history_scenarios(load, pv, day_index, method, cfg):
    """计划日只能看到严格早于该日的完整历史日。"""
    if day_index < 1:
        raise ValueError("1月1日没有历史数据，应使用 cold-start 规则")

    window = (
        cfg["mean_days"]
        if method == "mean7_battery"
        else cfg["scenario_days"]
    )
    first = max(0, day_index - window)

    # 功率(kW) -> 每10分钟电量(kWh)
    scenario = (
        (load[first:day_index] - pv[first:day_index])
        * cfg["interval_minutes"] / 60
    )

    if method == "mean7_battery":
        scenario = scenario.mean(axis=0, keepdims=True)

    return scenario, first


# ============================================================
# 日前计划优化
# ============================================================

def plan_day(scenarios, prices, initial_kwh, cfg, battery=True):
    """
    历史情景近似目标：
        min p·g + 5 p·E[h]

    每个历史情景共享同一份日前购电和计划储能轨迹。
    若无储能，逐时经验分位数给出解析最优解。
    """
    a = np.asarray(scenarios, dtype=float)
    p = np.asarray(prices, dtype=float)

    require(
        a.ndim == 2
        and a.shape[1] == len(p)
        and np.isfinite(a).all(),
        "Invalid scenarios",
    )
    require(
        np.isfinite(p).all() and (p > 0).all(),
        "Prices must be positive",
    )

    s, n = a.shape

    if not battery:
        # min g + M * mean((net-g)+), M=5
        # 对应经验 (1-1/M)=80% 分位数。
        rank = max(
            0,
            int(
                np.ceil(
                    (1 - 1 / cfg["emergency_multiplier"]) * s
                )
            ) - 1,
        )
        g = np.maximum(np.sort(a, axis=0)[rank], 0)
        h = np.maximum(a - g, 0)

        return {
            "grid": g,
            "charge": np.zeros(n),
            "discharge": np.zeros(n),
            "storage": np.full(n + 1, initial_kwh),
            "expected_emergency": h.mean(axis=0),
            "objective": float(
                p @ g
                + cfg["emergency_multiplier"] * (h @ p).mean()
            ),
            "solver": "empirical quantile",
            "plan_residual": 0.0,
        }

    # x = [g(n), c(n), d(n), E(n+1), h(s*n)]
    h0 = 4 * n + 1
    m = 4 * n + 1 + s * n

    obj = np.zeros(m)
    obj[:n] = p
    obj[h0:] = np.tile(
        cfg["emergency_multiplier"] * p / s,
        s,
    )

    ec = cfg["charge_efficiency"]
    ed = cfg["discharge_efficiency"]
    mc = (
        cfg["max_charge_power_kw"]
        * cfg["interval_minutes"] / 60
    )
    md = (
        cfg["max_discharge_power_kw"]
        * cfg["interval_minutes"] / 60
    )
    lo = cfg["storage_min_kwh"]
    hi = cfg["storage_max_kwh"]

    lower = np.zeros(m)
    upper = np.full(m, np.inf)

    upper[n:2*n] = mc
    upper[2*n:3*n] = md

    lower[3*n:4*n+1] = lo
    upper[3*n:4*n+1] = hi

    lower[3*n] = initial_kwh
    upper[3*n] = initial_kwh

    lower[4*n] = cfg["planned_terminal_kwh"]
    upper[4*n] = cfg["planned_terminal_kwh"]

    t = np.arange(n)

    # E[t+1]-E[t]-eta_c*c+d/eta_d = 0
    eq = coo_matrix(
        (
            np.r_[
                -ec * np.ones(n),
                np.ones(n) / ed,
                -np.ones(n),
                np.ones(n),
            ],
            (
                np.tile(t, 4),
                np.r_[
                    n + t,
                    2*n + t,
                    3*n + t,
                    3*n + t + 1,
                ],
            ),
        ),
        shape=(n, m),
    ).tocsr()

    # g + h_s + d - c >= net_s
    # -> -g + c - d - h_s <= -net_s
    rows = np.arange(s * n)
    slots = np.tile(t, s)

    ub = coo_matrix(
        (
            np.r_[
                -np.ones(s*n),
                np.ones(s*n),
                -np.ones(s*n),
                -np.ones(s*n),
            ],
            (
                np.tile(rows, 4),
                np.r_[
                    slots,
                    n + slots,
                    2*n + slots,
                    h0 + rows,
                ],
            ),
        ),
        shape=(s*n, m),
    ).tocsr()

    b = -a.ravel()

    lp = linprog(
        obj,
        A_ub=ub,
        b_ub=b,
        A_eq=eq,
        b_eq=np.zeros(n),
        bounds=np.column_stack([lower, upper]),
        method="highs",
    )

    if not lp.success:
        raise RuntimeError("Q2 LP failed: " + lp.message)

    x = lp.x
    solver = "HiGHS LP"

    # LP 若出现同一时段同时充放电，则用二进制变量转 MILP。
    overlap = (
        (x[n:2*n] > cfg["tolerance"])
        & (x[2*n:3*n] > cfg["tolerance"])
    )

    if overlap.any():
        mode = coo_matrix(
            (
                np.r_[
                    np.ones(n),
                    -mc * np.ones(n),
                    np.ones(n),
                    md * np.ones(n),
                ],
                (
                    np.r_[t, t, n+t, n+t],
                    np.r_[n+t, m+t, 2*n+t, m+t],
                ),
            ),
            shape=(2*n, m+n),
        ).tocsr()

        res = milp(
            np.r_[obj, np.zeros(n)],
            integrality=np.r_[np.zeros(m), np.ones(n)],
            bounds=Bounds(
                np.r_[lower, np.zeros(n)],
                np.r_[upper, np.ones(n)],
            ),
            constraints=[
                LinearConstraint(
                    hstack([eq, csr_matrix((n, n))]),
                    0,
                    0,
                ),
                LinearConstraint(
                    hstack([ub, csr_matrix((s*n, n))]),
                    -np.inf,
                    b,
                ),
                LinearConstraint(
                    mode,
                    -np.inf,
                    np.r_[
                        np.zeros(n),
                        md * np.ones(n),
                    ],
                ),
            ],
            options={"mip_rel_gap": 1e-8},
        )

        if not res.success:
            raise RuntimeError("Q2 MILP failed: " + res.message)

        x = res.x[:m]
        solver = "HiGHS MILP"

    plan_residual = float(
        max(
            np.max(ub @ x - b),
            np.max(np.abs(eq @ x)),
            0,
        )
    )

    require(
        plan_residual <= cfg["tolerance"],
        f"Planning constraints failed: residual={plan_residual}",
    )

    return {
        "grid": np.maximum(x[:n], 0),
        "charge": np.maximum(x[n:2*n], 0),
        "discharge": np.maximum(x[2*n:3*n], 0),
        "storage": x[3*n:4*n+1],
        "expected_emergency": (
            x[h0:].reshape(s, n).mean(axis=0)
        ),
        "objective": float(obj @ x),
        "solver": solver,
        "plan_residual": plan_residual,
    }


# ============================================================
# 实时因果执行
# ============================================================

def execute_day(
    grid,
    actual_load,
    actual_pv,
    initial_kwh,
    cfg,
    battery=True,
):
    """
    每个时段只读取：
      - 已经在00:00确定的计划购电量
      - 当前时段真实负载
      - 当前时段真实光伏
      - 当前储能状态

    不读取当天未来时段真实值。
    """
    n = len(grid)
    e = np.empty(n + 1)
    e[0] = initial_kwh

    c = np.zeros(n)
    d = np.zeros(n)
    h = np.zeros(n)
    w = np.zeros(n)

    ec = cfg["charge_efficiency"]
    ed = cfg["discharge_efficiency"]

    mc = (
        cfg["max_charge_power_kw"]
        * cfg["interval_minutes"] / 60
        if battery
        else 0
    )
    md = (
        cfg["max_discharge_power_kw"]
        * cfg["interval_minutes"] / 60
        if battery
        else 0
    )

    for t in range(n):
        surplus = grid[t] + actual_pv[t] - actual_load[t]

        if surplus >= 0:
            c[t] = min(
                surplus,
                mc,
                max(
                    0,
                    (
                        cfg["storage_max_kwh"] - e[t]
                    ) / ec,
                ),
            )
            w[t] = surplus - c[t]

        else:
            d[t] = min(
                -surplus,
                md,
                max(
                    0,
                    (
                        e[t] - cfg["storage_min_kwh"]
                    ) * ed,
                ),
            )
            h[t] = -surplus - d[t]

        e[t+1] = e[t] + ec * c[t] - d[t] / ed

    return {
        "charge": c,
        "discharge": d,
        "emergency": h,
        "unused": w,
        "storage": e,
    }


# ============================================================
# 全年滚动策略
# ============================================================

def run_strategy(
    load,
    pv,
    prices,
    dates,
    method,
    cfg=None,
    progress=True,
):
    cfg = dict(CFG if cfg is None else cfg)

    require(method in METHODS, f"Unknown method: {method}")
    require(
        load.shape == pv.shape
        and load.shape[1] == 144
        and len(prices) == 144,
        "Expect day × 144 inputs",
    )

    state = cfg["initial_storage_kwh"]
    frames = []
    diary = []

    battery = method != "saa28_no_battery"

    for k, date in enumerate(dates):

        if k == 0:
            # 1月1日没有过去数据，只用于冷启动与1月预热。
            plan = {
                "grid": np.zeros(144),
                "charge": np.zeros(144),
                "discharge": np.zeros(144),
                "storage": np.full(145, state),
                "expected_emergency": np.zeros(144),
                "objective": 0.0,
                "solver": "cold-start",
                "plan_residual": 0.0,
            }
            first = 0

        else:
            scenarios, first = history_scenarios(
                load, pv, k, method, cfg
            )
            plan = plan_day(
                scenarios,
                prices,
                state,
                cfg,
                battery=battery,
            )

        actual_load = (
            load[k] * cfg["interval_minutes"] / 60
        )
        actual_pv = (
            pv[k] * cfg["interval_minutes"] / 60
        )

        actual = execute_day(
            plan["grid"],
            actual_load,
            actual_pv,
            state,
            cfg,
            battery=(battery and k > 0),
        )

        e = actual["storage"]
        state = float(e[-1])  # 实际SOC连续传递，不人为重置。

        intervals = pd.date_range(
            date,
            periods=144,
            freq="10min",
        )

        frame = pd.DataFrame(
            {
                "date": str(pd.Timestamp(date).date()),
                "slot": np.arange(1, 145),
                "interval_start": intervals,
                "interval_end": (
                    intervals + pd.Timedelta(minutes=10)
                ),
                "price_yuan_per_kwh": prices,
                "load_kwh": actual_load,
                "pv_kwh": actual_pv,
                "planned_grid_kwh": plan["grid"],
                "planned_charge_kwh": plan["charge"],
                "planned_discharge_kwh": plan["discharge"],
                "planned_storage_start_kwh": (
                    plan["storage"][:-1]
                ),
                "planned_storage_end_kwh": (
                    plan["storage"][1:]
                ),
                "forecast_emergency_kwh": (
                    plan["expected_emergency"]
                ),
                "charge_kwh": actual["charge"],
                "discharge_kwh": actual["discharge"],
                "emergency_kwh": actual["emergency"],
                "unused_energy_kwh": actual["unused"],
                "storage_start_kwh": e[:-1],
                "storage_end_kwh": e[1:],
            }
        )

        frame["planned_cost_yuan"] = (
            prices * plan["grid"]
        )
        frame["emergency_cost_yuan"] = (
            cfg["emergency_multiplier"]
            * prices
            * actual["emergency"]
        )
        frame["total_cost_yuan"] = (
            frame.planned_cost_yuan
            + frame.emergency_cost_yuan
        )
        frame["evaluation"] = (
            pd.Timestamp(date) >= EVALUATION_START
        )

        frames.append(frame)

        diary.append(
            {
                "date": str(pd.Timestamp(date).date()),
                "solver": plan["solver"],
                "history_first_date": (
                    None
                    if k == 0
                    else str(pd.Timestamp(dates[first]).date())
                ),
                "history_last_date": (
                    None
                    if k == 0
                    else str(pd.Timestamp(dates[k-1]).date())
                ),
                "history_days": (
                    0 if k == 0 else k - first
                ),
                "objective_surrogate_yuan": (
                    plan["objective"]
                ),
                "planning_constraint_residual": (
                    plan["plan_residual"]
                ),
                "initial_actual_kwh": float(e[0]),
                "terminal_actual_kwh": float(e[-1]),
            }
        )

        if progress and (
            (k + 1) % 30 == 0
            or k == len(dates) - 1
        ):
            print(
                f"{method}: {k+1}/{len(dates)} days",
                flush=True,
            )

    return (
        pd.concat(frames, ignore_index=True),
        pd.DataFrame(diary),
    )


# ============================================================
# 汇总与物理检查
# ============================================================

def daily_summary(dispatch):
    sumcols = [
        "planned_grid_kwh",
        "emergency_kwh",
        "planned_cost_yuan",
        "emergency_cost_yuan",
        "total_cost_yuan",
        "charge_kwh",
        "discharge_kwh",
        "unused_energy_kwh",
        "load_kwh",
        "pv_kwh",
    ]

    out = dispatch.groupby("date")[sumcols].sum()

    out["initial_storage_kwh"] = (
        dispatch.groupby("date").storage_start_kwh.first()
    )
    out["terminal_storage_kwh"] = (
        dispatch.groupby("date").storage_end_kwh.last()
    )
    out["emergency_slots"] = (
        dispatch.assign(
            event=dispatch.emergency_kwh > 1e-6
        )
        .groupby("date")
        .event.sum()
    )

    return out.reset_index()


def verify_physics(dispatch, cfg=None):
    cfg = CFG if cfg is None else cfg
    x = dispatch
    tol = cfg["tolerance"]

    def maximum(a):
        return float(np.max(np.abs(np.asarray(a))))

    residual = (
        x.planned_grid_kwh
        + x.pv_kwh
        + x.discharge_kwh
        + x.emergency_kwh
        - x.load_kwh
        - x.charge_kwh
        - x.unused_energy_kwh
    )

    evolution = (
        x.storage_end_kwh
        - x.storage_start_kwh
        - cfg["charge_efficiency"] * x.charge_kwh
        + x.discharge_kwh / cfg["discharge_efficiency"]
    )

    values = {
        "energy_balance_kwh": maximum(residual),
        "battery_update_kwh": maximum(evolution),
        "continuous_state_kwh": maximum(
            x.storage_start_kwh.to_numpy()[1:]
            - x.storage_end_kwh.to_numpy()[:-1]
        ),
        "soc_lower_violation_kwh": max(
            0.0,
            float(
                cfg["storage_min_kwh"]
                - x.storage_end_kwh.min()
            ),
        ),
        "soc_upper_violation_kwh": max(
            0.0,
            float(
                x.storage_end_kwh.max()
                - cfg["storage_max_kwh"]
            ),
        ),
        "charge_power_violation_kw": max(
            0.0,
            float(
                x.charge_kwh.max()
                * 60 / cfg["interval_minutes"]
                - cfg["max_charge_power_kw"]
            ),
        ),
        "discharge_power_violation_kw": max(
            0.0,
            float(
                x.discharge_kwh.max()
                * 60 / cfg["interval_minutes"]
                - cfg["max_discharge_power_kw"]
            ),
        ),
        "simultaneous_charge_discharge_slots": int(
            (
                (x.charge_kwh > tol)
                & (x.discharge_kwh > tol)
            ).sum()
        ),
        "planned_cost_error_yuan": maximum(
            x.planned_cost_yuan
            - x.price_yuan_per_kwh
            * x.planned_grid_kwh
        ),
        "emergency_cost_error_yuan": maximum(
            x.emergency_cost_yuan
            - cfg["emergency_multiplier"]
            * x.price_yuan_per_kwh
            * x.emergency_kwh
        ),
        "total_cost_error_yuan": maximum(
            x.total_cost_yuan
            - x.planned_cost_yuan
            - x.emergency_cost_yuan
        ),
        "negative_flow_violation_kwh": max(
            0.0,
            -float(
                x[
                    [
                        "planned_grid_kwh",
                        "charge_kwh",
                        "discharge_kwh",
                        "emergency_kwh",
                        "unused_energy_kwh",
                    ]
                ].min().min()
            ),
        ),
    }

    failures = {
        k: v
        for k, v in values.items()
        if v > tol
    }

    require(
        not failures,
        "Physical checks failed: " + str(failures),
    )
    require(
        np.isfinite(
            x.select_dtypes("number").to_numpy()
        ).all(),
        "Non-finite numeric output",
    )

    starts = pd.to_datetime(
        x.interval_start
    ).iloc[1:].to_numpy()
    ends = pd.to_datetime(
        x.interval_end
    ).iloc[:-1].to_numpy()

    require(
        np.array_equal(starts, ends),
        "Time axis is not continuous",
    )

    return values


def emergency_events(dispatch):
    """将连续的10分钟应急购电槽合并成题目要求的时间段。"""
    rows = []

    for day, x in dispatch.groupby("date", sort=True):
        x = x.sort_values("slot")
        a = x.emergency_kwh.to_numpy()
        t = 0

        while t < len(a):
            if a[t] <= 1e-6:
                t += 1
                continue

            start = t

            while t < len(a) and a[t] > 1e-6:
                t += 1

            rows.append(
                {
                    "date": day,
                    "interval": (
                        f"{clock(start*10)}-{clock(t*10)}"
                    ),
                    "emergency_kwh": float(
                        a[start:t].sum()
                    ),
                    "start_slot": start + 1,
                    "end_slot": t,
                }
            )

    return pd.DataFrame(
        rows,
        columns=[
            "date",
            "interval",
            "emergency_kwh",
            "start_slot",
            "end_slot",
        ],
    )


def validate_primary(
    dispatch,
    diary,
    load,
    pv,
    prices,
    dates,
    cfg,
):
    """最终提交前的核心独立检查。"""
    tol = cfg["tolerance"]

    require(
        len(dispatch) == 52560
        and dispatch.date.nunique() == 365,
        "Full-year trace incomplete",
    )

    ev = dispatch[dispatch.evaluation]

    require(
        len(ev) == 48096
        and ev.date.nunique() == 334,
        "Evaluation trace must contain 334 days × 144 slots",
    )

    require(
        np.array_equal(
            dispatch.price_yuan_per_kwh.to_numpy(),
            np.tile(prices, 365),
        ),
        "Attachment-1 prices changed",
    )

    require(
        np.allclose(
            dispatch.load_kwh.to_numpy(),
            load.ravel() / 6,
            rtol=0,
            atol=1e-10,
        ),
        "Load data changed",
    )
    require(
        np.allclose(
            dispatch.pv_kwh.to_numpy(),
            pv.ravel() / 6,
            rtol=0,
            atol=1e-10,
        ),
        "PV data changed",
    )

    close(
        dispatch.storage_start_kwh.iloc[0],
        cfg["initial_storage_kwh"],
        "Jan-1 initial storage",
        tol,
    )

    daily_start = (
        dispatch.groupby("date")
        .storage_start_kwh.first()
        .to_numpy()
    )
    daily_end = (
        dispatch.groupby("date")
        .storage_end_kwh.last()
        .to_numpy()
    )

    require(
        np.max(np.abs(daily_start[1:] - daily_end[:-1]))
        <= tol,
        "Day-boundary SOC is discontinuous",
    )

    hist_last = pd.to_datetime(
        diary.history_last_date.iloc[1:]
    )
    plan_dates = pd.to_datetime(diary.date.iloc[1:])

    require(
        bool((hist_last < plan_dates).all()),
        "Future-date information leakage detected",
    )
    require(
        bool((diary.history_days <= 28).all()),
        "History window exceeds 28 days",
    )

    close(
        dispatch.planned_cost_yuan.sum(),
        np.dot(
            dispatch.price_yuan_per_kwh,
            dispatch.planned_grid_kwh,
        ),
        "Fixed purchase settlement",
        1e-6,
    )

    close(
        dispatch.emergency_cost_yuan.sum(),
        cfg["emergency_multiplier"]
        * np.dot(
            dispatch.price_yuan_per_kwh,
            dispatch.emergency_kwh,
        ),
        "Emergency settlement",
        1e-6,
    )

    verify_physics(dispatch, cfg)

    # 每天计划时的初始SOC必须等于当天实测开始SOC。
    planned_start = (
        dispatch.groupby("date")
        .planned_storage_start_kwh.first()
        .to_numpy()
    )

    require(
        np.max(np.abs(planned_start - daily_start))
        <= tol,
        "Planning state differs from measured day-start SOC",
    )

    # 除冷启动日外，日前优化均以6000kWh为计划末态目标。
    planned_end = (
        dispatch.groupby("date")
        .planned_storage_end_kwh.last()
        .to_numpy()
    )
    require(
        np.max(
            np.abs(
                planned_end[1:]
                - cfg["planned_terminal_kwh"]
            )
        ) <= tol,
        "Planned terminal SOC target failed",
    )

    require(
        (
            diary.planning_constraint_residual
            <= tol
        ).all(),
        "Some daily planning problems violate constraints",
    )

    # 未来真实数据突变不允许改变某一更早计划日的历史情景。
    origin = 150
    s1, _ = history_scenarios(
        load,
        pv,
        origin,
        PRIMARY,
        cfg,
    )

    changed_load = load.copy()
    changed_pv = pv.copy()
    changed_load[origin:] *= 11
    changed_pv[origin:] = 0

    s2, _ = history_scenarios(
        changed_load,
        changed_pv,
        origin,
        PRIMARY,
        cfg,
    )

    require(
        np.array_equal(s1, s2),
        "Future actual mutation changed historical scenarios",
    )

    # 执行规则前72槽不应读取后72槽未来真实负载。
    base_load = load[100] / 6
    base_pv = pv[100] / 6

    original = execute_day(
        np.full(144, 600.0),
        base_load,
        base_pv,
        6000,
        cfg,
    )

    future_load = base_load.copy()
    future_load[72:] *= 10

    mutated = execute_day(
        np.full(144, 600.0),
        future_load,
        base_pv,
        6000,
        cfg,
    )

    for key in [
        "charge",
        "discharge",
        "emergency",
        "unused",
    ]:
        require(
            np.array_equal(
                original[key][:72],
                mutated[key][:72],
            ),
            f"Execution future lookup detected: {key}",
        )

    require(
        np.array_equal(
            original["storage"][:73],
            mutated["storage"][:73],
        ),
        "Execution future lookup detected: storage",
    )

    return True


# ============================================================
# Excel 模板导出
# ============================================================

def _header_key(value):
    if value is None:
        return ""
    return (
        str(value)
        .replace("\n", "")
        .replace("\r", "")
        .replace(" ", "")
        .strip()
    )


def _copy_row_style(ws, source_row, target_row, max_col):
    """模板行数不足时，复制一行基础样式。"""
    if source_row == target_row:
        return

    if ws.row_dimensions[source_row].height is not None:
        ws.row_dimensions[target_row].height = (
            ws.row_dimensions[source_row].height
        )

    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)

        if src.has_style:
            dst._style = copy(src._style)

        if src.number_format:
            dst.number_format = src.number_format

        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)


def _ensure_rows(ws, needed_last_row):
    if ws.max_row >= needed_last_row:
        return

    source_row = max(2, ws.max_row)

    for row in range(ws.max_row + 1, needed_last_row + 1):
        _copy_row_style(
            ws,
            source_row,
            row,
            ws.max_column,
        )


def _clear_values(ws, start_row, end_row, columns):
    end_row = min(end_row, ws.max_row)
    for row in range(start_row, end_row + 1):
        for col in columns:
            ws.cell(row, col).value = None


def _plan_sheet_columns(ws):
    headers = {
        col: _header_key(ws.cell(1, col).value)
        for col in range(1, ws.max_column + 1)
    }

    date_col = next(
        (
            c
            for c, h in headers.items()
            if h.startswith("日期")
        ),
        None,
    )
    total_col = next(
        (
            c
            for c, h in headers.items()
            if "全天计划购电量" in h
        ),
        None,
    )
    cost_col = next(
        (
            c
            for c, h in headers.items()
            if "全天计划购电费" in h
        ),
        None,
    )

    interval_cols = []

    pattern = re_compile_interval()

    for c, h in headers.items():
        if pattern.fullmatch(h):
            interval_cols.append((c, h))

    require(date_col is not None, "计划购电量 sheet 缺少日期列")
    require(total_col is not None, "计划购电量 sheet 缺少全天计划购电量列")
    require(cost_col is not None, "计划购电量 sheet 缺少全天计划购电费列")
    require(
        len(interval_cols) == 144,
        f"计划购电量 sheet 应有144个时段列，实际 {len(interval_cols)}",
    )

    interval_cols.sort(
        key=lambda item: _interval_start_minutes(item[1])
    )

    expected = [
        interval_label(i)
        for i in range(1, 145)
    ]
    actual = [h for _, h in interval_cols]

    require(
        actual == expected,
        "计划购电量 sheet 的144个时间段标签不符合00:00-24:00十分钟顺序",
    )

    return date_col, interval_cols, total_col, cost_col


def re_compile_interval():
    import re
    return re.compile(
        r"\d{2}:\d{2}-\d{2}:\d{2}"
    )


def _interval_start_minutes(text):
    left = text.split("-", 1)[0]
    h, m = map(int, left.split(":"))
    return h * 60 + m


def _battery_sheet_columns(ws):
    headers = {
        col: _header_key(ws.cell(1, col).value)
        for col in range(1, ws.max_column + 1)
    }

    def find(predicate, label):
        col = next(
            (c for c, h in headers.items() if predicate(h)),
            None,
        )
        require(col is not None, f"充放电量 sheet 缺少 {label} 列")
        return col

    date_col = find(
        lambda h: h == "日期" or h.startswith("日期"),
        "日期",
    )
    interval_col = find(
        lambda h: "时间段" in h,
        "时间段",
    )
    charge_col = find(
        lambda h: "充电量" in h and "放电量" not in h,
        "充电量",
    )
    discharge_col = find(
        lambda h: "放电量" in h,
        "放电量",
    )
    time_col = find(
        lambda h: h == "时刻",
        "时刻",
    )
    storage_col = find(
        lambda h: "储电量" in h,
        "储电量",
    )

    return (
        date_col,
        interval_col,
        charge_col,
        discharge_col,
        time_col,
        storage_col,
    )


def _emergency_sheet_columns(ws):
    headers = {
        col: _header_key(ws.cell(1, col).value)
        for col in range(1, ws.max_column + 1)
    }

    def find(predicate, label):
        col = next(
            (c for c, h in headers.items() if predicate(h)),
            None,
        )
        require(col is not None, f"紧急购电量 sheet 缺少 {label} 列")
        return col

    date_col = find(
        lambda h: h == "日期" or h.startswith("日期"),
        "日期",
    )
    interval_col = find(
        lambda h: "购电时间段" in h or "时间段" in h,
        "紧急购电时间段",
    )
    amount_col = find(
        lambda h: "紧急购电量" in h,
        "紧急购电量",
    )

    return date_col, interval_col, amount_col


def build_submission_tables(primary_dispatch):
    ev = primary_dispatch[
        primary_dispatch.evaluation
    ].copy()

    require(
        ev.date.nunique() == 334,
        "Submission period must contain 334 dates",
    )

    dates = pd.date_range(
        EVALUATION_START,
        EVALUATION_END,
        freq="D",
    )

    # ---------- 计划购电 ----------
    daily = daily_summary(ev).set_index("date")

    plan_rows = []

    for date in dates:
        day = str(date.date())
        x = ev[ev.date == day].sort_values("slot")

        require(
            len(x) == 144,
            f"{day}: planned purchase does not contain 144 slots",
        )

        plan_rows.append(
            {
                "date": day,
                "grid": x.planned_grid_kwh.to_numpy(float),
                "total_grid": float(
                    x.planned_grid_kwh.sum()
                ),
                "cost": float(
                    x.planned_cost_yuan.sum()
                ),
            }
        )

    # ---------- 充放电 ----------
    battery_rows = []

    for date in dates:
        day = str(date.date())
        x = ev[ev.date == day].sort_values("slot")

        for block in range(6):
            start_slot = block * 24 + 1
            end_slot = (block + 1) * 24

            z = x[
                (x.slot >= start_slot)
                & (x.slot <= end_slot)
            ]

            battery_rows.append(
                {
                    "date": day,
                    "interval": (
                        f"{clock(block*240)}-"
                        f"{clock((block+1)*240)}"
                    ),
                    "charge_kwh": float(
                        z.charge_kwh.sum()
                    ),
                    "discharge_kwh": float(
                        z.discharge_kwh.sum()
                    ),
                    "time": (
                        "00:00"
                        if block == 0
                        else (
                            "24:00"
                            if block == 1
                            else None
                        )
                    ),
                    "storage_kwh": (
                        float(
                            x.storage_start_kwh.iloc[0]
                        )
                        if block == 0
                        else (
                            float(
                                x.storage_end_kwh.iloc[-1]
                            )
                            if block == 1
                            else None
                        )
                    ),
                }
            )

    # ---------- 紧急购电 ----------
    raw_events = emergency_events(ev)
    emergency_rows = []

    for date in dates:
        day = str(date.date())
        rows = raw_events[raw_events.date == day]

        if len(rows) == 0:
            emergency_rows.append(
                {
                    "date": day,
                    "interval": "无",
                    "emergency_kwh": 0.0,
                }
            )
        else:
            for _, row in rows.iterrows():
                emergency_rows.append(
                    {
                        "date": day,
                        "interval": row["interval"],
                        "emergency_kwh": float(
                            row["emergency_kwh"]
                        ),
                    }
                )

    return plan_rows, battery_rows, emergency_rows


def export_result2(template_path, output_path, primary_dispatch):
    template_path = Path(template_path).resolve()
    output_path = Path(output_path).resolve()

    require(
        template_path.exists(),
        f"Missing result2 template: {template_path}",
    )
    require(
        template_path != output_path,
        "禁止直接覆盖附件5原始 result2.xlsx 模板",
    )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    # 先复制模板，再仅填写结果单元格。
    shutil.copy2(template_path, output_path)

    wb = load_workbook(output_path)

    required_sheets = {
        "计划购电量",
        "充放电量",
        "紧急购电量",
    }
    require(
        required_sheets.issubset(set(wb.sheetnames)),
        "result2.xlsx 模板必须包含：计划购电量、充放电量、紧急购电量",
    )

    plan_rows, battery_rows, emergency_rows = (
        build_submission_tables(primary_dispatch)
    )

    # --------------------------------------------------------
    # 1) 计划购电量
    # --------------------------------------------------------
    ws = wb["计划购电量"]

    (
        date_col,
        interval_cols,
        total_col,
        cost_col,
    ) = _plan_sheet_columns(ws)

    last_row = 1 + len(plan_rows)
    _ensure_rows(ws, last_row)

    for i, item in enumerate(plan_rows, start=2):
        ws.cell(i, date_col).value = pd.Timestamp(
            item["date"]
        ).to_pydatetime()

        for slot, (col, _) in enumerate(
            interval_cols,
            start=1,
        ):
            ws.cell(i, col).value = float(
                item["grid"][slot - 1]
            )

        ws.cell(i, total_col).value = item["total_grid"]
        ws.cell(i, cost_col).value = item["cost"]

    # 若使用的并非空模板，清掉评价期之后残留的旧数据。
    if ws.max_row > last_row:
        _clear_values(
            ws,
            last_row + 1,
            ws.max_row,
            [
                date_col,
                *[c for c, _ in interval_cols],
                total_col,
                cost_col,
            ],
        )

    # --------------------------------------------------------
    # 2) 充放电量
    # --------------------------------------------------------
    ws = wb["充放电量"]

    (
        date_col,
        interval_col,
        charge_col,
        discharge_col,
        time_col,
        storage_col,
    ) = _battery_sheet_columns(ws)

    last_row = 1 + len(battery_rows)
    _ensure_rows(ws, last_row)

    for i, item in enumerate(battery_rows, start=2):
        ws.cell(i, date_col).value = pd.Timestamp(
            item["date"]
        ).to_pydatetime()
        ws.cell(i, interval_col).value = item["interval"]
        ws.cell(i, charge_col).value = item["charge_kwh"]
        ws.cell(i, discharge_col).value = item["discharge_kwh"]
        ws.cell(i, time_col).value = item["time"]
        ws.cell(i, storage_col).value = item["storage_kwh"]

    if ws.max_row > last_row:
        _clear_values(
            ws,
            last_row + 1,
            ws.max_row,
            [
                date_col,
                interval_col,
                charge_col,
                discharge_col,
                time_col,
                storage_col,
            ],
        )

    # --------------------------------------------------------
    # 3) 紧急购电量
    # --------------------------------------------------------
    ws = wb["紧急购电量"]

    (
        date_col,
        interval_col,
        amount_col,
    ) = _emergency_sheet_columns(ws)

    last_row = 1 + len(emergency_rows)
    _ensure_rows(ws, last_row)

    for i, item in enumerate(emergency_rows, start=2):
        ws.cell(i, date_col).value = pd.Timestamp(
            item["date"]
        ).to_pydatetime()
        ws.cell(i, interval_col).value = item["interval"]
        ws.cell(i, amount_col).value = item["emergency_kwh"]

    if ws.max_row > last_row:
        _clear_values(
            ws,
            last_row + 1,
            ws.max_row,
            [date_col, interval_col, amount_col],
        )

    wb.save(output_path)

    return {
        "plan_rows": len(plan_rows),
        "battery_rows": len(battery_rows),
        "emergency_rows": len(emergency_rows),
    }


# ============================================================
# 保存后的 Excel 独立核验
# ============================================================

def _as_date_string(value):
    if value is None:
        return None
    try:
        return str(pd.Timestamp(value).date())
    except Exception:
        return str(value).strip()


def verify_result2(output_path, primary_dispatch):
    output_path = Path(output_path)

    require(
        output_path.exists(),
        f"Output workbook not found: {output_path}",
    )

    wb = load_workbook(
        output_path,
        data_only=True,
        read_only=False,
    )

    for sheet in [
        "计划购电量",
        "充放电量",
        "紧急购电量",
    ]:
        require(
            sheet in wb.sheetnames,
            f"Saved workbook missing sheet: {sheet}",
        )

    plan_rows, battery_rows, emergency_rows = (
        build_submission_tables(primary_dispatch)
    )

    tol = CFG["tolerance"]

    # ---------- 计划购电 ----------
    ws = wb["计划购电量"]

    (
        date_col,
        interval_cols,
        total_col,
        cost_col,
    ) = _plan_sheet_columns(ws)

    saved_grid_total = 0.0
    saved_cost_total = 0.0

    for i, expected in enumerate(plan_rows, start=2):
        require(
            _as_date_string(
                ws.cell(i, date_col).value
            ) == expected["date"],
            f"计划购电量 row {i}: date mismatch",
        )

        values = []

        for col, _ in interval_cols:
            value = ws.cell(i, col).value
            require(
                value is not None,
                f"计划购电量 row {i}: blank interval cell",
            )
            values.append(float(value))

        require(
            np.allclose(
                values,
                expected["grid"],
                rtol=0,
                atol=tol,
            ),
            f"计划购电量 row {i}: 144-slot values mismatch",
        )

        close(
            ws.cell(i, total_col).value,
            expected["total_grid"],
            f"计划购电量 row {i}: daily grid",
            tol,
        )
        close(
            ws.cell(i, cost_col).value,
            expected["cost"],
            f"计划购电量 row {i}: daily cost",
            tol,
        )

        saved_grid_total += sum(values)
        saved_cost_total += float(
            ws.cell(i, cost_col).value
        )

    # ---------- 充放电 ----------
    ws = wb["充放电量"]

    (
        date_col,
        interval_col,
        charge_col,
        discharge_col,
        time_col,
        storage_col,
    ) = _battery_sheet_columns(ws)

    saved_charge = 0.0
    saved_discharge = 0.0

    for i, expected in enumerate(battery_rows, start=2):
        require(
            _as_date_string(
                ws.cell(i, date_col).value
            ) == expected["date"],
            f"充放电量 row {i}: date mismatch",
        )
        require(
            str(ws.cell(i, interval_col).value).strip()
            == expected["interval"],
            f"充放电量 row {i}: interval mismatch",
        )

        close(
            ws.cell(i, charge_col).value,
            expected["charge_kwh"],
            f"充放电量 row {i}: charge",
            tol,
        )
        close(
            ws.cell(i, discharge_col).value,
            expected["discharge_kwh"],
            f"充放电量 row {i}: discharge",
            tol,
        )

        saved_time = ws.cell(i, time_col).value
        expected_time = expected["time"]

        if expected_time is None:
            require(
                saved_time is None
                or str(saved_time).strip() == "",
                f"充放电量 row {i}: unexpected time label",
            )
        else:
            require(
                str(saved_time).strip() == expected_time,
                f"充放电量 row {i}: time label mismatch",
            )

        saved_storage = ws.cell(i, storage_col).value

        if expected["storage_kwh"] is None:
            require(
                saved_storage is None
                or str(saved_storage).strip() == "",
                f"充放电量 row {i}: unexpected storage value",
            )
        else:
            close(
                saved_storage,
                expected["storage_kwh"],
                f"充放电量 row {i}: storage",
                tol,
            )

        saved_charge += float(
            ws.cell(i, charge_col).value
        )
        saved_discharge += float(
            ws.cell(i, discharge_col).value
        )

    # ---------- 紧急购电 ----------
    ws = wb["紧急购电量"]

    (
        date_col,
        interval_col,
        amount_col,
    ) = _emergency_sheet_columns(ws)

    saved_emergency = 0.0

    for i, expected in enumerate(emergency_rows, start=2):
        require(
            _as_date_string(
                ws.cell(i, date_col).value
            ) == expected["date"],
            f"紧急购电量 row {i}: date mismatch",
        )
        require(
            str(ws.cell(i, interval_col).value).strip()
            == expected["interval"],
            f"紧急购电量 row {i}: interval mismatch",
        )
        close(
            ws.cell(i, amount_col).value,
            expected["emergency_kwh"],
            f"紧急购电量 row {i}: amount",
            tol,
        )

        saved_emergency += float(
            ws.cell(i, amount_col).value
        )

    ev = primary_dispatch[
        primary_dispatch.evaluation
    ]

    close(
        saved_grid_total,
        ev.planned_grid_kwh.sum(),
        "Saved workbook total planned purchase",
        1e-5,
    )
    close(
        saved_cost_total,
        ev.planned_cost_yuan.sum(),
        "Saved workbook total planned cost",
        1e-5,
    )
    close(
        saved_charge,
        ev.charge_kwh.sum(),
        "Saved workbook total battery charge",
        1e-5,
    )
    close(
        saved_discharge,
        ev.discharge_kwh.sum(),
        "Saved workbook total battery discharge",
        1e-5,
    )
    close(
        saved_emergency,
        ev.emergency_kwh.sum(),
        "Saved workbook total emergency purchase",
        1e-5,
    )

    return {
        "status": "PASS",
        "evaluation_days": 334,
        "plan_rows": len(plan_rows),
        "battery_rows": len(battery_rows),
        "emergency_rows": len(emergency_rows),
        "planned_purchase_kwh": float(
            ev.planned_grid_kwh.sum()
        ),
        "emergency_purchase_kwh": float(
            ev.emergency_kwh.sum()
        ),
        "planned_cost_yuan": float(
            ev.planned_cost_yuan.sum()
        ),
        "emergency_cost_yuan": float(
            ev.emergency_cost_yuan.sum()
        ),
        "total_cost_yuan": float(
            ev.total_cost_yuan.sum()
        ),
    }


# ============================================================
# 可选策略比较（仅终端显示）
# ============================================================

def compare_strategies(results):
    rows = []

    for name, (allx, _) in results.items():
        x = allx[allx.evaluation]
        d = daily_summary(x)

        rows.append(
            {
                "strategy": name,
                "days": len(d),
                "planned_grid_kwh": float(
                    x.planned_grid_kwh.sum()
                ),
                "emergency_kwh": float(
                    x.emergency_kwh.sum()
                ),
                "planned_cost_yuan": float(
                    x.planned_cost_yuan.sum()
                ),
                "emergency_cost_yuan": float(
                    x.emergency_cost_yuan.sum()
                ),
                "total_cost_yuan": float(
                    x.total_cost_yuan.sum()
                ),
                "emergency_days": int(
                    (d.emergency_slots > 0).sum()
                ),
            }
        )

    return pd.DataFrame(rows)


# ============================================================
# 主程序
# ============================================================

def main():
    parser = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawTextHelpFormatter,
    )

    parser.add_argument(
        "--actual",
        type=Path,
        default=DEFAULT_ACTUAL,
        help="全年实际负载/光伏 CSV",
    )
    parser.add_argument(
        "--price",
        type=Path,
        default=DEFAULT_PRICE,
        help="144个固定电价 CSV",
    )
    parser.add_argument(
        "--template",
        type=Path,
        default=None,
        help="官方 result2.xlsx；默认自动查找附件5",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="最终 result2.xlsx 输出位置",
    )
    parser.add_argument(
        "--compare",
        action="store_true",
        help="额外运行两个对照策略并仅在终端打印比较",
    )
    parser.add_argument(
        "--quiet",
        action="store_true",
        help="关闭逐月进度显示",
    )

    args = parser.parse_args()

    started = time.perf_counter()

    actual_path = args.actual.resolve()
    price_path = args.price.resolve()
    output_path = args.output.resolve()

    template_path = (
        args.template.resolve()
        if args.template is not None
        else default_template_path().resolve()
    )

    require(
        actual_path.exists(),
        f"Missing actual input: {actual_path}",
    )
    require(
        price_path.exists(),
        f"Missing price input: {price_path}",
    )
    require(
        template_path.exists(),
        f"Missing result2 template: {template_path}",
    )

    print("Loading Q2 inputs...")
    load, pv, prices, dates = load_inputs(
        actual_path,
        price_path,
    )

    print(
        "Running primary strategy: "
        "28-day historical SAA + battery"
    )

    primary_dispatch, primary_diary = run_strategy(
        load,
        pv,
        prices,
        dates,
        PRIMARY,
        CFG,
        progress=not args.quiet,
    )

    print("Checking physical and causal constraints...")

    validate_primary(
        primary_dispatch,
        primary_diary,
        load,
        pv,
        prices,
        dates,
        CFG,
    )

    print("Writing official result2.xlsx...")

    export_result2(
        template_path,
        output_path,
        primary_dispatch,
    )

    print("Re-opening and verifying saved workbook...")

    check = verify_result2(
        output_path,
        primary_dispatch,
    )

    if args.compare:
        print("\nRunning optional comparison strategies...")

        results = {
            PRIMARY: (
                primary_dispatch,
                primary_diary,
            )
        }

        for method in METHODS:
            if method == PRIMARY:
                continue

            results[method] = run_strategy(
                load,
                pv,
                prices,
                dates,
                method,
                CFG,
                progress=not args.quiet,
            )

        comparison = compare_strategies(results)

        print("\nStrategy comparison:")
        print(
            comparison.to_string(
                index=False,
                float_format=lambda v: f"{v:.4f}",
            )
        )

    elapsed = time.perf_counter() - started

    print("\n" + "=" * 68)
    print("Q2 COMPLETED: PASS")
    print("Primary strategy: 28-day historical SAA + battery")
    print(
        f"Evaluation: 2025-02-01 to 2025-12-31 "
        f"({check['evaluation_days']} days)"
    )
    print(
        f"Planned purchase: "
        f"{check['planned_purchase_kwh']:.4f} kWh"
    )
    print(
        f"Emergency purchase: "
        f"{check['emergency_purchase_kwh']:.4f} kWh"
    )
    print(
        f"Planned cost: "
        f"{check['planned_cost_yuan']:.4f} yuan"
    )
    print(
        f"Emergency cost: "
        f"{check['emergency_cost_yuan']:.4f} yuan"
    )
    print(
        f"Total cost: "
        f"{check['total_cost_yuan']:.4f} yuan"
    )
    print("Physical/causal checks: PASS")
    print("Saved-XLSX verification: PASS")
    print(f"Output: {output_path}")
    print("Generated artifacts: result2.xlsx only")
    print(f"Elapsed: {elapsed:.1f} seconds")
    print("=" * 68)


if __name__ == "__main__":
    main()
