"""
CUMCM 2026 C题 Q2：单文件最终版

默认运行：
    python q2.py

作用：
1. 读取 q2/input/q2_actual_load_pv.csv 和 q2/input/q2_fixed_prices.csv；
2. 每天 00:00 仅使用此前已经结束的历史日数据构造情景；
3. 使用 28 日历史情景 SAA + 储能作为最终策略；
4. 按因果实时规则执行储能与紧急购电；
5. 完成物理约束、费用、历史信息泄漏等检查；
6. 直接填写附件5中的官方 result2.xlsx；
7. 严格保留官方 result2.xlsx 的工作表、表头、日期排版和时间标签，并在保存后重新核验。

默认只生成：
    q2/results/result2.xlsx

不生成图片，不生成中间 CSV/JSON/Markdown，不依赖 parameters.json、
q2_core.py、run_q2.py、check_q2.py 或 JavaScript 导出脚本。

官方模板的“计划购电量”从 0:10-0:20 开始，以 0:00-0:10+1 结束。
当前模型沿用前一问的“时刻样本代表此前10分钟平均功率”解释，因此导出时
按官方标签将同一天的 144 个物理区间循环平移一格：
    0:10-0:20      <- 物理 0:10-0:20
    ...
    23:50-0:00+1   <- 物理 23:50-24:00
    0:00-0:10+1    <- 物理 0:00-0:10
这样每个自然日的 144 个十分钟物理区间恰好出现一次，全天计划量和费用不变。

可选：
    python q2.py --compare
仅额外在终端比较 7日均值、28日SAA+储能、28日SAA无储能，
仍不生成中间文件。
"""

from __future__ import annotations

from copy import copy
from pathlib import Path
import argparse
import math
import shutil
import time

import numpy as np
import pandas as pd
from scipy.optimize import Bounds, LinearConstraint, linprog, milp
from scipy.sparse import coo_matrix, csr_matrix, hstack

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise ImportError(
        "需要 openpyxl 才能填写 result2.xlsx。请运行: pip install openpyxl"
    ) from exc


# ============================================================
# 路径
# ============================================================

Q2_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = Q2_DIR.parent

INPUT_DIR = Q2_DIR / "input"
DEFAULT_ACTUAL = INPUT_DIR / "q2_actual_load_pv.csv"
DEFAULT_PRICE = INPUT_DIR / "q2_fixed_prices.csv"
DEFAULT_OUTPUT = Q2_DIR / "results" / "result2.xlsx"


def default_template_path() -> Path:
    """优先使用整理后的 templates；没有时直接读取附件5中的官方原模板。"""
    candidates = [
        PROJECT_ROOT / "Data_preprocessed" / "templates" / "result2.xlsx",
        PROJECT_ROOT / "Data_preprocessed" / "raw" / "附件5" / "result2.xlsx",
    ]
    for path in candidates:
        if path.exists():
            return path
    return candidates[0]


# ============================================================
# 固定模型参数
# ============================================================

CFG = {
    "interval_minutes": 10,
    "charge_efficiency": 0.9,
    "discharge_efficiency": 0.9,
    "storage_min_kwh": 1200.0,
    "storage_max_kwh": 10800.0,
    "initial_storage_kwh": 6000.0,
    "planned_terminal_kwh": 6000.0,
    "max_charge_power_kw": 5000.0,
    "max_discharge_power_kw": 5000.0,
    "emergency_multiplier": 5.0,
    "scenario_days": 28,
    "mean_days": 7,
    "tolerance": 1e-6,
}

PRIMARY = "saa28_battery"
METHODS = ("mean7_battery", "saa28_battery", "saa28_no_battery")
EVALUATION_START = pd.Timestamp("2025-02-01")
EVALUATION_END = pd.Timestamp("2025-12-31")


# ============================================================
# 通用检查
# ============================================================

def require(condition, message):
    if not bool(condition):
        raise ValueError(message)


def close(actual, expected, label, tol=1e-6):
    actual = float(actual)
    expected = float(expected)
    if not (math.isfinite(actual) and math.isfinite(expected)):
        raise ValueError(f"{label}: non-finite value")
    if abs(actual - expected) > tol:
        raise ValueError(
            f"{label}: {actual} != {expected}, tolerance={tol}"
        )


def clock(minute: int) -> str:
    return f"{minute // 60:02d}:{minute % 60:02d}"


def interval_label(slot: int) -> str:
    start = (slot - 1) * 10
    end = slot * 10
    return f"{clock(start)}-{clock(end)}"


# ============================================================
# 输入
# ============================================================

def load_inputs(actual_path: Path, price_path: Path):
    frame = pd.read_csv(actual_path, float_precision="round_trip")
    price = pd.read_csv(price_path, float_precision="round_trip")

    require(
        list(frame.columns) == ["date", "slot", "load_kw", "pv_kw"],
        "q2_actual_load_pv.csv 列名必须为 date, slot, load_kw, pv_kw",
    )
    require(
        list(price.columns) == ["slot", "price_yuan_per_kwh"],
        "q2_fixed_prices.csv 列名必须为 slot, price_yuan_per_kwh",
    )
    require(len(frame) == 52560, "实际负载/光伏数据应为 365×144=52560 行")
    require(len(price) == 144, "固定电价应为 144 个十分钟时段")
    require(
        frame.groupby("date").size().eq(144).all(),
        "每天必须恰有 144 个十分钟时段",
    )
    require(
        not frame.duplicated(["date", "slot"]).any(),
        "存在重复 date-slot",
    )

    frame = frame.sort_values(["date", "slot"]).reset_index(drop=True)
    dates = pd.DatetimeIndex(sorted(pd.to_datetime(frame.date.unique())))

    require(
        dates.equals(pd.date_range("2025-01-01", "2025-12-31")),
        "日期范围必须完整覆盖 2025-01-01 至 2025-12-31",
    )
    require(
        np.array_equal(
            frame.slot.to_numpy(),
            np.tile(np.arange(1, 145), 365),
        ),
        "slot 必须在每天按 1..144 排列",
    )
    require(
        np.array_equal(price.slot.to_numpy(), np.arange(1, 145)),
        "电价 slot 必须为 1..144",
    )

    numeric = frame[["load_kw", "pv_kw"]].to_numpy(float)
    p = price.price_yuan_per_kwh.to_numpy(float)

    require(np.isfinite(numeric).all(), "负载/光伏存在非有限数值")
    require(np.isfinite(p).all() and (p > 0).all(), "电价必须为有限正数")
    require((numeric >= 0).all(), "负载和光伏功率不能为负")

    load = frame.load_kw.to_numpy(float).reshape(365, 144)
    pv = frame.pv_kw.to_numpy(float).reshape(365, 144)

    return load, pv, p, dates


# ============================================================
# 仅使用历史信息构造情景
# ============================================================

def history_scenarios(load, pv, day_index, method, cfg):
    """计划日只能看到严格早于该日的完整历史日。"""
    if day_index < 1:
        raise ValueError("1月1日没有历史数据，应使用 cold-start 规则")

    window = (
        cfg["mean_days"]
        if method == "mean7_battery"
        else cfg["scenario_days"]
    )
    first = max(0, day_index - window)

    # 功率(kW) -> 每10分钟电量(kWh)
    scenario = (
        (load[first:day_index] - pv[first:day_index])
        * cfg["interval_minutes"] / 60
    )

    if method == "mean7_battery":
        scenario = scenario.mean(axis=0, keepdims=True)

    return scenario, first


# ============================================================
# 日前计划优化
# ============================================================

def plan_day(scenarios, prices, initial_kwh, cfg, battery=True):
    """
    历史情景近似目标：
        min p·g + 5 p·E[h]

    每个历史情景共享同一份日前购电和计划储能轨迹。
    若无储能，逐时经验分位数给出解析最优解。
    """
    a = np.asarray(scenarios, dtype=float)
    p = np.asarray(prices, dtype=float)

    require(
        a.ndim == 2
        and a.shape[1] == len(p)
        and np.isfinite(a).all(),
        "Invalid scenarios",
    )
    require(
        np.isfinite(p).all() and (p > 0).all(),
        "Prices must be positive",
    )

    s, n = a.shape

    if not battery:
        # min g + M * mean((net-g)+), M=5
        # 对应经验 (1-1/M)=80% 分位数。
        rank = max(
            0,
            int(
                np.ceil(
                    (1 - 1 / cfg["emergency_multiplier"]) * s
                )
            ) - 1,
        )
        g = np.maximum(np.sort(a, axis=0)[rank], 0)
        h = np.maximum(a - g, 0)

        return {
            "grid": g,
            "charge": np.zeros(n),
            "discharge": np.zeros(n),
            "storage": np.full(n + 1, initial_kwh),
            "expected_emergency": h.mean(axis=0),
            "objective": float(
                p @ g
                + cfg["emergency_multiplier"] * (h @ p).mean()
            ),
            "solver": "empirical quantile",
            "plan_residual": 0.0,
        }

    # x = [g(n), c(n), d(n), E(n+1), h(s*n)]
    h0 = 4 * n + 1
    m = 4 * n + 1 + s * n

    obj = np.zeros(m)
    obj[:n] = p
    obj[h0:] = np.tile(
        cfg["emergency_multiplier"] * p / s,
        s,
    )

    ec = cfg["charge_efficiency"]
    ed = cfg["discharge_efficiency"]
    mc = (
        cfg["max_charge_power_kw"]
        * cfg["interval_minutes"] / 60
    )
    md = (
        cfg["max_discharge_power_kw"]
        * cfg["interval_minutes"] / 60
    )
    lo = cfg["storage_min_kwh"]
    hi = cfg["storage_max_kwh"]

    lower = np.zeros(m)
    upper = np.full(m, np.inf)

    upper[n:2*n] = mc
    upper[2*n:3*n] = md

    lower[3*n:4*n+1] = lo
    upper[3*n:4*n+1] = hi

    lower[3*n] = initial_kwh
    upper[3*n] = initial_kwh

    lower[4*n] = cfg["planned_terminal_kwh"]
    upper[4*n] = cfg["planned_terminal_kwh"]

    t = np.arange(n)

    # E[t+1]-E[t]-eta_c*c+d/eta_d = 0
    eq = coo_matrix(
        (
            np.r_[
                -ec * np.ones(n),
                np.ones(n) / ed,
                -np.ones(n),
                np.ones(n),
            ],
            (
                np.tile(t, 4),
                np.r_[
                    n + t,
                    2*n + t,
                    3*n + t,
                    3*n + t + 1,
                ],
            ),
        ),
        shape=(n, m),
    ).tocsr()

    # g + h_s + d - c >= net_s
    # -> -g + c - d - h_s <= -net_s
    rows = np.arange(s * n)
    slots = np.tile(t, s)

    ub = coo_matrix(
        (
            np.r_[
                -np.ones(s*n),
                np.ones(s*n),
                -np.ones(s*n),
                -np.ones(s*n),
            ],
            (
                np.tile(rows, 4),
                np.r_[
                    slots,
                    n + slots,
                    2*n + slots,
                    h0 + rows,
                ],
            ),
        ),
        shape=(s*n, m),
    ).tocsr()

    b = -a.ravel()

    lp = linprog(
        obj,
        A_ub=ub,
        b_ub=b,
        A_eq=eq,
        b_eq=np.zeros(n),
        bounds=np.column_stack([lower, upper]),
        method="highs",
    )

    if not lp.success:
        raise RuntimeError("Q2 LP failed: " + lp.message)

    x = lp.x
    solver = "HiGHS LP"

    # LP 若出现同一时段同时充放电，则用二进制变量转 MILP。
    overlap = (
        (x[n:2*n] > cfg["tolerance"])
        & (x[2*n:3*n] > cfg["tolerance"])
    )

    if overlap.any():
        mode = coo_matrix(
            (
                np.r_[
                    np.ones(n),
                    -mc * np.ones(n),
                    np.ones(n),
                    md * np.ones(n),
                ],
                (
                    np.r_[t, t, n+t, n+t],
                    np.r_[n+t, m+t, 2*n+t, m+t],
                ),
            ),
            shape=(2*n, m+n),
        ).tocsr()

        res = milp(
            np.r_[obj, np.zeros(n)],
            integrality=np.r_[np.zeros(m), np.ones(n)],
            bounds=Bounds(
                np.r_[lower, np.zeros(n)],
                np.r_[upper, np.ones(n)],
            ),
            constraints=[
                LinearConstraint(
                    hstack([eq, csr_matrix((n, n))]),
                    0,
                    0,
                ),
                LinearConstraint(
                    hstack([ub, csr_matrix((s*n, n))]),
                    -np.inf,
                    b,
                ),
                LinearConstraint(
                    mode,
                    -np.inf,
                    np.r_[
                        np.zeros(n),
                        md * np.ones(n),
                    ],
                ),
            ],
            options={"mip_rel_gap": 1e-8},
        )

        if not res.success:
            raise RuntimeError("Q2 MILP failed: " + res.message)

        x = res.x[:m]
        solver = "HiGHS MILP"

    plan_residual = float(
        max(
            np.max(ub @ x - b),
            np.max(np.abs(eq @ x)),
            0,
        )
    )

    require(
        plan_residual <= cfg["tolerance"],
        f"Planning constraints failed: residual={plan_residual}",
    )

    return {
        "grid": np.maximum(x[:n], 0),
        "charge": np.maximum(x[n:2*n], 0),
        "discharge": np.maximum(x[2*n:3*n], 0),
        "storage": x[3*n:4*n+1],
        "expected_emergency": (
            x[h0:].reshape(s, n).mean(axis=0)
        ),
        "objective": float(obj @ x),
        "solver": solver,
        "plan_residual": plan_residual,
    }


# ============================================================
# 实时因果执行
# ============================================================

def execute_day(
    grid,
    actual_load,
    actual_pv,
    initial_kwh,
    cfg,
    battery=True,
):
    """
    每个时段只读取：
      - 已经在00:00确定的计划购电量
      - 当前时段真实负载
      - 当前时段真实光伏
      - 当前储能状态

    不读取当天未来时段真实值。
    """
    n = len(grid)
    e = np.empty(n + 1)
    e[0] = initial_kwh

    c = np.zeros(n)
    d = np.zeros(n)
    h = np.zeros(n)
    w = np.zeros(n)

    ec = cfg["charge_efficiency"]
    ed = cfg["discharge_efficiency"]

    mc = (
        cfg["max_charge_power_kw"]
        * cfg["interval_minutes"] / 60
        if battery
        else 0
    )
    md = (
        cfg["max_discharge_power_kw"]
        * cfg["interval_minutes"] / 60
        if battery
        else 0
    )

    for t in range(n):
        surplus = grid[t] + actual_pv[t] - actual_load[t]

        if surplus >= 0:
            c[t] = min(
                surplus,
                mc,
                max(
                    0,
                    (
                        cfg["storage_max_kwh"] - e[t]
                    ) / ec,
                ),
            )
            w[t] = surplus - c[t]

        else:
            d[t] = min(
                -surplus,
                md,
                max(
                    0,
                    (
                        e[t] - cfg["storage_min_kwh"]
                    ) * ed,
                ),
            )
            h[t] = -surplus - d[t]

        e[t+1] = e[t] + ec * c[t] - d[t] / ed

    return {
        "charge": c,
        "discharge": d,
        "emergency": h,
        "unused": w,
        "storage": e,
    }


# ============================================================
# 全年滚动策略
# ============================================================

def run_strategy(
    load,
    pv,
    prices,
    dates,
    method,
    cfg=None,
    progress=True,
):
    cfg = dict(CFG if cfg is None else cfg)

    require(method in METHODS, f"Unknown method: {method}")
    require(
        load.shape == pv.shape
        and load.shape[1] == 144
        and len(prices) == 144,
        "Expect day × 144 inputs",
    )

    state = cfg["initial_storage_kwh"]
    frames = []
    diary = []

    battery = method != "saa28_no_battery"

    for k, date in enumerate(dates):

        if k == 0:
            # 1月1日没有过去数据，只用于冷启动与1月预热。
            plan = {
                "grid": np.zeros(144),
                "charge": np.zeros(144),
                "discharge": np.zeros(144),
                "storage": np.full(145, state),
                "expected_emergency": np.zeros(144),
                "objective": 0.0,
                "solver": "cold-start",
                "plan_residual": 0.0,
            }
            first = 0

        else:
            scenarios, first = history_scenarios(
                load, pv, k, method, cfg
            )
            plan = plan_day(
                scenarios,
                prices,
                state,
                cfg,
                battery=battery,
            )

        actual_load = (
            load[k] * cfg["interval_minutes"] / 60
        )
        actual_pv = (
            pv[k] * cfg["interval_minutes"] / 60
        )

        actual = execute_day(
            plan["grid"],
            actual_load,
            actual_pv,
            state,
            cfg,
            battery=(battery and k > 0),
        )

        e = actual["storage"]
        state = float(e[-1])  # 实际SOC连续传递，不人为重置。

        intervals = pd.date_range(
            date,
            periods=144,
            freq="10min",
        )

        frame = pd.DataFrame(
            {
                "date": str(pd.Timestamp(date).date()),
                "slot": np.arange(1, 145),
                "interval_start": intervals,
                "interval_end": (
                    intervals + pd.Timedelta(minutes=10)
                ),
                "price_yuan_per_kwh": prices,
                "load_kwh": actual_load,
                "pv_kwh": actual_pv,
                "planned_grid_kwh": plan["grid"],
                "planned_charge_kwh": plan["charge"],
                "planned_discharge_kwh": plan["discharge"],
                "planned_storage_start_kwh": (
                    plan["storage"][:-1]
                ),
                "planned_storage_end_kwh": (
                    plan["storage"][1:]
                ),
                "forecast_emergency_kwh": (
                    plan["expected_emergency"]
                ),
                "charge_kwh": actual["charge"],
                "discharge_kwh": actual["discharge"],
                "emergency_kwh": actual["emergency"],
                "unused_energy_kwh": actual["unused"],
                "storage_start_kwh": e[:-1],
                "storage_end_kwh": e[1:],
            }
        )

        frame["planned_cost_yuan"] = (
            prices * plan["grid"]
        )
        frame["emergency_cost_yuan"] = (
            cfg["emergency_multiplier"]
            * prices
            * actual["emergency"]
        )
        frame["total_cost_yuan"] = (
            frame.planned_cost_yuan
            + frame.emergency_cost_yuan
        )
        frame["evaluation"] = (
            pd.Timestamp(date) >= EVALUATION_START
        )

        frames.append(frame)

        diary.append(
            {
                "date": str(pd.Timestamp(date).date()),
                "solver": plan["solver"],
                "history_first_date": (
                    None
                    if k == 0
                    else str(pd.Timestamp(dates[first]).date())
                ),
                "history_last_date": (
                    None
                    if k == 0
                    else str(pd.Timestamp(dates[k-1]).date())
                ),
                "history_days": (
                    0 if k == 0 else k - first
                ),
                "objective_surrogate_yuan": (
                    plan["objective"]
                ),
                "planning_constraint_residual": (
                    plan["plan_residual"]
                ),
                "initial_actual_kwh": float(e[0]),
                "terminal_actual_kwh": float(e[-1]),
            }
        )

        if progress and (
            (k + 1) % 30 == 0
            or k == len(dates) - 1
        ):
            print(
                f"{method}: {k+1}/{len(dates)} days",
                flush=True,
            )

    return (
        pd.concat(frames, ignore_index=True),
        pd.DataFrame(diary),
    )


# ============================================================
# 汇总与物理检查
# ============================================================

def daily_summary(dispatch):
    sumcols = [
        "planned_grid_kwh",
        "emergency_kwh",
        "planned_cost_yuan",
        "emergency_cost_yuan",
        "total_cost_yuan",
        "charge_kwh",
        "discharge_kwh",
        "unused_energy_kwh",
        "load_kwh",
        "pv_kwh",
    ]

    out = dispatch.groupby("date")[sumcols].sum()

    out["initial_storage_kwh"] = (
        dispatch.groupby("date").storage_start_kwh.first()
    )
    out["terminal_storage_kwh"] = (
        dispatch.groupby("date").storage_end_kwh.last()
    )
    out["emergency_slots"] = (
        dispatch.assign(
            event=dispatch.emergency_kwh > 1e-6
        )
        .groupby("date")
        .event.sum()
    )

    return out.reset_index()


def verify_physics(dispatch, cfg=None):
    cfg = CFG if cfg is None else cfg
    x = dispatch
    tol = cfg["tolerance"]

    def maximum(a):
        return float(np.max(np.abs(np.asarray(a))))

    residual = (
        x.planned_grid_kwh
        + x.pv_kwh
        + x.discharge_kwh
        + x.emergency_kwh
        - x.load_kwh
        - x.charge_kwh
        - x.unused_energy_kwh
    )

    evolution = (
        x.storage_end_kwh
        - x.storage_start_kwh
        - cfg["charge_efficiency"] * x.charge_kwh
        + x.discharge_kwh / cfg["discharge_efficiency"]
    )

    values = {
        "energy_balance_kwh": maximum(residual),
        "battery_update_kwh": maximum(evolution),
        "continuous_state_kwh": maximum(
            x.storage_start_kwh.to_numpy()[1:]
            - x.storage_end_kwh.to_numpy()[:-1]
        ),
        "soc_lower_violation_kwh": max(
            0.0,
            float(
                cfg["storage_min_kwh"]
                - x.storage_end_kwh.min()
            ),
        ),
        "soc_upper_violation_kwh": max(
            0.0,
            float(
                x.storage_end_kwh.max()
                - cfg["storage_max_kwh"]
            ),
        ),
        "charge_power_violation_kw": max(
            0.0,
            float(
                x.charge_kwh.max()
                * 60 / cfg["interval_minutes"]
                - cfg["max_charge_power_kw"]
            ),
        ),
        "discharge_power_violation_kw": max(
            0.0,
            float(
                x.discharge_kwh.max()
                * 60 / cfg["interval_minutes"]
                - cfg["max_discharge_power_kw"]
            ),
        ),
        "simultaneous_charge_discharge_slots": int(
            (
                (x.charge_kwh > tol)
                & (x.discharge_kwh > tol)
            ).sum()
        ),
        "planned_cost_error_yuan": maximum(
            x.planned_cost_yuan
            - x.price_yuan_per_kwh
            * x.planned_grid_kwh
        ),
        "emergency_cost_error_yuan": maximum(
            x.emergency_cost_yuan
            - cfg["emergency_multiplier"]
            * x.price_yuan_per_kwh
            * x.emergency_kwh
        ),
        "total_cost_error_yuan": maximum(
            x.total_cost_yuan
            - x.planned_cost_yuan
            - x.emergency_cost_yuan
        ),
        "negative_flow_violation_kwh": max(
            0.0,
            -float(
                x[
                    [
                        "planned_grid_kwh",
                        "charge_kwh",
                        "discharge_kwh",
                        "emergency_kwh",
                        "unused_energy_kwh",
                    ]
                ].min().min()
            ),
        ),
    }

    failures = {
        k: v
        for k, v in values.items()
        if v > tol
    }

    require(
        not failures,
        "Physical checks failed: " + str(failures),
    )
    require(
        np.isfinite(
            x.select_dtypes("number").to_numpy()
        ).all(),
        "Non-finite numeric output",
    )

    starts = pd.to_datetime(
        x.interval_start
    ).iloc[1:].to_numpy()
    ends = pd.to_datetime(
        x.interval_end
    ).iloc[:-1].to_numpy()

    require(
        np.array_equal(starts, ends),
        "Time axis is not continuous",
    )

    return values


def emergency_events(dispatch):
    """将连续的10分钟应急购电槽合并成题目要求的时间段。"""
    rows = []

    for day, x in dispatch.groupby("date", sort=True):
        x = x.sort_values("slot")
        a = x.emergency_kwh.to_numpy()
        t = 0

        while t < len(a):
            if a[t] <= 1e-6:
                t += 1
                continue

            start = t

            while t < len(a) and a[t] > 1e-6:
                t += 1

            rows.append(
                {
                    "date": day,
                    "interval": (
                        f"{clock(start*10)}-{clock(t*10)}"
                    ),
                    "emergency_kwh": float(
                        a[start:t].sum()
                    ),
                    "start_slot": start + 1,
                    "end_slot": t,
                }
            )

    return pd.DataFrame(
        rows,
        columns=[
            "date",
            "interval",
            "emergency_kwh",
            "start_slot",
            "end_slot",
        ],
    )


def validate_primary(
    dispatch,
    diary,
    load,
    pv,
    prices,
    dates,
    cfg,
):
    """最终提交前的核心独立检查。"""
    tol = cfg["tolerance"]

    require(
        len(dispatch) == 52560
        and dispatch.date.nunique() == 365,
        "Full-year trace incomplete",
    )

    ev = dispatch[dispatch.evaluation]

    require(
        len(ev) == 48096
        and ev.date.nunique() == 334,
        "Evaluation trace must contain 334 days × 144 slots",
    )

    require(
        np.array_equal(
            dispatch.price_yuan_per_kwh.to_numpy(),
            np.tile(prices, 365),
        ),
        "Attachment-1 prices changed",
    )

    require(
        np.allclose(
            dispatch.load_kwh.to_numpy(),
            load.ravel() / 6,
            rtol=0,
            atol=1e-10,
        ),
        "Load data changed",
    )
    require(
        np.allclose(
            dispatch.pv_kwh.to_numpy(),
            pv.ravel() / 6,
            rtol=0,
            atol=1e-10,
        ),
        "PV data changed",
    )

    close(
        dispatch.storage_start_kwh.iloc[0],
        cfg["initial_storage_kwh"],
        "Jan-1 initial storage",
        tol,
    )

    daily_start = (
        dispatch.groupby("date")
        .storage_start_kwh.first()
        .to_numpy()
    )
    daily_end = (
        dispatch.groupby("date")
        .storage_end_kwh.last()
        .to_numpy()
    )

    require(
        np.max(np.abs(daily_start[1:] - daily_end[:-1]))
        <= tol,
        "Day-boundary SOC is discontinuous",
    )

    hist_last = pd.to_datetime(
        diary.history_last_date.iloc[1:]
    )
    plan_dates = pd.to_datetime(diary.date.iloc[1:])

    require(
        bool((hist_last < plan_dates).all()),
        "Future-date information leakage detected",
    )
    require(
        bool((diary.history_days <= 28).all()),
        "History window exceeds 28 days",
    )

    close(
        dispatch.planned_cost_yuan.sum(),
        np.dot(
            dispatch.price_yuan_per_kwh,
            dispatch.planned_grid_kwh,
        ),
        "Fixed purchase settlement",
        1e-6,
    )

    close(
        dispatch.emergency_cost_yuan.sum(),
        cfg["emergency_multiplier"]
        * np.dot(
            dispatch.price_yuan_per_kwh,
            dispatch.emergency_kwh,
        ),
        "Emergency settlement",
        1e-6,
    )

    verify_physics(dispatch, cfg)

    # 每天计划时的初始SOC必须等于当天实测开始SOC。
    planned_start = (
        dispatch.groupby("date")
        .planned_storage_start_kwh.first()
        .to_numpy()
    )

    require(
        np.max(np.abs(planned_start - daily_start))
        <= tol,
        "Planning state differs from measured day-start SOC",
    )

    # 除冷启动日外，日前优化均以6000kWh为计划末态目标。
    planned_end = (
        dispatch.groupby("date")
        .planned_storage_end_kwh.last()
        .to_numpy()
    )
    require(
        np.max(
            np.abs(
                planned_end[1:]
                - cfg["planned_terminal_kwh"]
            )
        ) <= tol,
        "Planned terminal SOC target failed",
    )

    require(
        (
            diary.planning_constraint_residual
            <= tol
        ).all(),
        "Some daily planning problems violate constraints",
    )

    # 未来真实数据突变不允许改变某一更早计划日的历史情景。
    origin = 150
    s1, _ = history_scenarios(
        load,
        pv,
        origin,
        PRIMARY,
        cfg,
    )

    changed_load = load.copy()
    changed_pv = pv.copy()
    changed_load[origin:] *= 11
    changed_pv[origin:] = 0

    s2, _ = history_scenarios(
        changed_load,
        changed_pv,
        origin,
        PRIMARY,
        cfg,
    )

    require(
        np.array_equal(s1, s2),
        "Future actual mutation changed historical scenarios",
    )

    # 执行规则前72槽不应读取后72槽未来真实负载。
    base_load = load[100] / 6
    base_pv = pv[100] / 6

    original = execute_day(
        np.full(144, 600.0),
        base_load,
        base_pv,
        6000,
        cfg,
    )

    future_load = base_load.copy()
    future_load[72:] *= 10

    mutated = execute_day(
        np.full(144, 600.0),
        future_load,
        base_pv,
        6000,
        cfg,
    )

    for key in [
        "charge",
        "discharge",
        "emergency",
        "unused",
    ]:
        require(
            np.array_equal(
                original[key][:72],
                mutated[key][:72],
            ),
            f"Execution future lookup detected: {key}",
        )

    require(
        np.array_equal(
            original["storage"][:73],
            mutated["storage"][:73],
        ),
        "Execution future lookup detected: storage",
    )

    return True


# ============================================================
# Excel 模板导出：严格按照官方 result2.xlsx
# ============================================================

OFFICIAL_SHEETS = ["计划购电量", "充放电量", "紧急购电量"]

OFFICIAL_BATTERY_INTERVALS = [
    "0:00-4:00",
    "4:00-8:00",
    "8:00-12:00",
    "12:00-16:00",
    "16:00-20:00",
    "20:00-24:00",
]


def _header_text(value):
    if value is None:
        return ""
    return str(value).replace("\n", "").replace("\r", "").strip()


def _snapshot_row_style(ws, row, max_col):
    """保存官方模板某一行的样式，供扩展模板行时复制。"""
    return {
        "height": ws.row_dimensions[row].height,
        "cells": [copy(ws.cell(row, col)._style) for col in range(1, max_col + 1)],
    }


def _apply_row_style(ws, row, snapshot):
    if snapshot["height"] is not None:
        ws.row_dimensions[row].height = snapshot["height"]
    for col, style in enumerate(snapshot["cells"], start=1):
        ws.cell(row, col)._style = copy(style)


def _as_date_string(value):
    if value is None:
        return None
    try:
        return str(pd.Timestamp(value).date())
    except Exception:
        return str(value).strip()


def _time_text(value):
    """把 Excel 中的 0、time(0,0)、'24:00' 统一成可比较文本。"""
    if value is None:
        return None
    if isinstance(value, (int, float)) and abs(float(value)) < 1e-12:
        return "0:00"
    if hasattr(value, "hour") and hasattr(value, "minute"):
        return f"{int(value.hour)}:{int(value.minute):02d}"
    text = str(value).strip()
    if text in ("00:00", "0:00", "00:00:00", "0:00:00"):
        return "0:00"
    return text


def _official_plan_layout(ws):
    """
    官方模板固定布局：
      A       日期\时间
      B:EO    144个十分钟计划购电量
      EP      全天购电量
      EQ      全天购电费
    """
    require(ws.max_column >= 147, "计划购电量工作表列数不足")

    require(
        _header_text(ws.cell(1, 1).value) == r"日期\时间",
        "计划购电量!A1 必须保持官方标题“日期\\时间”",
    )

    interval_headers = [
        _header_text(ws.cell(1, col).value)
        for col in range(2, 146)
    ]

    require(len(interval_headers) == 144, "官方计划购电量必须有144个时段")
    require(
        interval_headers[0] == "0:10-0:20",
        f"官方模板首时段异常: {interval_headers[0]!r}",
    )
    require(
        interval_headers[-2] == "23:50-0:00+1",
        f"官方模板倒数第二时段异常: {interval_headers[-2]!r}",
    )
    require(
        interval_headers[-1] == "0:00-0:10+1",
        f"官方模板末时段异常: {interval_headers[-1]!r}",
    )
    require(
        _header_text(ws.cell(1, 146).value) == "全天计划购电量",
        "计划购电量!EP1 必须保持官方标题“全天计划购电量”",
    )
    require(
        _header_text(ws.cell(1, 147).value) == "全天计划购电费",
        "计划购电量!EQ1 必须保持官方标题“全天计划购电费”",
    )

    return 1, list(range(2, 146)), 146, 147


def _official_battery_layout(ws):
    expected = ["日期", "时间段", "充电量", "放电量", "时刻", "储电量"]
    actual = [_header_text(ws.cell(1, c).value) for c in range(1, 7)]
    require(actual == expected, f"充放电量表头被修改: {actual}")
    return tuple(range(1, 7))


def _official_emergency_layout(ws):
    expected = ["日期", "购电时间段", "购电量"]
    actual = [_header_text(ws.cell(1, c).value) for c in range(1, 4)]
    require(actual == expected, f"紧急购电量表头被修改: {actual}")
    return tuple(range(1, 4))


def build_submission_tables(primary_dispatch):
    """
    生成与官方 result2.xlsx 一一对应的数据。

    重要：
    当前内部 dispatch 的 slot 1..144 表示物理区间
      00:00-00:10, 00:10-00:20, ..., 23:50-24:00。
    官方模板的144列标签则是
      0:10-0:20, ..., 23:50-0:00+1, 0:00-0:10+1。
    因此按既定“此前10分钟平均功率”口径循环平移一格，仅改变展示位置，
    不改变同一自然日的全天计划购电量和计划购电费。
    """
    ev = primary_dispatch[primary_dispatch.evaluation].copy()

    require(
        ev.date.nunique() == 334 and len(ev) == 334 * 144,
        "提交期必须为 2025-02-01 至 2025-12-31，共334天×144时段",
    )

    dates = pd.date_range(EVALUATION_START, EVALUATION_END, freq="D")

    # ---------- 计划购电量 ----------
    plan_rows = []

    for date in dates:
        day = str(date.date())
        x = ev[ev.date == day].sort_values("slot")
        require(len(x) == 144, f"{day}: 缺少144个计划购电时段")

        physical_grid = x.planned_grid_kwh.to_numpy(float)

        # 官方模板首列是 0:10-0:20，因此物理 slot2 放首列；
        # 最后一列 0:00-0:10+1 放物理 slot1。
        official_grid = np.r_[physical_grid[1:], physical_grid[0]]

        plan_rows.append(
            {
                "date": day,
                "grid": official_grid,
                "total_grid": float(physical_grid.sum()),
                "cost": float(x.planned_cost_yuan.sum()),
            }
        )

    # ---------- 充放电量 ----------
    # 官方模板：每个日期固定6行；日期仅第一行填写。
    battery_rows = []

    for date in dates:
        day = str(date.date())
        x = ev[ev.date == day].sort_values("slot")
        require(len(x) == 144, f"{day}: 实际执行数据不完整")

        for block in range(6):
            z = x.iloc[block * 24:(block + 1) * 24]

            battery_rows.append(
                {
                    "date": day if block == 0 else None,
                    "interval": OFFICIAL_BATTERY_INTERVALS[block],
                    "charge_kwh": float(z.charge_kwh.sum()),
                    "discharge_kwh": float(z.discharge_kwh.sum()),
                    # 严格沿用官方模板：第一行0:00，第二行24:00。
                    "time": 0.0 if block == 0 else ("24:00" if block == 1 else None),
                    "storage_kwh": (
                        float(x.storage_start_kwh.iloc[0])
                        if block == 0
                        else (
                            float(x.storage_end_kwh.iloc[-1])
                            if block == 1
                            else None
                        )
                    ),
                }
            )

    # ---------- 紧急购电量 ----------
    # 官方表4：同一天多个事件时，日期只在第一条出现。
    # 模板示例为每个日期预留3行，因此每日至少保留3行；
    # 超过3个事件时自动扩展。
    raw_events = emergency_events(ev)
    emergency_rows = []

    for date in dates:
        day = str(date.date())
        events = raw_events[raw_events.date == day].reset_index(drop=True)
        row_count = max(3, len(events))

        for j in range(row_count):
            if j < len(events):
                interval = str(events.loc[j, "interval"])
                amount = float(events.loc[j, "emergency_kwh"])
            else:
                interval = None
                amount = None

            emergency_rows.append(
                {
                    "date": day if j == 0 else None,
                    "interval": interval,
                    "emergency_kwh": amount,
                    "is_last_for_day": j == row_count - 1,
                }
            )

    return plan_rows, battery_rows, emergency_rows


def export_result2(template_path, output_path, primary_dispatch):
    """从官方空模板生成最终 result2.xlsx，不修改官方表头和工作表名称。"""
    template_path = Path(template_path).resolve()
    output_path = Path(output_path).resolve()

    require(template_path.exists(), f"Missing result2 template: {template_path}")
    require(template_path != output_path, "禁止直接覆盖官方 result2.xlsx 空模板")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)

    wb = load_workbook(output_path)
    require(
        wb.sheetnames == OFFICIAL_SHEETS,
        f"官方 result2.xlsx 工作表应严格为 {OFFICIAL_SHEETS}，实际为 {wb.sheetnames}",
    )

    plan_rows, battery_rows, emergency_rows = build_submission_tables(primary_dispatch)

    # --------------------------------------------------------
    # 1) 计划购电量：官方模板本身已经有334个日期行
    # --------------------------------------------------------
    ws = wb["计划购电量"]
    date_col, interval_cols, total_col, cost_col = _official_plan_layout(ws)

    require(ws.max_row >= 335, "官方计划购电量模板缺少 334 个日期行")

    expected_dates = pd.date_range(EVALUATION_START, EVALUATION_END, freq="D")

    for i, (date, item) in enumerate(zip(expected_dates, plan_rows), start=2):
        # 日期由官方模板提供，不自行改写。
        require(
            _as_date_string(ws.cell(i, date_col).value) == str(date.date()),
            f"计划购电量第{i}行官方日期异常",
        )

        for value, col in zip(item["grid"], interval_cols):
            ws.cell(i, col).value = float(value)

        ws.cell(i, total_col).value = float(item["total_grid"])
        ws.cell(i, cost_col).value = float(item["cost"])

    # --------------------------------------------------------
    # 2) 充放电量：按官方6行/日的格式扩展至334天
    # --------------------------------------------------------
    ws = wb["充放电量"]
    _official_battery_layout(ws)

    # 先保存官方第一个完整6行日期块的样式。
    battery_styles = [
        _snapshot_row_style(ws, row, 6)
        for row in range(2, 8)
    ]

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for idx, item in enumerate(battery_rows):
        row = idx + 2
        block = idx % 6
        _apply_row_style(ws, row, battery_styles[block])

        if item["date"] is not None:
            ws.cell(row, 1).value = pd.Timestamp(item["date"]).to_pydatetime()
        else:
            ws.cell(row, 1).value = None

        ws.cell(row, 2).value = item["interval"]
        ws.cell(row, 3).value = float(item["charge_kwh"])
        ws.cell(row, 4).value = float(item["discharge_kwh"])
        ws.cell(row, 5).value = item["time"]
        ws.cell(row, 6).value = item["storage_kwh"]

    # --------------------------------------------------------
    # 3) 紧急购电量：日期仅每组第一行出现；无应急时留空
    # --------------------------------------------------------
    ws = wb["紧急购电量"]
    _official_emergency_layout(ws)

    emergency_styles = {
        "first": _snapshot_row_style(ws, 2, 3),
        "middle": _snapshot_row_style(ws, 3, 3),
        "last": _snapshot_row_style(ws, 4, 3),
    }

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    group_position = 0
    for idx, item in enumerate(emergency_rows):
        row = idx + 2

        # 根据当前日期组内位置选择官方示例行样式。
        if item["date"] is not None:
            group_position = 0

        if group_position == 0 and not item["is_last_for_day"]:
            style = emergency_styles["first"]
        elif item["is_last_for_day"]:
            style = emergency_styles["last"]
        else:
            style = emergency_styles["middle"]

        _apply_row_style(ws, row, style)

        if item["date"] is not None:
            ws.cell(row, 1).value = pd.Timestamp(item["date"]).to_pydatetime()
        else:
            ws.cell(row, 1).value = None

        ws.cell(row, 2).value = item["interval"]
        ws.cell(row, 3).value = item["emergency_kwh"]

        group_position += 1
        if item["is_last_for_day"]:
            group_position = 0

    wb.save(output_path)

    return {
        "status": "WRITTEN",
        "plan_rows": len(plan_rows),
        "battery_rows": len(battery_rows),
        "emergency_rows": len(emergency_rows),
        "template_mapping": "official shifted labels; same-day cyclic one-slot mapping",
    }


# ============================================================
# 保存后的 Excel 独立核验
# ============================================================

def verify_result2(output_path, template_path, primary_dispatch):
    output_path = Path(output_path).resolve()
    template_path = Path(template_path).resolve()

    require(output_path.exists(), f"Output workbook not found: {output_path}")

    saved = load_workbook(output_path, data_only=True, read_only=False)
    template = load_workbook(template_path, data_only=True, read_only=False)

    require(
        saved.sheetnames == OFFICIAL_SHEETS,
        f"保存后的工作表名称/顺序改变: {saved.sheetnames}",
    )
    require(
        template.sheetnames == OFFICIAL_SHEETS,
        "传入的模板不是官方 result2.xlsx 三工作表结构",
    )

    plan_rows, battery_rows, emergency_rows = build_submission_tables(primary_dispatch)
    tol = CFG["tolerance"]

    # ---------- 计划购电量 ----------
    ws = saved["计划购电量"]
    tws = template["计划购电量"]

    date_col, interval_cols, total_col, cost_col = _official_plan_layout(ws)
    _official_plan_layout(tws)

    # 官方表头必须逐格保持不变。
    for col in range(1, 148):
        require(
            ws.cell(1, col).value == tws.cell(1, col).value,
            f"计划购电量表头被修改: col={col}",
        )

    saved_grid_total = 0.0
    saved_cost_total = 0.0

    for i, expected in enumerate(plan_rows, start=2):
        require(
            _as_date_string(ws.cell(i, date_col).value) == expected["date"],
            f"计划购电量第{i}行日期错误",
        )
        require(
            _as_date_string(ws.cell(i, date_col).value)
            == _as_date_string(tws.cell(i, date_col).value),
            f"计划购电量第{i}行日期与官方模板不一致",
        )

        actual_grid = np.array(
            [float(ws.cell(i, col).value) for col in interval_cols],
            dtype=float,
        )

        require(
            np.allclose(actual_grid, expected["grid"], rtol=0, atol=tol),
            f"计划购电量第{i}行144个时段与模型映射不一致",
        )

        close(
            ws.cell(i, total_col).value,
            expected["total_grid"],
            f"计划购电量第{i}行全天计划购电量",
            tol,
        )
        close(
            ws.cell(i, cost_col).value,
            expected["cost"],
            f"计划购电量第{i}行全天计划购电费",
            tol,
        )

        saved_grid_total += float(actual_grid.sum())
        saved_cost_total += float(ws.cell(i, cost_col).value)

    # ---------- 充放电量 ----------
    ws = saved["充放电量"]
    tws = template["充放电量"]
    _official_battery_layout(ws)

    for col in range(1, 7):
        require(
            ws.cell(1, col).value == tws.cell(1, col).value,
            f"充放电量表头被修改: col={col}",
        )

    require(
        ws.max_row == 1 + len(battery_rows),
        f"充放电量行数错误: {ws.max_row}",
    )

    saved_charge = 0.0
    saved_discharge = 0.0

    for i, expected in enumerate(battery_rows, start=2):
        actual_date = _as_date_string(ws.cell(i, 1).value)
        require(actual_date == expected["date"], f"充放电量第{i}行日期排版错误")

        require(
            _header_text(ws.cell(i, 2).value) == expected["interval"],
            f"充放电量第{i}行时间段错误",
        )
        close(ws.cell(i, 3).value, expected["charge_kwh"], f"充放电量第{i}行充电量", tol)
        close(ws.cell(i, 4).value, expected["discharge_kwh"], f"充放电量第{i}行放电量", tol)

        expected_time = None if expected["time"] is None else (
            "0:00" if expected["time"] == 0.0 else "24:00"
        )
        require(
            _time_text(ws.cell(i, 5).value) == expected_time,
            f"充放电量第{i}行时刻错误",
        )

        if expected["storage_kwh"] is None:
            require(
                ws.cell(i, 6).value is None,
                f"充放电量第{i}行不应填写储电量",
            )
        else:
            close(
                ws.cell(i, 6).value,
                expected["storage_kwh"],
                f"充放电量第{i}行储电量",
                tol,
            )

        saved_charge += float(ws.cell(i, 3).value)
        saved_discharge += float(ws.cell(i, 4).value)

    # ---------- 紧急购电量 ----------
    ws = saved["紧急购电量"]
    tws = template["紧急购电量"]
    _official_emergency_layout(ws)

    for col in range(1, 4):
        require(
            ws.cell(1, col).value == tws.cell(1, col).value,
            f"紧急购电量表头被修改: col={col}",
        )

    require(
        ws.max_row == 1 + len(emergency_rows),
        f"紧急购电量行数错误: {ws.max_row}",
    )

    saved_emergency = 0.0

    for i, expected in enumerate(emergency_rows, start=2):
        require(
            _as_date_string(ws.cell(i, 1).value) == expected["date"],
            f"紧急购电量第{i}行日期排版错误",
        )

        actual_interval = ws.cell(i, 2).value
        actual_amount = ws.cell(i, 3).value

        if expected["interval"] is None:
            require(actual_interval is None, f"紧急购电量第{i}行应为空时间段")
            require(actual_amount is None, f"紧急购电量第{i}行应为空购电量")
        else:
            require(
                _header_text(actual_interval) == expected["interval"],
                f"紧急购电量第{i}行时间段错误",
            )
            close(
                actual_amount,
                expected["emergency_kwh"],
                f"紧急购电量第{i}行购电量",
                tol,
            )
            saved_emergency += float(actual_amount)

    ev = primary_dispatch[primary_dispatch.evaluation]

    # 循环平移只改变列位置，不得改变总量。
    close(
        saved_grid_total,
        ev.planned_grid_kwh.sum(),
        "最终Excel全天计划购电总量",
        1e-5,
    )
    close(
        saved_cost_total,
        ev.planned_cost_yuan.sum(),
        "最终Excel计划购电总费用",
        1e-5,
    )
    close(
        saved_charge,
        ev.charge_kwh.sum(),
        "最终Excel充电总量",
        1e-5,
    )
    close(
        saved_discharge,
        ev.discharge_kwh.sum(),
        "最终Excel放电总量",
        1e-5,
    )
    close(
        saved_emergency,
        ev.emergency_kwh.sum(),
        "最终Excel紧急购电总量",
        1e-5,
    )

    return {
        "status": "PASS",
        "evaluation_days": 334,
        "plan_rows": len(plan_rows),
        "battery_rows": len(battery_rows),
        "emergency_rows": len(emergency_rows),
        "planned_purchase_kwh": float(ev.planned_grid_kwh.sum()),
        "emergency_purchase_kwh": float(ev.emergency_kwh.sum()),
        "planned_cost_yuan": float(ev.planned_cost_yuan.sum()),
        "emergency_cost_yuan": float(ev.emergency_cost_yuan.sum()),
        "total_cost_yuan": float(ev.total_cost_yuan.sum()),
        "template_mapping": "official shifted labels; same-day cyclic one-slot mapping",
    }


# ============================================================
# 可选策略比较（仅终端显示）
# ============================================================

def compare_strategies(results):
    rows = []

    for name, (allx, _) in results.items():
        x = allx[allx.evaluation]
        d = daily_summary(x)

        rows.append(
            {
                "strategy": name,
                "days": len(d),
                "planned_grid_kwh": float(
                    x.planned_grid_kwh.sum()
                ),
                "emergency_kwh": float(
                    x.emergency_kwh.sum()
                ),
                "planned_cost_yuan": float(
                    x.planned_cost_yuan.sum()
                ),
                "emergency_cost_yuan": float(
                    x.emergency_cost_yuan.sum()
                ),
                "total_cost_yuan": float(
                    x.total_cost_yuan.sum()
                ),
                "emergency_days": int(
                    (d.emergency_slots > 0).sum()
                ),
            }
        )

    return pd.DataFrame(rows)


# ============================================================
# 主程序
# ============================================================

def main():
    parser = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawTextHelpFormatter,
    )

    parser.add_argument(
        "--actual",
        type=Path,
        default=DEFAULT_ACTUAL,
        help="全年实际负载/光伏 CSV",
    )
    parser.add_argument(
        "--price",
        type=Path,
        default=DEFAULT_PRICE,
        help="144个固定电价 CSV",
    )
    parser.add_argument(
        "--template",
        type=Path,
        default=None,
        help="官方 result2.xlsx；默认自动查找附件5",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="最终 result2.xlsx 输出位置",
    )
    parser.add_argument(
        "--compare",
        action="store_true",
        help="额外运行两个对照策略并仅在终端打印比较",
    )
    parser.add_argument(
        "--quiet",
        action="store_true",
        help="关闭逐月进度显示",
    )

    args = parser.parse_args()

    started = time.perf_counter()

    actual_path = args.actual.resolve()
    price_path = args.price.resolve()
    output_path = args.output.resolve()

    template_path = (
        args.template.resolve()
        if args.template is not None
        else default_template_path().resolve()
    )

    require(
        actual_path.exists(),
        f"Missing actual input: {actual_path}",
    )
    require(
        price_path.exists(),
        f"Missing price input: {price_path}",
    )
    require(
        template_path.exists(),
        f"Missing result2 template: {template_path}",
    )

    print("Loading Q2 inputs...")
    load, pv, prices, dates = load_inputs(
        actual_path,
        price_path,
    )

    print(
        "Running primary strategy: "
        "28-day historical SAA + battery"
    )

    primary_dispatch, primary_diary = run_strategy(
        load,
        pv,
        prices,
        dates,
        PRIMARY,
        CFG,
        progress=not args.quiet,
    )

    print("Checking physical and causal constraints...")

    validate_primary(
        primary_dispatch,
        primary_diary,
        load,
        pv,
        prices,
        dates,
        CFG,
    )

    print("Writing official result2.xlsx...")

    export_result2(
        template_path,
        output_path,
        primary_dispatch,
    )

    print("Re-opening and verifying saved workbook...")

    check = verify_result2(
        output_path,
        template_path,
        primary_dispatch,
    )

    if args.compare:
        print("\nRunning optional comparison strategies...")

        results = {
            PRIMARY: (
                primary_dispatch,
                primary_diary,
            )
        }

        for method in METHODS:
            if method == PRIMARY:
                continue

            results[method] = run_strategy(
                load,
                pv,
                prices,
                dates,
                method,
                CFG,
                progress=not args.quiet,
            )

        comparison = compare_strategies(results)

        print("\nStrategy comparison:")
        print(
            comparison.to_string(
                index=False,
                float_format=lambda v: f"{v:.4f}",
            )
        )

    elapsed = time.perf_counter() - started

    print("\n" + "=" * 68)
    print("Q2 COMPLETED: PASS")
    print("Primary strategy: 28-day historical SAA + battery")
    print(
        f"Evaluation: 2025-02-01 to 2025-12-31 "
        f"({check['evaluation_days']} days)"
    )
    print(
        f"Planned purchase: "
        f"{check['planned_purchase_kwh']:.4f} kWh"
    )
    print(
        f"Emergency purchase: "
        f"{check['emergency_purchase_kwh']:.4f} kWh"
    )
    print(
        f"Planned cost: "
        f"{check['planned_cost_yuan']:.4f} yuan"
    )
    print(
        f"Emergency cost: "
        f"{check['emergency_cost_yuan']:.4f} yuan"
    )
    print(
        f"Total cost: "
        f"{check['total_cost_yuan']:.4f} yuan"
    )
    print("Physical/causal checks: PASS")
    print("Saved-XLSX verification: PASS")
    print("Official template layout/header/date formatting: PRESERVED")
    print("Plan time columns: official 0:10...0:10+1 cyclic mapping")
    print(f"Output: {output_path}")
    print("Generated artifacts: result2.xlsx only")
    print(f"Elapsed: {elapsed:.1f} seconds")
    print("=" * 68)


if __name__ == "__main__":
    main()
