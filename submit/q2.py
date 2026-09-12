from __future__ import annotations

from pathlib import Path
from copy import copy
import argparse
import json
import time

import numpy as np
import pandas as pd
from scipy.optimize import linprog, Bounds, LinearConstraint, milp
from scipy.sparse import coo_matrix, csr_matrix, hstack
from sklearn.linear_model import Ridge
from sklearn.preprocessing import StandardScaler
from sklearn.pipeline import make_pipeline

SCRIPT_DIR = Path(__file__).resolve().parent

# ============================================================
# Q2 CONFIGURATION
# ============================================================
FORECAST = {
    "window_days": 60,          # maximum rolling window; warm-up uses all available history
    "k_load": 3,
    "k_pv": 2,
    "alpha_load": 0.3,
    "alpha_pv": 10.0,
    "scenario_days": 28,        # maximum paired residual scenario window
}
Q_GRID = np.round(np.arange(0.05, 0.951, 0.05), 2)
DEV_START, DEV_END = "2025-05-01", "2025-08-31"
RUN_START, RUN_END = "2025-02-01", "2025-12-31"

CFG = {
    "eta_c": 0.9,
    "eta_d": 0.9,
    "storage_min_kwh": 1200.0,
    "storage_max_kwh": 10800.0,
    "initial_storage_kwh": 6000.0,
    "max_charge_kw": 5000.0,
    "max_discharge_kw": 5000.0,
    "emergency_multiplier": 5.0,
    "interval_hours": 1.0 / 6.0,
}

def require(cond, msg):
    if not cond:
        raise AssertionError(msg)

def resolve_existing(candidates):
    for p in candidates:
        p = Path(p)
        if p.exists():
            return p
    return Path(candidates[0])

def default_paths():
    actual = resolve_existing([
        SCRIPT_DIR / "Data" / "附件2.xlsx",
        SCRIPT_DIR / "input" / "q2_actual_load_pv.csv",
        SCRIPT_DIR / "q2_actual_load_pv.csv",
        SCRIPT_DIR.parent / "q2" / "input" / "q2_actual_load_pv.csv",
    ])
    price = resolve_existing([
        SCRIPT_DIR / "Data" / "附件1.xlsx",
        SCRIPT_DIR / "input" / "q2_fixed_prices.csv",
        SCRIPT_DIR / "q2_fixed_prices.csv",
        SCRIPT_DIR.parent / "q2" / "input" / "q2_fixed_prices.csv",
    ])
    template = resolve_existing([
        SCRIPT_DIR / "Data" / "附件5" / "result2.xlsx",
        SCRIPT_DIR / "input" / "result2_template.xlsx",
        SCRIPT_DIR / "result2_template.xlsx",
        SCRIPT_DIR.parent / "Data_preprocessed" / "raw" / "附件5" / "result2.xlsx",
    ])
    return actual, price, template

def _read_raw_xlsx_inputs(actual_path, price_path):
    """Read the original competition workbooks without an intermediate CSV."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError(
            "读取原始 Excel 需要 openpyxl；请安装 openpyxl>=3.1"
        ) from e

    def sheet_rows(path, sheet_name):
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet_name not in wb.sheetnames:
                raise ValueError(
                    f"{path.name} 缺少工作表 {sheet_name!r}，实际工作表: {wb.sheetnames}"
                )
            return [list(row) for row in wb[sheet_name].iter_rows(values_only=True)]
        finally:
            wb.close()

    def daily_sheet(path, sheet_name, label):
        rows = sheet_rows(path, sheet_name)
        require(len(rows) == 366 and all(len(r) >= 145 for r in rows),
                f"{path.name}/{sheet_name} 应为 365 天×144 时段，且首列为日期")
        dates = pd.DatetimeIndex([pd.Timestamp(r[0]) for r in rows[1:]])
        require(dates.equals(pd.date_range("2025-01-01", "2025-12-31", freq="D")),
                f"{label} 日期缺失、重复或顺序错误")
        values = np.asarray([r[1:145] for r in rows[1:]], dtype=float)
        require(np.isfinite(values).all(), f"{label} 存在缺失或非有限值")
        require((values >= 0).all(), f"{label} 存在负值")
        return dates, values

    dates_load, load = daily_sheet(actual_path, "小区负载", "小区负载")
    dates_pv, pv = daily_sheet(actual_path, "光伏发电实际功率", "光伏实际功率")
    require(dates_load.equals(dates_pv), "负载和光伏日期轴不一致")

    rows = sheet_rows(price_path, "Sheet1")
    require(len(rows) >= 145 and all(len(r) >= 2 for r in rows[:145]),
            f"{price_path.name}/Sheet1 应包含 144 个电价记录")
    prices = np.asarray([r[1] for r in rows[1:145]], dtype=float)
    require(np.isfinite(prices).all() and (prices > 0).all(), "电价数据非法")
    return dates_load, load, pv, prices


def load_inputs(actual_path, price_path):
    if Path(actual_path).suffix.lower() in {".xlsx", ".xlsm"}:
        return _read_raw_xlsx_inputs(Path(actual_path), Path(price_path))

    x = pd.read_csv(actual_path, encoding="utf-8-sig")
    x["date"] = pd.to_datetime(x["date"])
    x = x.sort_values(["date", "slot"]).reset_index(drop=True)

    require(len(x) == 365 * 144, "q2_actual_load_pv.csv 必须是 365×144=52560 行")
    require(x[["load_kw", "pv_kw"]].notna().all().all(), "Load/PV 存在缺失值")
    require((x[["load_kw", "pv_kw"]] >= 0).all().all(), "Load/PV 存在负值")
    require(x[["date", "slot"]].duplicated().sum() == 0, "存在重复 date-slot")
    counts = x.groupby("date")["slot"].nunique()
    require((counts == 144).all(), "存在某日不是 144 个时段")

    dates = pd.DatetimeIndex(x["date"].drop_duplicates())
    load = x["load_kw"].to_numpy(float).reshape(-1, 144)
    pv = x["pv_kw"].to_numpy(float).reshape(-1, 144)

    p = pd.read_csv(price_path, encoding="utf-8-sig").sort_values("slot")
    require(len(p) == 144, "q2_fixed_prices.csv 必须为 144 行")
    prices = p["price_yuan_per_kwh"].to_numpy(float)
    require(np.isfinite(prices).all() and (prices > 0).all(), "电价数据非法")
    return dates, load, pv, prices

def pca_basis(curves, k):
    mu = curves.mean(axis=0)
    _, _, vt = np.linalg.svd(curves - mu, full_matrices=False)
    return mu, vt[:k]

def weekday7(date):
    x = np.zeros(7)
    x[date.weekday()] = 1.0
    return x

def annual1(date):
    a = 2.0 * np.pi * date.dayofyear / 365.25
    return np.array([np.sin(a), np.cos(a)])

def forecast_scenarios(dates, load, pv):
    """
    Causal rolling forecast.

    Important:
    - W=60 is the maximum history length, not available on Feb-01.
    - During warm-up, all available completed history is used.
    - Residual scenario bank uses at most the latest 28 completed days.
    - Today's actual Load/PV is appended to the residual bank only AFTER today's
      scenario is constructed, so there is no same-day look-ahead leakage.
    """
    W = FORECAST["window_days"]
    KL = FORECAST["k_load"]
    KP = FORECAST["k_pv"]
    AL = FORECAST["alpha_load"]
    AP = FORECAST["alpha_pv"]
    S = FORECAST["scenario_days"]

    pred_l = np.full_like(load, np.nan)
    pred_p = np.full_like(pv, np.nan)
    scenarios = {}
    residuals = []

    # Need at least 7 lag days + a small training set.
    for d in range(14, len(dates)):
        h0 = max(0, d - W)

        mu_l, phi_l = pca_basis(load[h0:d], KL)
        mu_p, phi_p = pca_basis(pv[h0:d], KP)

        sl = (load[h0:d] - mu_l) @ phi_l.T
        sp = (pv[h0:d] - mu_p) @ phi_p.T

        train_days = np.arange(max(h0 + 7, 7), d)

        def f_load(j):
            jj = j - h0
            return np.r_[weekday7(dates[j]), sl[jj - 1]]

        def f_pv(j):
            jj = j - h0
            return np.r_[
                annual1(dates[j]),
                sp[jj - 1],
                sp[jj - 7],
                sp[jj - 7:jj].mean(axis=0),
            ]

        XL = np.vstack([f_load(j) for j in train_days])
        XP = np.vstack([f_pv(j) for j in train_days])
        yL = np.vstack([sl[j - h0] for j in train_days])
        yP = np.vstack([sp[j - h0] for j in train_days])

        model_l = make_pipeline(StandardScaler(), Ridge(alpha=AL))
        model_p = make_pipeline(StandardScaler(), Ridge(alpha=AP))
        model_l.fit(XL, yL)
        model_p.fit(XP, yP)

        sh_l = np.atleast_1d(model_l.predict(f_load(d)[None, :])[0])
        sh_p = np.atleast_1d(model_p.predict(f_pv(d)[None, :])[0])

        pred_l[d] = np.maximum(mu_l + sh_l @ phi_l, 0)
        pred_p[d] = np.maximum(mu_p + sh_p @ phi_p, 0)

        if residuals:
            recent = residuals[-S:]
            eL = np.array([r[1] for r in recent])
            eP = np.array([r[2] for r in recent])
            # Keep same-day Load/PV residual pairing; clip physical Load/PV to nonnegative.
            scenarios[d] = (
                np.maximum(pred_l[d][None, :] + eL, 0)
                - np.maximum(pred_p[d][None, :] + eP, 0)
            )

        residuals.append((d, load[d] - pred_l[d], pv[d] - pred_p[d]))

    return pred_l, pred_p, scenarios

def solve_dispatch(
    net_kwh,
    prices,
    initial_storage_kwh=None,
    terminal_target_kwh=None,
    terminal_penalty_yuan_per_kwh=0.0,
):
    """Daily dispatch LP with a carried storage state; MILP removes c/d overlap."""
    net = np.asarray(net_kwh, float)
    n = len(net)
    require(n == 144, "每日必须 144 个时段")

    eta_c, eta_d = CFG["eta_c"], CFG["eta_d"]
    cmax = CFG["max_charge_kw"] * CFG["interval_hours"]
    dmax = CFG["max_discharge_kw"] * CFG["interval_hours"]
    emin, emax = CFG["storage_min_kwh"], CFG["storage_max_kwh"]

    # x = [grid, charge, discharge, unused, E_0...E_144, dev_plus, dev_minus]
    # The last two variables linearize the absolute terminal-state deviation.
    base_count = 5 * n + 1
    dev_plus, dev_minus = base_count, base_count + 1
    count = base_count + 2
    obj = np.zeros(count)
    obj[:n] = prices
    if terminal_target_kwh is not None:
        require(terminal_penalty_yuan_per_kwh >= 0, "终点偏离惩罚系数不能为负")
        obj[dev_plus] = terminal_penalty_yuan_per_kwh
        obj[dev_minus] = terminal_penalty_yuan_per_kwh

    rows, cols, data = [], [], []
    rhs = np.zeros(2 * n)
    t = np.arange(n)

    # planned energy balance: g + d - c - u = net
    for rr, cc, dd in [
        (t, t, np.ones(n)),
        (t, 2*n+t, np.ones(n)),
        (t, n+t, -np.ones(n)),
        (t, 3*n+t, -np.ones(n)),
    ]:
        rows.extend(rr); cols.extend(cc); data.extend(dd)
    rhs[:n] = net

    # storage recursion
    r = n + t
    for rr, cc, dd in [
        (r, 4*n+t+1, np.ones(n)),
        (r, 4*n+t, -np.ones(n)),
        (r, n+t, -eta_c*np.ones(n)),
        (r, 2*n+t, np.ones(n)/eta_d),
    ]:
        rows.extend(rr); cols.extend(cc); data.extend(dd)

    if terminal_target_kwh is not None:
        rows.extend([2*n, 2*n])
        cols.extend([5*n, dev_plus])
        data.extend([1.0, -1.0])
        rows.append(2*n)
        cols.append(dev_minus)
        data.append(1.0)
        rhs = np.r_[rhs, float(terminal_target_kwh)]

    Aeq = coo_matrix((data, (rows, cols)), shape=(len(rhs), count)).tocsr()

    lower = np.zeros(count)
    upper = np.full(count, np.inf)
    upper[n:2*n] = cmax
    upper[2*n:3*n] = dmax
    lower[4*n:] = emin
    upper[4*n:] = emax
    lower[dev_plus:] = 0.0
    upper[dev_plus:] = np.inf
    if initial_storage_kwh is None:
        initial_storage_kwh = CFG["initial_storage_kwh"]
    require(
        CFG["storage_min_kwh"] <= initial_storage_kwh <= CFG["storage_max_kwh"],
        "初始储能电量超出允许范围",
    )
    lower[4*n] = upper[4*n] = initial_storage_kwh

    res = linprog(
        obj,
        A_eq=Aeq,
        b_eq=rhs,
        bounds=np.column_stack([lower, upper]),
        method="highs",
    )
    require(res.success, "LP 求解失败: " + res.message)
    x = res.x
    solver = "LP"

    overlap = (x[n:2*n] > 1e-7) & (x[2*n:3*n] > 1e-7)
    if overlap.any():
        # binary z_t: c_t <= cmax*z_t, d_t <= dmax*(1-z_t)
        m = count + n
        mode_rows = np.r_[t, t, n+t, n+t]
        mode_cols = np.r_[n+t, count+t, 2*n+t, count+t]
        mode_data = np.r_[
            np.ones(n), -cmax*np.ones(n),
            np.ones(n), dmax*np.ones(n),
        ]
        mode = coo_matrix((mode_data, (mode_rows, mode_cols)), shape=(2*n, m)).tocsr()

        res2 = milp(
            np.r_[obj, np.zeros(n)],
            integrality=np.r_[np.zeros(count), np.ones(n)],
            bounds=Bounds(
                np.r_[lower, np.zeros(n)],
                np.r_[upper, np.ones(n)],
            ),
            constraints=[
                LinearConstraint(
                    hstack([Aeq, csr_matrix((len(rhs), n))]),
                    rhs, rhs,
                ),
                LinearConstraint(
                    mode,
                    -np.inf,
                    np.r_[np.zeros(n), dmax*np.ones(n)],
                ),
            ],
            options={"mip_rel_gap": 1e-8},
        )
        require(res2.success, "MILP 求解失败: " + res2.message)
        x = res2.x[:count]
        solver = "MILP"

    return {
        "grid": np.maximum(x[:n], 0),
        "charge": np.maximum(x[n:2*n], 0),
        "discharge": np.maximum(x[2*n:3*n], 0),
        "unused_plan": np.maximum(x[3*n:4*n], 0),
        "storage": x[4*n:5*n+1],
        "solver": solver,
        "max_plan_balance_residual": float(np.max(np.abs(Aeq @ x - rhs))),
    }

def execute_fixed_plan(plan, actual_net_kwh):
    supply = plan["grid"] + plan["discharge"] - plan["charge"]
    emergency = np.maximum(actual_net_kwh - supply, 0)
    unused = np.maximum(supply - actual_net_kwh, 0)
    balance = (
        plan["grid"] + plan["discharge"] + emergency
        - actual_net_kwh - plan["charge"] - unused
    )
    require(np.max(np.abs(balance)) < 1e-6, "实际执行能量平衡失败")
    return emergency, unused

def run_q(
    dates,
    load,
    pv,
    prices,
    scenarios,
    day_idx,
    q,
    keep_slots=False,
    terminal_target_kwh=None,
    terminal_penalty_yuan_per_kwh=0.0,
):
    daily, slots = [], []
    storage_kwh = CFG["initial_storage_kwh"]

    for d in day_idx:
        qnet_kw = np.quantile(scenarios[d], q, axis=0)
        plan = solve_dispatch(
            qnet_kw * CFG["interval_hours"], prices,
            initial_storage_kwh=storage_kwh,
            terminal_target_kwh=terminal_target_kwh,
            terminal_penalty_yuan_per_kwh=terminal_penalty_yuan_per_kwh,
        )
        storage_kwh = float(plan["storage"][-1])

        actual_net_kwh = (load[d] - pv[d]) * CFG["interval_hours"]
        emergency, unused = execute_fixed_plan(plan, actual_net_kwh)

        planned_cost = float(prices @ plan["grid"])
        emergency_cost = float(
            CFG["emergency_multiplier"] * prices @ emergency
        )

        daily.append({
            "q": float(q),
            "date": str(dates[d].date()),
            "planned_grid_kwh": float(plan["grid"].sum()),
            "planned_charge_kwh": float(plan["charge"].sum()),
            "planned_discharge_kwh": float(plan["discharge"].sum()),
            "emergency_kwh": float(emergency.sum()),
            "unused_kwh": float(unused.sum()),
            "planned_cost_yuan": planned_cost,
            "emergency_cost_yuan": emergency_cost,
            "total_cost_yuan": planned_cost + emergency_cost,
            "emergency_slots": int((emergency > 1e-7).sum()),
            "solver": plan["solver"],
            "max_balance_residual": plan["max_plan_balance_residual"],
            "min_storage_kwh": float(plan["storage"].min()),
            "max_storage_kwh": float(plan["storage"].max()),
        })

        if keep_slots:
            for t in range(144):
                slots.append({
                    "date": str(dates[d].date()),
                    "slot": t + 1,
                    "forecast_net_q85_kw": float(qnet_kw[t]),
                    "actual_load_kw": float(load[d, t]),
                    "actual_pv_kw": float(pv[d, t]),
                    "actual_net_kwh": float(actual_net_kwh[t]),
                    "price_yuan_per_kwh": float(prices[t]),
                    "planned_grid_kwh": float(plan["grid"][t]),
                    "planned_charge_kwh": float(plan["charge"][t]),
                    "planned_discharge_kwh": float(plan["discharge"][t]),
                    "planned_storage_start_kwh": float(plan["storage"][t]),
                    "planned_storage_end_kwh": float(plan["storage"][t+1]),
                    "emergency_kwh": float(emergency[t]),
                    "unused_kwh": float(unused[t]),
                })

    return pd.DataFrame(daily), pd.DataFrame(slots)

def emergency_events(slot_df):
    """Merge consecutive 10-minute emergency-purchase slots."""
    rows = []

    def clock(minute):
        if minute == 1440:
            return "24:00"
        return f"{(minute // 60) % 24}:{minute % 60:02d}"

    for date, x in slot_df.groupby("date", sort=True):
        a = x.sort_values("slot")["emergency_kwh"].to_numpy()
        t = 0
        while t < 144:
            if a[t] <= 1e-7:
                t += 1
                continue
            start = t
            while t < 144 and a[t] > 1e-7:
                t += 1
            rows.append({
                "date": date,
                "interval": f"{clock(start*10)}-{clock(t*10)}",
                "emergency_kwh": float(a[start:t].sum()),
                "start_slot": start + 1,
                "end_slot": t,
            })
    return pd.DataFrame(rows)

def scan_q_star(
    dates,
    load,
    pv,
    prices,
    scenarios,
    terminal_target_kwh=None,
    terminal_penalty_yuan_per_kwh=0.0,
):
    dev = np.where(
        (dates >= pd.Timestamp(DEV_START))
        & (dates <= pd.Timestamp(DEV_END))
    )[0]
    require(all(d in scenarios for d in dev), "开发期场景缺失")

    rows = []
    for q in Q_GRID:
        ddf, _ = run_q(
            dates, load, pv, prices, scenarios, dev, float(q), False
            , terminal_target_kwh, terminal_penalty_yuan_per_kwh
        )
        rows.append({
            "q": float(q),
            "planned_cost_yuan": float(ddf["planned_cost_yuan"].sum()),
            "emergency_cost_yuan": float(ddf["emergency_cost_yuan"].sum()),
            "total_cost_yuan": float(ddf["total_cost_yuan"].sum()),
            "emergency_kwh": float(ddf["emergency_kwh"].sum()),
            "unused_kwh": float(ddf["unused_kwh"].sum()),
            "emergency_days": int((ddf["emergency_kwh"] > 1e-7).sum()),
        })

    scan = pd.DataFrame(rows)
    q_star = float(scan.loc[scan["total_cost_yuan"].idxmin(), "q"])
    scan["delta_cost_yuan"] = (
        scan["total_cost_yuan"] - scan["total_cost_yuan"].min()
    )
    scan["delta_pct"] = (
        scan["delta_cost_yuan"] / scan["total_cost_yuan"].min() * 100
    )
    return q_star, scan

def export_result2(template_path, output_path, daily, slots, events):
    """
    Fill the official result2.xlsx template.

    Official plan-sheet time order is:
        [slot2, slot3, ..., slot144, slot1]
    i.e. 0:10-0:20 ... 23:50-0:00+1, 0:00-0:10+1.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError(
            "生成 result2.xlsx 需要 openpyxl；竞赛环境中请安装 openpyxl>=3.1"
        ) from e

    wb = load_workbook(template_path)
    require(
        set(["计划购电量", "充放电量", "紧急购电量"]).issubset(wb.sheetnames),
        "result2 模板缺少规定工作表",
    )

    ws_plan = wb["计划购电量"]
    ws_bat = wb["充放电量"]
    ws_em = wb["紧急购电量"]

    dates = sorted(daily["date"].unique())
    by_date = {
        d: slots[slots["date"] == d].sort_values("slot").reset_index(drop=True)
        for d in dates
    }
    daily_map = daily.set_index("date")

    # -------- 计划购电量 --------
    require(len(dates) == 334, "最终结果必须包含 334 天")
    for i, d in enumerate(dates, start=2):
        g = by_date[d]["planned_grid_kwh"].to_numpy(float)
        require(len(g) == 144, f"{d} 不是 144 个时段")

        # Date cell
        ws_plan.cell(i, 1).value = pd.Timestamp(d).to_pydatetime()

        # B:EO = slot2...slot144,slot1
        rotated = np.r_[g[1:], g[:1]]
        for j, value in enumerate(rotated, start=2):
            ws_plan.cell(i, j).value = float(value)

        ws_plan.cell(i, 146).value = float(
            daily_map.loc[d, "planned_grid_kwh"]
        )
        ws_plan.cell(i, 147).value = float(
            daily_map.loc[d, "planned_cost_yuan"]
        )

    # -------- 充放电量 --------
    # Preserve prototype styles before deleting sample rows.
    prototype = []
    for r in range(2, 8):
        row_style = []
        for c in range(1, 7):
            cell = ws_bat.cell(r, c)
            row_style.append({
                "style": copy(cell._style),
                "number_format": cell.number_format,
                "alignment": copy(cell.alignment),
                "font": copy(cell.font),
                "fill": copy(cell.fill),
                "border": copy(cell.border),
                "protection": copy(cell.protection),
            })
        prototype.append(row_style)

    if ws_bat.max_row >= 2:
        ws_bat.delete_rows(2, ws_bat.max_row - 1)

    intervals = [
        "0:00-4:00", "4:00-8:00", "8:00-12:00",
        "12:00-16:00", "16:00-20:00", "20:00-24:00",
    ]

    out_row = 2
    for d in dates:
        x = by_date[d]
        day_start = float(x.iloc[0]["planned_storage_start_kwh"])
        day_end = float(x.iloc[-1]["planned_storage_end_kwh"])

        for b in range(6):
            z = x.iloc[b*24:(b+1)*24]
            values = [
                pd.Timestamp(d).to_pydatetime() if b == 0 else None,
                intervals[b],
                float(z["planned_charge_kwh"].sum()),
                float(z["planned_discharge_kwh"].sum()),
                "0:00" if b == 0 else ("24:00" if b == 1 else None),
                day_start if b == 0 else (day_end if b == 1 else None),
            ]

            for c, value in enumerate(values, start=1):
                cell = ws_bat.cell(out_row, c)
                cell.value = value
                st = prototype[b][c-1]
                cell._style = copy(st["style"])
                cell.number_format = st["number_format"]
                cell.alignment = copy(st["alignment"])
                cell.font = copy(st["font"])
                cell.fill = copy(st["fill"])
                cell.border = copy(st["border"])
                cell.protection = copy(st["protection"])

            out_row += 1

    # -------- 紧急购电量 --------
    em_style = []
    for c in range(1, 4):
        cell = ws_em.cell(2, c)
        em_style.append({
            "style": copy(cell._style),
            "number_format": cell.number_format,
            "alignment": copy(cell.alignment),
            "font": copy(cell.font),
            "fill": copy(cell.fill),
            "border": copy(cell.border),
            "protection": copy(cell.protection),
        })

    if ws_em.max_row >= 2:
        ws_em.delete_rows(2, ws_em.max_row - 1)

    events_by_date = {
        d: events[events["date"] == d].reset_index(drop=True)
        for d in dates
    }

    out_row = 2
    for d in dates:
        e = events_by_date[d]
        # Keep at least three rows/day to match the official template's layout convention.
        n = max(3, len(e))
        for j in range(n):
            values = [
                pd.Timestamp(d).to_pydatetime() if j == 0 else None,
                str(e.iloc[j]["interval"]) if j < len(e) else None,
                float(e.iloc[j]["emergency_kwh"]) if j < len(e) else None,
            ]
            for c, value in enumerate(values, start=1):
                cell = ws_em.cell(out_row, c)
                cell.value = value
                st = em_style[c-1]
                cell._style = copy(st["style"])
                cell.number_format = st["number_format"]
                cell.alignment = copy(st["alignment"])
                cell.font = copy(st["font"])
                cell.fill = copy(st["fill"])
                cell.border = copy(st["border"])
                cell.protection = copy(st["protection"])
            out_row += 1

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

def verify_final(daily, slots, q_star):
    require(len(slots) == 334 * 144, "最终逐时数据不是 334×144")
    require(slots["date"].nunique() == 334, "最终日期数不是 334")

    overlap = np.minimum(
        slots["planned_charge_kwh"].to_numpy(float),
        slots["planned_discharge_kwh"].to_numpy(float),
    )
    require(overlap.max() <= 1e-7, "发现同时充放电")

    max_slot_energy = CFG["max_charge_kw"] * CFG["interval_hours"]
    require(
        slots["planned_charge_kwh"].max() <= max_slot_energy + 1e-7,
        "充电功率越界",
    )
    require(
        slots["planned_discharge_kwh"].max() <= max_slot_energy + 1e-7,
        "放电功率越界",
    )

    smin = min(
        slots["planned_storage_start_kwh"].min(),
        slots["planned_storage_end_kwh"].min(),
    )
    smax = max(
        slots["planned_storage_start_kwh"].max(),
        slots["planned_storage_end_kwh"].max(),
    )
    require(smin >= CFG["storage_min_kwh"] - 1e-6, "储能下界越界")
    require(smax <= CFG["storage_max_kwh"] + 1e-6, "储能上界越界")
    require(daily["max_balance_residual"].max() < 1e-6, "平衡残差过大")

    # The daily plan has no same-day terminal equality, but its state is
    # continuous across adjacent days.
    day_state = (
        slots.groupby("date", sort=True)
        .agg(day_start=("planned_storage_start_kwh", "first"),
             day_end=("planned_storage_end_kwh", "last"))
    )
    require(
        abs(float(day_state.iloc[0]["day_start"]) - CFG["initial_storage_kwh"]) < 1e-6,
        "首日初始储能电量不正确",
    )
    require(
        np.max(np.abs(day_state["day_start"].to_numpy()[1:] -
                      day_state["day_end"].to_numpy()[:-1])) < 1e-6,
        "相邻日期储能状态未连续传递",
    )

def run_all(
    actual_path,
    price_path,
    template_path,
    outdir,
    excel_output,
    terminal_target_kwh=None,
    terminal_penalty_yuan_per_kwh=0.0,
):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    dates, load, pv, prices = load_inputs(actual_path, price_path)
    pred_l, pred_p, scenarios = forecast_scenarios(dates, load, pv)

    q_star, scan = scan_q_star(
        dates, load, pv, prices, scenarios,
        terminal_target_kwh=terminal_target_kwh,
        terminal_penalty_yuan_per_kwh=terminal_penalty_yuan_per_kwh,
    )

    run_idx = np.where(
        (dates >= pd.Timestamp(RUN_START))
        & (dates <= pd.Timestamp(RUN_END))
    )[0]
    require(all(d in scenarios for d in run_idx), "最终回放期场景缺失")

    daily, slots = run_q(
        dates, load, pv, prices, scenarios, run_idx, q_star, True
        , terminal_target_kwh, terminal_penalty_yuan_per_kwh
    )
    events = emergency_events(slots)

    verify_final(daily, slots, q_star)

    export_result2(
        template_path=template_path,
        output_path=excel_output,
        daily=daily,
        slots=slots,
        events=events,
    )

    report = {
        "status": "PASS",
        "q_star": q_star,
        "development_period": [DEV_START, DEV_END],
        "strategy_period": [RUN_START, RUN_END],
        "days": int(daily["date"].nunique()),
        "slots": int(len(slots)),
        "planned_purchase_kwh": float(daily["planned_grid_kwh"].sum()),
        "emergency_purchase_kwh": float(daily["emergency_kwh"].sum()),
        "unused_energy_kwh": float(daily["unused_kwh"].sum()),
        "planned_cost_yuan": float(daily["planned_cost_yuan"].sum()),
        "emergency_cost_yuan": float(daily["emergency_cost_yuan"].sum()),
        "total_cost_yuan": float(daily["total_cost_yuan"].sum()),
        "result2": str(Path(excel_output).resolve()),
    }
    return report

def main():
    default_actual, default_price, default_template = default_paths()

    ap = argparse.ArgumentParser(
        description="CUMCM 2026 C题 第二问最终程序"
    )
    ap.add_argument("--actual", type=Path, default=default_actual)
    ap.add_argument("--price", type=Path, default=default_price)
    ap.add_argument("--template", type=Path, default=default_template)
    ap.add_argument(
        "--outdir",
        type=Path,
        default=SCRIPT_DIR / "results",
    )
    ap.add_argument(
        "--excel",
        type=Path,
        default=SCRIPT_DIR / "results" / "result2.xlsx",
    )
    args = ap.parse_args()

    for p, label in [
        (args.actual, "actual"),
        (args.price, "price"),
        (args.template, "template"),
    ]:
        if not p.exists():
            raise FileNotFoundError(f"{label} 文件不存在: {p}")

    t0 = time.perf_counter()
    report = run_all(
        args.actual,
        args.price,
        args.template,
        args.outdir,
        args.excel,
    )

    print(json.dumps(report, ensure_ascii=False, indent=2))
    print(f"运行耗时: {time.perf_counter() - t0:.2f} s")

if __name__ == "__main__":
    main()
