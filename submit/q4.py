"""Question 4 final runner: quadratic PCA continuation-value strategy only.

The script re-runs Q2 and Q3 with the volatile prices in attachment 4 and
writes the two official result workbooks.  It deliberately keeps no cache or
diagnostic output in the submission directory.
"""
from __future__ import annotations

import argparse
import sys
import time
from copy import copy
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.optimize import linprog, minimize_scalar, nnls
from scipy.sparse import csc_matrix, csr_matrix, eye, hstack, vstack
from sklearn.decomposition import PCA
from sklearn.linear_model import Ridge
from sklearn.preprocessing import StandardScaler

ROOT = Path(__file__).resolve().parent
DT = 1.0 / 6.0
ETA = 0.9
EMIN, EMAX, E0 = 1200.0, 10800.0, 6000.0
CAP = EMAX - EMIN
POWER = 5000.0
Q = 0.85
RUN_START = pd.Timestamp("2025-02-01")
RUN_END = pd.Timestamp("2025-12-31")
VALUE_COMPONENTS = 5
VALUE_ALPHA = 10.0
TOL = 1e-5

# The terminal-value experiment used an isolated solver environment.  Use it
# when available, while retaining the exact LP fallback for ordinary setups.
for _dep in (
    ROOT.parent / "CUMCM_2026" / ".tmp" / "q4_terminal_value" / "deps",
    ROOT / ".tmp" / "q4_terminal_value" / "deps",
):
    if _dep.exists():
        sys.path.insert(0, str(_dep))
try:
    import clarabel
except ImportError:
    clarabel = None


def require(condition, message):
    if not bool(condition):
        raise ValueError(message)


def read_prices(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [list(row) for row in wb["Sheet1"].iter_rows(values_only=True)]
    finally:
        wb.close()
    require(len(rows) == 366 and all(len(row) >= 145 for row in rows),
            "附件4格式错误")
    dates = pd.DatetimeIndex(pd.to_datetime([row[0] for row in rows[1:]]))
    require(dates.equals(pd.date_range("2025-01-01", "2025-12-31")),
            "附件4日期轴错误")
    prices = np.asarray([row[1:145] for row in rows[1:]], dtype=float)
    require(np.isfinite(prices).all() and (prices > 0).all(),
            "附件4电价含非法值")
    return dates, prices


def reward_value(reward, x):
    if reward is None:
        return 0.0
    return reward[0] * x - 0.5 * reward[1] * x * x


def q4_arrays(net, prices, initial, terminal=None, dt=DT):
    """Build the scaled daily QP used by the terminal-value experiment."""
    net = np.asarray(net, dtype=float)
    prices = np.asarray(prices, dtype=float)
    n = len(net)
    require(net.shape == prices.shape and n > 0, "日调度输入长度错误")
    m = 5 * n + 1
    t = np.arange(n)
    rows = np.r_[t, t, t, t, n + t, n + t]
    cols = np.r_[t, n + t, 2 * n + t, 3 * n + t,
                 4 * n + t + 1, 4 * n + t]
    vals = np.r_[np.ones(n), -np.ones(n), np.ones(n), -np.ones(n),
                 np.ones(n), -np.ones(n)]
    # Storage recursion coefficients for c and d.
    rows = np.r_[rows, n + t, n + t]
    cols = np.r_[cols, n + t, 2 * n + t]
    vals = np.r_[vals, -ETA * np.ones(n), np.ones(n) / ETA]
    A = csr_matrix((vals, (rows, cols)), shape=(2 * n, m))
    rhs = np.r_[net / 1000.0, np.zeros(n)]

    lo = np.zeros(m)
    hi = np.full(m, np.inf)
    hi[n:3 * n] = POWER * dt / 1000.0
    lo[4 * n:5 * n + 1] = EMIN / 1000.0
    hi[4 * n:5 * n + 1] = EMAX / 1000.0
    lo[4 * n] = hi[4 * n] = initial / 1000.0
    if terminal is not None:
        lo[5 * n] = hi[5 * n] = float(terminal) / 1000.0
    return A, rhs, lo, hi, prices, n


def clean_q4_plan(x, net, prices, initial, reward, seconds, solver):
    n = len(net)
    g, c, d, unused = [np.asarray(x[j * n:(j + 1) * n]) * 1000.0
                        for j in range(4)]
    e = np.asarray(x[4 * n:5 * n + 1]) * 1000.0
    for a in (g, c, d, unused):
        require(a.min() >= -0.01, "二次调度产生负决策")
        a[a < 0] = 0
    # Remove the zero-net-energy cycle that can occur on a flat LP face.
    overlap = np.minimum(c, d / ETA ** 2)
    c -= overlap
    d -= ETA ** 2 * overlap
    unused += (1 - ETA ** 2) * overlap
    balance = g + d - c - unused - net
    recursion = np.diff(e) - ETA * c + d / ETA
    require(max(abs(balance).max(), abs(recursion).max()) < 0.01,
            "二次调度物理约束校验失败")
    return {
        "grid": g, "charge": c, "discharge": d, "unused": unused,
        "storage": e, "solver": solver, "seconds": seconds,
        "objective": float(prices @ g - reward_value(reward, e[-1] - EMIN)),
    }


def solve_q4_daily(net, prices, initial=E0, reward=None, terminal=None, dt=DT):
    """Solve the winning concave terminal-reward daily scheduling problem."""
    start = time.perf_counter()
    A, rhs, lo, hi, prices, n = q4_arrays(net, prices, initial, terminal, dt)
    reward = reward or (0.0, 0.0)
    q = np.zeros(len(lo))
    q[:n] = prices
    q[5 * n] = -(reward[0] + reward[1] * EMIN)
    diag = np.zeros(len(lo))
    # E is represented in MWh in this QP; the original reward uses kWh.
    diag[5 * n] = reward[1] * 1000.0

    if reward[1] > 1e-12 and clarabel is not None:
        upper = np.flatnonzero(np.isfinite(hi))
        lower = np.flatnonzero(np.isfinite(lo))
        AA = vstack([A, eye(len(lo), format="csc")[upper, :],
                     -eye(len(lo), format="csc")[lower, :]], format="csc")
        bb = np.r_[rhs, hi[upper], -lo[lower]]
        P = csc_matrix((diag, (np.arange(len(lo)), np.arange(len(lo)))),
                       shape=(len(lo), len(lo)))
        settings = clarabel.DefaultSettings()
        settings.verbose = False
        settings.tol_gap_abs = 1e-8
        settings.tol_gap_rel = 1e-8
        settings.tol_feas = 1e-8
        settings.max_iter = 200
        cones = [clarabel.ZeroConeT(len(rhs)),
                 clarabel.NonnegativeConeT(len(bb) - len(rhs))]
        solved = clarabel.DefaultSolver(P, q, AA, bb, cones, settings).solve()
        if str(solved.status) in ("Solved", "AlmostSolved"):
            return clean_q4_plan(np.asarray(solved.x), net, prices, initial,
                                 reward, time.perf_counter() - start, "Clarabel")

    def solve_at(end_kwh):
        lower, upper = lo.copy(), hi.copy()
        lower[5 * n] = upper[5 * n] = float(end_kwh) / 1000.0
        result = linprog(q, A_eq=A, b_eq=rhs,
                         bounds=np.c_[lower, upper], method="highs")
        require(result.success, "二次调度LP回退失败: " + result.message)
        plan = clean_q4_plan(result.x, net, prices, initial, reward,
                             time.perf_counter() - start, "LP")
        plan["objective"] = float(prices @ plan["grid"] -
                                   reward_value(reward, end_kwh - EMIN))
        return plan

    if terminal is not None:
        return solve_at(terminal)
    lower = max(EMIN, initial - POWER * dt * n / ETA)
    upper = min(EMAX, initial + ETA * POWER * dt * n)
    cache = {}

    def objective(end_kwh):
        key = round(float(end_kwh), 7)
        if key not in cache:
            cache[key] = solve_at(key)
        return cache[key]["objective"]

    result = minimize_scalar(objective, bounds=(lower, upper), method="bounded",
                             options={"xatol": 1e-3, "maxiter": 50})
    for end in (lower, upper, result.x):
        objective(end)
    return min(cache.values(), key=lambda x: x["objective"])


def historical_future(data, d):
    ref = np.arange(d - 6, d)
    return {k: np.asarray(data[k][ref], dtype=float).copy()
            for k in ("prices", "load", "pv")}


def value_features(future):
    p = future["prices"].reshape(6, 24, 6).mean(axis=2)
    net = (future["load"] - future["pv"]).reshape(6, 24, 6).mean(axis=2)
    parts = []
    for k in range(6):
        parts.extend([p[k].mean(), p[k].min(), p[k].max(),
                      np.maximum(net[k], 0).sum(),
                      np.maximum(-net[k], 0).sum()])
    parts.extend([p[0, :6].mean(), np.maximum(net[0, :6], 0).sum(),
                  np.argmin(p[0])])
    return {"price": p.ravel(), "net": net.ravel(), "economic": np.asarray(parts)}


def teacher(future):
    net = (future["load"] - future["pv"]).reshape(6, 24, 6).mean(axis=2).ravel()
    prices = future["prices"].reshape(6, 24, 6).mean(axis=2).ravel()
    costs = np.asarray([
        solve_q4_daily(net, prices, EMIN + x, reward=(0.0, 0.0),
                       terminal=E0, dt=1.0)["objective"]
        for x in np.linspace(0.0, CAP, 7)
    ])
    values = costs[0] - costs
    grid = np.linspace(0.0, CAP, 7)
    z = grid / CAP
    weights, _ = nnls(np.column_stack([z, z - 0.5 * z * z]), values / CAP)
    a = float(weights.sum())
    b = float(weights[1] / CAP)
    return {"values": values, "a": a, "b": b,
            "marginal_low": a, "marginal_high": a - b * CAP,
            "features": value_features(future)}


class ValueModel:
    def __init__(self):
        self.components = VALUE_COMPONENTS
        self.alpha = VALUE_ALPHA

    def fit(self, rows):
        self.train_indices = tuple(r["index"] for r in rows)
        self.pca_price = PCA(n_components=self.components, svd_solver="full")
        self.pca_net = PCA(n_components=self.components, svd_solver="full")
        price = self.pca_price.fit_transform(np.stack([r["features"]["price"] for r in rows]))
        net = self.pca_net.fit_transform(np.stack([r["features"]["net"] for r in rows]))
        X = np.column_stack([price, net])
        self.scaler = StandardScaler().fit(X)
        y = np.asarray([[r["marginal_low"], r["marginal_high"]] for r in rows])
        self.regressor = Ridge(alpha=self.alpha).fit(self.scaler.transform(X), y)
        return self

    def predict(self, rows):
        require(min(r["index"] for r in rows) > max(self.train_indices),
                "价值模型使用了未来标签")
        price = self.pca_price.transform(np.stack([r["features"]["price"] for r in rows]))
        net = self.pca_net.transform(np.stack([r["features"]["net"] for r in rows]))
        y = self.regressor.predict(self.scaler.transform(np.column_stack([price, net])))
        rewards = []
        for low, high in y:
            low, high = float(low), float(high)
            if low < high:
                low = high = 0.5 * (low + high)
            low, high = max(0.0, low), max(0.0, high)
            rewards.append((low, (low - high) / CAP))
        return rewards


def build_labels(dates, load, pv, prices):
    labels = {}
    for d in range(7, len(dates)):
        future = historical_future({"load": load, "pv": pv, "prices": prices}, d)
        labels[d] = teacher(future) | {
            "index": d, "date": dates[d], "features": value_features(future)
        }
        if d % 30 == 0:
            print(f"价值标签 {d}/{len(dates) - 1}", flush=True)
    return labels


def monthly_rewards(labels, dates, indices):
    result = {}
    for month in sorted({dates[d].month for d in indices}):
        month_idx = [d for d in indices if dates[d].month == month]
        first = month_idx[0]
        train = [labels[d] for d in sorted(labels) if d < first]
        require(len(train) >= VALUE_COMPONENTS + 1, "价值模型历史标签不足")
        result.update(zip(month_idx, ValueModel().fit(train).predict(
            [labels[d] for d in month_idx])))
    return result


def execute(grid, actual_net, initial):
    n = len(grid)
    e = np.empty(n + 1)
    e[0] = initial
    charge = np.zeros(n)
    discharge = np.zeros(n)
    emergency = np.zeros(n)
    unused = np.zeros(n)
    mc = POWER * DT
    for t in range(n):
        surplus = grid[t] - actual_net[t]
        if surplus >= 0:
            charge[t] = min(surplus, mc, max(0.0, (EMAX - e[t]) / ETA))
            unused[t] = surplus - charge[t]
        else:
            discharge[t] = min(-surplus, mc, max(0.0, (e[t] - EMIN) * ETA))
            emergency[t] = -surplus - discharge[t]
        e[t + 1] = e[t] + ETA * charge[t] - discharge[t] / ETA
    require(e.min() >= EMIN - 1e-6 and e.max() <= EMAX + 1e-6,
            "实际执行储能越界")
    return charge, discharge, emergency, unused, e


def read_all_inputs():
    import q2
    import q3

    dates, load, pv, fixed_prices, forecasts = q3.read_inputs(
        ROOT / "Data" / "附件2.xlsx",
        ROOT / "Data" / "附件1.xlsx",
        ROOT / "Data" / "附件3.xlsx",
    )
    _, prices = read_prices(ROOT / "Data" / "附件4.xlsx")
    pred_load, pred_pv, scenarios = q2.forecast_scenarios(dates, load, pv)
    return dates, load, pv, fixed_prices, prices, forecasts, pred_load, pred_pv, scenarios


def summarize(x):
    return {
        "date": x["date"].iloc[0],
        "planned_kwh": float(x["planned_grid_kwh"].sum()),
        "updated_kwh": float(x["updated_grid_kwh"].sum()),
        "emergency_kwh": float(x["emergency_kwh"].sum()),
        "planned_cost": float(x["planned_cost"].sum()),
        "adjustment_cost": float(x["adjustment_cost"].sum()),
        "emergency_cost": float(x["emergency_cost"].sum()),
        "total_cost": float(x[["planned_cost", "adjustment_cost",
                               "emergency_cost"]].to_numpy().sum()),
        "state_start": float(x["storage_start_kwh"].iloc[0]),
        "state_end": float(x["storage_end_kwh"].iloc[-1]),
    }


def run_q4_2(dates, load, pv, prices, scenarios, decision_prices, rewards):
    import q2

    daily, details = [], []
    state = E0
    for d in np.flatnonzero((dates >= RUN_START) & (dates <= RUN_END)):
        net = np.quantile(scenarios[d], Q, axis=0) * DT
        plan = solve_q4_daily(net, decision_prices[d], state, rewards[d])
        charge, discharge, emergency, unused, actual_storage = execute(
            plan["grid"], (load[d] - pv[d]) * DT, state)
        actual_price = prices[d]
        date = str(dates[d].date())
        x = pd.DataFrame({
            "date": date, "slot": np.arange(1, 145),
            "price_forecast": decision_prices[d], "price": actual_price,
            "actual_load_kwh": load[d] * DT, "actual_pv_kwh": pv[d] * DT,
            "planned_grid_kwh": plan["grid"], "updated_grid_kwh": plan["grid"],
            "planned_charge_kwh": plan["charge"],
            "planned_discharge_kwh": plan["discharge"],
            "charge_kwh": charge, "discharge_kwh": discharge,
            "emergency_kwh": emergency, "unused_kwh": unused,
            "planned_storage_start_kwh": plan["storage"][:-1],
            "planned_storage_end_kwh": plan["storage"][1:],
            "storage_start_kwh": actual_storage[:-1],
            "storage_end_kwh": actual_storage[1:],
            "planned_cost": actual_price * plan["grid"],
            "adjustment_cost": np.zeros(144),
            "emergency_cost": 5.0 * actual_price * emergency,
        })
        daily.append(summarize(x))
        details.append(x)
        state = float(actual_storage[-1])
    return pd.DataFrame(daily), pd.concat(details, ignore_index=True)


def run_q4_3(dates, load, pv, fixed_prices, prices, forecasts, pred_load,
             rewards):
    import q3

    daily, details = [], []
    state = E0
    first_forecast_day = 14
    for d, day in enumerate(dates):
        in_eval = RUN_START <= day <= RUN_END
        decision_price = fixed_prices if not in_eval else prices[d - 7]
        baseline_a, baseline_w = q3.baseline2_deterministic_scenario(
            d, 0, pred_load, forecasts, load, first_forecast_day)
        if d == 0:
            baseline = np.zeros(144)
        elif not in_eval:
            baseline = q3.solve_problem(
                baseline_a, decision_price, state, risk=True,
                weights=baseline_w)["grid"]
        else:
            baseline = solve_q3_rewarded(
                baseline_a, decision_price, state, rewards[d])["grid"]

        updated = baseline.copy()
        blocks = []
        update_hours = (6, 12, 18)
        for hour in (0, 6, 12, 18):
            start, end = hour * 6, hour * 6 + 36
            if hour in update_hours and d > 0:
                a, weights = q3.issue_scenarios(
                    d, hour, pred_load, forecasts, load, pv,
                    first_forecast_day, True)
                if in_eval:
                    candidate = solve_q3_rewarded(
                        a, decision_price[start:], state, rewards[d],
                        fixed=baseline[start:], lock_from=36,
                        weights=weights)["grid"]
                else:
                    candidate = q3.solve_problem(
                        a, decision_price[start:], state, fixed=baseline[start:],
                        lock_from=36, risk=True, weights=weights)["grid"]
                updated[start:end] = candidate[:end - start]
            c, dis, emergency, unused, e = q3.execute(
                updated[start:end], load[d, start:end] * DT,
                pv[d, start:end] * DT, state, battery=(d > 0))
            state = float(e[-1])
            blocks.append((c, dis, emergency, unused, e))
        if not in_eval:
            continue
        c, dis, emergency, unused = [np.concatenate([b[j] for b in blocks])
                                     for j in range(4)]
        storage_start = np.concatenate([b[4][:-1] for b in blocks])
        storage_end = np.concatenate([b[4][1:] for b in blocks])
        actual_price = prices[d]
        adjustment = actual_price * (
            1.5 * np.maximum(updated - baseline, 0.0) +
            0.5 * np.maximum(baseline - updated, 0.0))
        date = str(day.date())
        x = pd.DataFrame({
            "date": date, "slot": np.arange(1, 145),
            "price_forecast": decision_price, "price": actual_price,
            "actual_load_kwh": load[d] * DT, "actual_pv_kwh": pv[d] * DT,
            "planned_grid_kwh": baseline, "updated_grid_kwh": updated,
            "charge_kwh": c, "discharge_kwh": dis,
            "emergency_kwh": emergency, "unused_kwh": unused,
            "storage_start_kwh": storage_start, "storage_end_kwh": storage_end,
            "planned_cost": actual_price * baseline,
            "adjustment_cost": adjustment,
            "emergency_cost": 5.0 * actual_price * emergency,
        })
        daily.append(summarize(x))
        details.append(x)
    return pd.DataFrame(daily), pd.concat(details, ignore_index=True)


def q3_problem_arrays(a, p, initial, fixed=None, lock_from=None, weights=None):
    a, p = np.asarray(a, float), np.asarray(p, float)
    scenarios, n = a.shape
    weights = np.ones(scenarios) / scenarios if weights is None else np.asarray(weights, float)
    require(len(weights) == scenarios and np.isclose(weights.sum(), 1),
            "Q3情景权重错误")
    h0 = 4 * n + 1
    if fixed is None:
        inc = dec = None
        m = h0 + scenarios * n
    else:
        inc = h0 + scenarios * n
        dec = inc + n
        m = dec + n
    q = np.zeros(m)
    if fixed is None:
        q[:n] = p
    else:
        q[inc:dec] = 1.5 * p
        q[dec:m] = 0.5 * p
    q[h0:h0 + scenarios * n] = np.repeat(weights, n) * np.tile(5.0 * p, scenarios)
    rows, cols, vals = [], [], []
    t = np.arange(n)
    rows.extend(np.r_[t, t, t, t])
    cols.extend(np.r_[n + t, 2 * n + t, 3 * n + t, 3 * n + t + 1])
    vals.extend(np.r_[-ETA * np.ones(n), np.ones(n) / ETA,
                      -np.ones(n), np.ones(n)])
    if fixed is not None:
        rows.extend(np.r_[n + t, n + t, n + t])
        cols.extend(np.r_[t, inc + t, dec + t])
        vals.extend(np.r_[np.ones(n), -np.ones(n), np.ones(n)])
        eq_rhs = np.r_[np.zeros(n), fixed]
    else:
        eq_rhs = np.zeros(n)
    Aeq = csr_matrix((vals, (rows, cols)), shape=(len(eq_rhs), m))
    scenario_rows = np.arange(scenarios * n)
    row_t = np.tile(t, scenarios)
    Aub = csr_matrix((np.r_[-np.ones(scenarios * n),
                             np.ones(scenarios * n),
                             -np.ones(scenarios * n),
                             -np.ones(scenarios * n)],
                      (np.tile(scenario_rows, 4),
                       np.r_[row_t, n + row_t, 2 * n + row_t,
                             h0 + scenario_rows])),
                     shape=(scenarios * n, m))
    bub = -a.ravel()
    lo = np.zeros(m)
    hi = np.full(m, np.inf)
    hi[n:2 * n] = POWER * DT
    hi[2 * n:3 * n] = POWER * DT
    lo[3 * n:4 * n + 1] = EMIN
    hi[3 * n:4 * n + 1] = EMAX
    lo[3 * n] = hi[3 * n] = initial
    if fixed is not None and lock_from is not None:
        lo[lock_from:n] = hi[lock_from:n] = fixed[lock_from:]
    return Aeq, eq_rhs, Aub, bub, lo, hi, q, m, n


def solve_q3_rewarded(a, p, initial, reward, fixed=None, lock_from=None,
                      weights=None):
    Aeq, eq_rhs, Aub, bub, lo, hi, q, m, n = q3_problem_arrays(
        a, p, initial, fixed, lock_from, weights)
    Pdiag = np.zeros(m)
    Pdiag[4 * n] = reward[1]
    q[4 * n] = -(reward[0] + reward[1] * EMIN)

    if reward[1] > 1e-12 and clarabel is not None:
        upper = np.flatnonzero(np.isfinite(hi))
        lower = np.flatnonzero(np.isfinite(lo))
        ident = eye(m, format="csc")
        AA = vstack([Aeq, Aub, ident[upper, :], -ident[lower, :]], format="csc")
        bb = np.r_[eq_rhs, bub, hi[upper], -lo[lower]]
        P = csc_matrix((Pdiag, (np.arange(m), np.arange(m))), shape=(m, m))
        settings = clarabel.DefaultSettings()
        settings.verbose = False
        settings.tol_gap_abs = 1e-8
        settings.tol_gap_rel = 1e-8
        settings.tol_feas = 1e-8
        settings.max_iter = 200
        cones = [clarabel.ZeroConeT(len(eq_rhs)),
                 clarabel.NonnegativeConeT(len(bb) - len(eq_rhs))]
        result = clarabel.DefaultSolver(P, q, AA, bb, cones, settings).solve()
        if str(result.status) in ("Solved", "AlmostSolved"):
            x = np.asarray(result.x)
            require(np.max(np.abs(Aeq @ x - eq_rhs)) < 0.02,
                    "Q3二次调度等式约束失败")
            return {"grid": np.maximum(x[:n], 0.0), "solver": "Clarabel"}

    def solve_at(end):
        low, high = lo.copy(), hi.copy()
        low[4 * n] = high[4 * n] = end
        result = linprog(q, A_ub=Aub, b_ub=bub, A_eq=Aeq, b_eq=eq_rhs,
                         bounds=np.c_[low, high], method="highs")
        require(result.success, "Q3二次调度LP回退失败: " + result.message)
        return result.x, float(result.fun - reward_value(reward, end - EMIN))

    # Fixed future grid values in Q3 adjustments can make the physical
    # reachability interval much narrower than the unconstrained battery
    # interval.  Derive the feasible interval from the actual LP constraints.
    end_objective = np.zeros(m)
    end_objective[4 * n] = 1.0
    low_result = linprog(end_objective, A_ub=Aub, b_ub=bub,
                         A_eq=Aeq, b_eq=eq_rhs,
                         bounds=np.c_[lo, hi], method="highs")
    high_result = linprog(-end_objective, A_ub=Aub, b_ub=bub,
                          A_eq=Aeq, b_eq=eq_rhs,
                          bounds=np.c_[lo, hi], method="highs")
    require(low_result.success and high_result.success,
            "Q3二次调度无法找到可行期末储能区间")
    low = float(low_result.x[4 * n])
    high = float(high_result.x[4 * n])
    cache = {}

    def objective(end):
        key = round(float(end), 6)
        if key not in cache:
            cache[key] = solve_at(key)
        return cache[key][1]

    optimum = minimize_scalar(objective, bounds=(low, high), method="bounded",
                              options={"xatol": 1e-3, "maxiter": 40})
    for end in (low, high, optimum.x):
        objective(end)
    x, _ = min(cache.values(), key=lambda v: v[1])
    return {"grid": np.maximum(x[:n], 0.0), "solver": "LP"}


def events_for_day(x):
    values = x.sort_values("slot")["emergency_kwh"].to_numpy()
    events, t = [], 0
    while t < 144:
        if values[t] <= 1e-7:
            t += 1
            continue
        start = t
        while t < 144 and values[t] > 1e-7:
            t += 1
        def clock(slot):
            minute = slot * 10
            return "24:00" if minute == 1440 else f"{minute // 60:02d}:{minute % 60:02d}"
        events.append((f"{clock(start)}-{clock(t)}",
                       float(values[start:t].sum())))
    return events


def copy_style(source, target):
    target._style = copy(source._style)
    target.number_format = source.number_format
    target.alignment = copy(source.alignment)
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.protection = copy(source.protection)


def export_result(template_path, output_path, daily, detail, q3_result=False):
    from openpyxl import load_workbook

    wb = load_workbook(template_path)
    dates = sorted(daily["date"].unique())
    dm = daily.set_index("date")
    by_date = {d: detail[detail.date.eq(d)].sort_values("slot") for d in dates}

    def write_plan(sheet, column, total_column):
        ws = wb[sheet]
        for i, date in enumerate(dates, 2):
            x = by_date[date]
            values = x[column].to_numpy(float)
            ws.cell(i, 1).value = pd.Timestamp(date).to_pydatetime()
            for j, value in enumerate(np.r_[values[1:], values[:1]], 2):
                ws.cell(i, j).value = float(value)
            ws.cell(i, 146).value = float(values.sum())
            ws.cell(i, 147).value = float(dm.loc[date, total_column])

    write_plan("计划购电量", "planned_grid_kwh", "planned_cost")
    if q3_result:
        write_plan("调整购电量", "updated_grid_kwh", "total_cost")

    ws = wb["充放电量"]
    prototype = [[copy(ws.cell(r, c)) for c in range(1, 7)] for r in range(2, 8)]
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)
    intervals = ["0:00-4:00", "4:00-8:00", "8:00-12:00",
                 "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    bat_cols = ("charge_kwh", "discharge_kwh", "storage_start_kwh",
                "storage_end_kwh")
    if not q3_result:
        bat_cols = ("planned_charge_kwh", "planned_discharge_kwh",
                    "planned_storage_start_kwh", "planned_storage_end_kwh")
    row = 2
    for date in dates:
        x = by_date[date]
        for block, interval in enumerate(intervals):
            z = x.iloc[block * 24:(block + 1) * 24]
            values = [
                pd.Timestamp(date).to_pydatetime() if block == 0 else None,
                interval, float(z[bat_cols[0]].sum()), float(z[bat_cols[1]].sum()),
                "0:00" if block == 0 else ("24:00" if block == 1 else None),
                float(x[bat_cols[2]].iloc[0] if block == 0
                      else x[bat_cols[3]].iloc[-1] if block == 1 else 0.0),
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row, col)
                cell.value = value
                copy_style(prototype[block][col - 1], cell)
            row += 1

    ws = wb["紧急购电量"]
    prototype = [copy(ws.cell(2, c)) for c in range(1, 4)]
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)
    row = 2
    for date in dates:
        events = events_for_day(by_date[date])
        for j in range(max(3, len(events))):
            values = [pd.Timestamp(date).to_pydatetime() if j == 0 else None,
                      events[j][0] if j < len(events) else None,
                      events[j][1] if j < len(events) else None]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row, col)
                cell.value = value
                copy_style(prototype[col - 1], cell)
            row += 1
    wb.save(output_path)


def verify(daily, detail, q3_result=False):
    require(len(daily) == 334 and len(detail) == 334 * 144,
            "结果日期或时段数量错误")
    balance = (detail.updated_grid_kwh + detail.actual_pv_kwh +
               detail.discharge_kwh + detail.emergency_kwh -
               detail.actual_load_kwh - detail.charge_kwh - detail.unused_kwh)
    storage = (detail.storage_end_kwh - detail.storage_start_kwh -
               ETA * detail.charge_kwh + detail.discharge_kwh / ETA)
    require(float(np.abs(balance).max()) < 1e-4, "供需平衡校验失败")
    require(float(np.abs(storage).max()) < 1e-4, "储能递推校验失败")
    require(detail.storage_start_kwh.min() >= EMIN - 1e-4 and
            detail.storage_end_kwh.max() <= EMAX + 1e-4, "储能边界校验失败")
    require(detail.charge_kwh.max() <= POWER * DT + 1e-4 and
            detail.discharge_kwh.max() <= POWER * DT + 1e-4, "储能功率校验失败")
    require(np.allclose(detail.planned_cost,
                        detail.price * detail.planned_grid_kwh),
            "计划费用复算失败")
    require(np.allclose(detail.emergency_cost,
                        5.0 * detail.price * detail.emergency_kwh),
            "紧急费用复算失败")
    if q3_result:
        expected = detail.price * (
            1.5 * np.maximum(detail.updated_grid_kwh - detail.planned_grid_kwh, 0) +
            0.5 * np.maximum(detail.planned_grid_kwh - detail.updated_grid_kwh, 0))
        require(np.allclose(detail.adjustment_cost, expected),
                "调整费用复算失败")
    return {"rows": len(detail), "max_balance": float(np.abs(balance).max()),
            "max_storage": float(np.abs(storage).max()),
            "total_cost": float(daily.total_cost.sum())}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=ROOT / "results")
    args = parser.parse_args()
    started = time.perf_counter()
    sys.path.insert(0, str(ROOT))
    (dates, load, pv, fixed_prices, prices, forecasts, pred_load, pred_pv,
     scenarios) = read_all_inputs()
    labels = build_labels(dates, load, pv, prices)
    eval_indices = list(np.flatnonzero((dates >= RUN_START) & (dates <= RUN_END)))
    rewards = monthly_rewards(labels, dates, eval_indices)
    decision_prices = np.vstack([prices[max(0, d - 7)] for d in range(len(dates))])

    q42_daily, q42_detail = run_q4_2(
        dates, load, pv, prices, scenarios, decision_prices, rewards)
    q43_daily, q43_detail = run_q4_3(
        dates, load, pv, fixed_prices, prices, forecasts, pred_load, rewards)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    out42 = args.output_dir / "result4-2.xlsx"
    out43 = args.output_dir / "result4-3.xlsx"
    export_result(ROOT / "Data" / "附件5" / "result4-2.xlsx", out42,
                  q42_daily, q42_detail, False)
    export_result(ROOT / "Data" / "附件5" / "result4-3.xlsx", out43,
                  q43_daily, q43_detail, True)
    check42 = verify(q42_daily, q42_detail)
    check43 = verify(q43_daily, q43_detail, True)
    print({"result4-2": str(out42.resolve()), "result4-3": str(out43.resolve()),
           "q4-2": check42, "q4-3": check43,
           "seconds": round(time.perf_counter() - started, 2)})


if __name__ == "__main__":
    main()
